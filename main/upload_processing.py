from __future__ import annotations

import base64
import io
import os
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable

import matplotlib
import numpy as np
from docx import Document
from pypdf import PdfReader

try:  # PyMuPDF is only required when a low-text/scanned PDF needs VLM parsing.
    import fitz  # type: ignore
except Exception:  # pragma: no cover - optional runtime dependency
    fitz = None

from deepem.assets import AssetManager, AssetRecord
from deepem.protocol import new_id

matplotlib.use("Agg")
from matplotlib.figure import Figure  # noqa: E402


ProgressCallback = Callable[[dict[str, Any]], None]


class UploadProcessingError(RuntimeError):
    pass


class OCRProvider:
    def extract_text(self, *, asset: AssetRecord, content: bytes) -> str:
        raise NotImplementedError


class NoopOCRProvider(OCRProvider):
    def extract_text(self, *, asset: AssetRecord, content: bytes) -> str:
        return ""


@dataclass(slots=True)
class RenderedPDFPage:
    page_number: int
    mime_type: str
    content: bytes


class VLMProvider:
    def extract_pdf_pages_as_markdown(
        self,
        *,
        asset: AssetRecord,
        pages: list[RenderedPDFPage],
        progress_callback: ProgressCallback | None = None,
        start_percent: int = 55,
        end_percent: int = 92,
    ) -> list[tuple[str, str | None, str | None]]:
        raise NotImplementedError


class NoopVLMProvider(VLMProvider):
    def extract_pdf_pages_as_markdown(
        self,
        *,
        asset: AssetRecord,
        pages: list[RenderedPDFPage],
        progress_callback: ProgressCallback | None = None,
        start_percent: int = 55,
        end_percent: int = 92,
    ) -> list[tuple[str, str | None, str | None]]:
        return []


class LLMVLMProvider(VLMProvider):
    """Use the same OpenAI-compatible chat LLM client for scanned-PDF OCR.

    The platform chat agent already sends messages through ``llm_client.complete``.
    For scanned PDFs we reuse that exact client and send each rendered PDF page as a
    multimodal message, asking the model to return Markdown only.
    """

    def __init__(self, llm_client: Any) -> None:
        self.llm_client = llm_client

    def extract_pdf_pages_as_markdown(
        self,
        *,
        asset: AssetRecord,
        pages: list[RenderedPDFPage],
        progress_callback: ProgressCallback | None = None,
        start_percent: int = 55,
        end_percent: int = 92,
    ) -> list[tuple[str, str | None, str | None]]:
        if not pages:
            return []
        if self.llm_client is None or self.llm_client.__class__.__name__ == "LocalWorkflowLLMClient":
            raise UploadProcessingError("扫描 PDF 需要可用的多模态 LLM 配置，当前未启用在线模型。")

        extracted: list[tuple[str, str | None, str | None]] = []
        total = len(pages)
        span = max(1, end_percent - start_percent)
        for index, page in enumerate(pages, start=1):
            current_percent = start_percent + int(span * (index - 1) / total)
            emit_progress(
                progress_callback,
                current_percent,
                "vlm_page_start",
                f"VLM 正在识别第 {page.page_number} 页（{index}/{total}）",
                file_name=asset.file_name,
                page=page.page_number,
            )
            try:
                markdown = self._recognize_page(asset=asset, page=page)
            except UploadProcessingError:
                raise
            except Exception as exc:
                raise UploadProcessingError(f"VLM 识别第 {page.page_number} 页失败：{exc}") from exc

            markdown = normalize_vlm_markdown(markdown)
            if markdown:
                page_markdown = f"## 第 {page.page_number} 页\n\n{markdown}" if total > 1 else markdown
                extracted.append((page_markdown, str(page.page_number), None))
            emit_progress(
                progress_callback,
                start_percent + int(span * index / total),
                "vlm_page_done",
                f"VLM 已完成第 {page.page_number} 页（{index}/{total}）",
                file_name=asset.file_name,
                page=page.page_number,
            )
        return extracted

    def _recognize_page(self, *, asset: AssetRecord, page: RenderedPDFPage) -> str:
        data_url = encode_image_as_data_url(page.mime_type, page.content)
        messages = [
            {
                "role": "system",
                "content": (
                    "你是严谨的文档 OCR 与版面转写助手。请从用户提供的 PDF 页面图片中提取全部可见文字，"
                    "统一输出 Markdown。保留标题层级、列表、表格、代码块、公式/符号和阅读顺序；"
                    "不要编造图片中不存在的内容；不要输出解释、寒暄或 JSON。"
                ),
            },
            {
                "role": "user",
                "content": [
                    {
                        "type": "text",
                        "text": (
                            f"文件名：{asset.file_name}\n页码：{page.page_number}\n"
                            "请将这页扫描 PDF 图片完整转写为 Markdown。"
                            "如果某处文字无法辨认，请用 `<!-- unreadable -->` 标注。"
                        ),
                    },
                    {"type": "image_url", "image_url": {"url": data_url}},
                ],
            },
        ]
        response = self.llm_client.complete(
            messages=messages,
            tools=[],
            temperature=0.0,
            generation_options={"enable_thinking": False},
        )
        return str(getattr(response, "content", "") or "")


@dataclass(slots=True)
class ParsedChunk:
    text: str
    chunk_index: int
    page: str | None = None
    sheet: str | None = None


@dataclass(slots=True)
class UploadProcessingResult:
    asset: AssetRecord
    upload_kind: str
    summary: str
    ocr_pending: bool = False
    preview_asset: AssetRecord | None = None
    database_id: str | None = None
    chunks: list[ParsedChunk] = field(default_factory=list)

    def to_client_payload(self, asset_manager: AssetManager) -> dict[str, Any]:
        payload = asset_manager.to_client_payload(self.asset)
        payload.update(
            {
                "summary": self.summary,
                "ocr_pending": self.ocr_pending,
                "database_id": self.database_id,
                "preview_asset_id": self.preview_asset.asset_id if self.preview_asset else None,
            }
        )
        if self.preview_asset is not None:
            payload["preview_url"] = asset_manager.public_url(self.preview_asset.asset_id)
        return payload


class UploadProcessor:
    IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif"}
    DOCUMENT_SUFFIXES = {".pdf", ".md", ".txt", ".doc", ".docx"}
    SIGNAL_SUFFIXES = {".bin"}

    def __init__(
        self,
        *,
        asset_manager: AssetManager,
        document_index: Any,
        ocr_provider: OCRProvider | None = None,
        vlm_provider: VLMProvider | None = None,
        llm_client: Any | None = None,
        text_threshold: int | None = None,
        pdf_render_dpi: int | None = None,
        pdf_vlm_max_pages: int | None = None,
    ) -> None:
        self.asset_manager = asset_manager
        self.document_index = document_index
        self.ocr_provider = ocr_provider or NoopOCRProvider()
        self.vlm_provider = vlm_provider or (LLMVLMProvider(llm_client) if llm_client is not None else NoopVLMProvider())
        self.text_threshold = int(text_threshold if text_threshold is not None else os.getenv("DEEPEM_DOCUMENT_TEXT_THRESHOLD", "80"))
        self.pdf_render_dpi = int(pdf_render_dpi if pdf_render_dpi is not None else os.getenv("DEEPEM_PDF_RENDER_DPI", "160"))
        self.pdf_vlm_max_pages = int(pdf_vlm_max_pages if pdf_vlm_max_pages is not None else os.getenv("DEEPEM_PDF_VLM_MAX_PAGES", "0"))

    def process_upload(
        self,
        *,
        file_name: str,
        content: bytes,
        conversation_id: str | None = None,
        progress_callback: ProgressCallback | None = None,
    ) -> UploadProcessingResult:
        suffix = Path(file_name).suffix.lower()
        emit_progress(progress_callback, 0, "received", f"已接收文件 {Path(file_name).name}", file_name=file_name)
        if suffix in self.IMAGE_SUFFIXES:
            return self._process_image(file_name=file_name, content=content, conversation_id=conversation_id, progress_callback=progress_callback)
        if suffix in self.DOCUMENT_SUFFIXES:
            return self._process_document(file_name=file_name, content=content, conversation_id=conversation_id, progress_callback=progress_callback)
        if suffix in self.SIGNAL_SUFFIXES:
            return self._process_signal(file_name=file_name, content=content, conversation_id=conversation_id, progress_callback=progress_callback)
        raise UploadProcessingError(f"暂不支持该文件类型：{suffix or '<无后缀>'}")

    def build_index_payloads(self, result: UploadProcessingResult) -> list[dict[str, Any]]:
        file_payload = {
            "document_id": f"{result.asset.asset_id}:file",
            "asset_id": result.asset.asset_id,
            "record_type": "file",
            "file_name": result.asset.file_name,
            "suffix": result.asset.suffix,
            "conversation_id": result.asset.conversation_id,
            "created_at": result.asset.created_at,
            "summary": result.summary,
            "text": result.summary,
            "mime_type": result.asset.mime_type,
            "upload_kind": result.upload_kind,
            "preview_asset_id": result.preview_asset.asset_id if result.preview_asset else None,
            "preview_url": self.asset_manager.public_url(result.preview_asset.asset_id) if result.preview_asset else None,
            "parse_method": result.asset.metadata.get("parse_method"),
            "text_char_count": result.asset.metadata.get("text_char_count"),
        }
        payloads = [file_payload]
        for chunk in result.chunks:
            payloads.append(
                {
                    "document_id": f"{result.asset.asset_id}:chunk:{chunk.chunk_index:04d}",
                    "asset_id": result.asset.asset_id,
                    "record_type": "chunk",
                    "file_name": result.asset.file_name,
                    "suffix": result.asset.suffix,
                    "conversation_id": result.asset.conversation_id,
                    "created_at": result.asset.created_at,
                    "summary": result.summary,
                    "text": chunk.text,
                    "chunk_index": chunk.chunk_index,
                    "page": chunk.page,
                    "sheet": chunk.sheet,
                    "mime_type": result.asset.mime_type,
                    "upload_kind": result.upload_kind,
                    "preview_asset_id": result.preview_asset.asset_id if result.preview_asset else None,
                    "preview_url": self.asset_manager.public_url(result.preview_asset.asset_id) if result.preview_asset else None,
                    "parse_method": result.asset.metadata.get("parse_method"),
                    "text_char_count": result.asset.metadata.get("text_char_count"),
                }
            )
        return payloads

    def _process_image(
        self,
        *,
        file_name: str,
        content: bytes,
        conversation_id: str | None,
        progress_callback: ProgressCallback | None = None,
    ) -> UploadProcessingResult:
        emit_progress(progress_callback, 10, "saving", f"正在保存图片 {Path(file_name).name}", file_name=file_name)
        asset = self.asset_manager.save_bytes(
            asset_id=new_id("asset"),
            file_name=file_name,
            content=content,
            upload_kind="image",
            conversation_id=conversation_id,
            metadata={"parse_method": "direct_image", "markdown": False},
        )
        summary = f"已上传图片 {asset.file_name}，可在当前轮直接进行多模态问答。"
        emit_progress(progress_callback, 100, "done", f"图片 {asset.file_name} 已上传", file_name=asset.file_name)
        return UploadProcessingResult(asset=asset, upload_kind="image", summary=summary)

    def _process_document(
        self,
        *,
        file_name: str,
        content: bytes,
        conversation_id: str | None,
        progress_callback: ProgressCallback | None = None,
    ) -> UploadProcessingResult:
        emit_progress(progress_callback, 5, "saving", f"正在保存文档 {Path(file_name).name}", file_name=file_name)
        asset = self.asset_manager.save_bytes(
            asset_id=new_id("asset"),
            file_name=file_name,
            content=content,
            upload_kind="document",
            conversation_id=conversation_id,
            metadata={},
        )
        suffix = asset.suffix.lower()
        extracted: list[tuple[str, str | None, str | None]] = []
        ocr_pending = False
        parse_method = "native"
        text_char_count = 0

        if suffix in {".txt", ".md"}:
            emit_progress(progress_callback, 25, "native_extract", f"正在原生解析 {asset.file_name}", file_name=asset.file_name)
            text = decode_text_content(content)
            markdown = normalize_markdown_text(text)
            extracted.append((markdown, None, None))
            text_char_count = count_visible_chars(markdown)
            parse_method = "native_markdown" if suffix == ".md" else "native_text"
            emit_progress(progress_callback, 65, "native_extract_done", f"原生解析完成，提取 {text_char_count} 个字符", file_name=asset.file_name)
        elif suffix == ".pdf":
            extracted, parse_method, text_char_count = self._parse_pdf(
                asset=asset,
                content=content,
                progress_callback=progress_callback,
            )
        elif suffix == ".docx":
            emit_progress(progress_callback, 25, "native_extract", f"正在原生解析 DOCX {asset.file_name}", file_name=asset.file_name)
            document = Document(io.BytesIO(content))
            markdown = docx_to_markdown(document)
            extracted.append((markdown, None, None))
            text_char_count = count_visible_chars(markdown)
            parse_method = "native_docx"
            emit_progress(progress_callback, 65, "native_extract_done", f"DOCX 原生解析完成，提取 {text_char_count} 个字符", file_name=asset.file_name)
        elif suffix == ".doc":
            raise UploadProcessingError("暂不支持 .doc 解析，请先转换为 .docx 或 PDF。")

        if not any(text.strip() for text, _, _ in extracted):
            ocr_pending = True
            emit_progress(progress_callback, 68, "ocr_fallback", f"原生解析未提取到文本，尝试 OCR 预留接口", file_name=asset.file_name)
            ocr_text = self.ocr_provider.extract_text(asset=asset, content=content)
            if ocr_text.strip():
                markdown = normalize_markdown_text(ocr_text)
                extracted = [(markdown, None, None)]
                text_char_count = count_visible_chars(markdown)
                parse_method = "ocr_provider"
                ocr_pending = False

        emit_progress(progress_callback, 93, "chunking", f"正在写入 Markdown chunk", file_name=asset.file_name)
        chunks: list[ParsedChunk] = []
        chunk_index = 0
        for text, page, sheet in extracted:
            for chunk_text in split_text_chunks(text):
                chunks.append(ParsedChunk(text=chunk_text, chunk_index=chunk_index, page=page, sheet=sheet))
                chunk_index += 1

        asset = self.asset_manager.update_metadata(
            asset.asset_id,
            {
                "parse_method": parse_method,
                "text_char_count": text_char_count,
                "text_threshold": self.text_threshold,
                "markdown": True,
                "chunk_count": len(chunks),
            },
        )
        summary = build_summary(asset.file_name, [item.text for item in chunks], fallback="文档已上传，但暂未提取到可检索文本。")
        emit_progress(progress_callback, 100, "done", f"{asset.file_name} 解析完成，生成 {len(chunks)} 个 chunk", file_name=asset.file_name)
        return UploadProcessingResult(
            asset=asset,
            upload_kind="document",
            summary=summary,
            ocr_pending=ocr_pending,
            chunks=chunks,
        )

    def _parse_pdf(
        self,
        *,
        asset: AssetRecord,
        content: bytes,
        progress_callback: ProgressCallback | None = None,
    ) -> tuple[list[tuple[str, str | None, str | None]], str, int]:
        emit_progress(progress_callback, 18, "pdf_native_extract", f"正在原生抽取 PDF 文本 {asset.file_name}", file_name=asset.file_name)
        native_pages = extract_pdf_native_pages(content=content, progress_callback=progress_callback, file_name=asset.file_name)
        native_char_count = sum(count_visible_chars(text) for text, _, _ in native_pages)
        emit_progress(
            progress_callback,
            42,
            "pdf_native_extract_done",
            f"PDF 原生抽取完成，共 {native_char_count} 个字符，阈值 {self.text_threshold}",
            file_name=asset.file_name,
            text_char_count=native_char_count,
            text_threshold=self.text_threshold,
        )
        if native_char_count >= self.text_threshold:
            emit_progress(progress_callback, 70, "pdf_native_selected", "文本量达到阈值，继续使用原生解析结果", file_name=asset.file_name)
            return native_pages, "native_pdf", native_char_count

        emit_progress(progress_callback, 48, "pdf_vlm_selected", "文本量低于阈值，判定为扫描/图片 PDF，准备渲染页面图片", file_name=asset.file_name)
        rendered_pages = render_pdf_pages(
            content=content,
            dpi=self.pdf_render_dpi,
            max_pages=self.pdf_vlm_max_pages,
            progress_callback=progress_callback,
            file_name=asset.file_name,
            start_percent=50,
            end_percent=55,
        )
        extracted = self.vlm_provider.extract_pdf_pages_as_markdown(
            asset=asset,
            pages=rendered_pages,
            progress_callback=progress_callback,
            start_percent=55,
            end_percent=92,
        )
        vlm_char_count = sum(count_visible_chars(text) for text, _, _ in extracted)
        if not extracted or vlm_char_count <= 0:
            raise UploadProcessingError("扫描 PDF 已渲染为图片，但 VLM 未返回可索引文本；请确认多模态模型配置可用。")
        return extracted, "vlm_pdf", vlm_char_count

    def _process_signal(
        self,
        *,
        file_name: str,
        content: bytes,
        conversation_id: str | None,
        progress_callback: ProgressCallback | None = None,
    ) -> UploadProcessingResult:
        emit_progress(progress_callback, 10, "saving", f"正在保存 BIN 文件 {Path(file_name).name}", file_name=file_name)
        asset = self.asset_manager.save_bytes(
            asset_id=new_id("asset"),
            file_name=file_name,
            content=content,
            upload_kind="signal_bin",
            conversation_id=conversation_id,
            metadata={},
        )
        emit_progress(progress_callback, 45, "signal_preview", "正在生成频谱预览", file_name=asset.file_name)
        preview_bytes, sample_count = build_signal_preview(content)
        preview_asset = self.asset_manager.save_bytes(
            asset_id=new_id("asset"),
            file_name=f"{Path(file_name).stem}_spectrogram.png",
            content=preview_bytes,
            upload_kind="signal_preview",
            conversation_id=conversation_id,
            metadata={"source_asset_id": asset.asset_id},
        )
        asset = self.asset_manager.update_metadata(asset.asset_id, {"preview_asset_id": preview_asset.asset_id})
        summary = (
            f"已上传二进制信号文件 {asset.file_name}。"
            f" 当前仅生成频谱预览，采样点数约 {sample_count}。"
            f" 预览图可通过 {self.asset_manager.public_url(preview_asset.asset_id)} 访问。"
        )
        chunks = [
            ParsedChunk(
                text=summary + f" 可在回答中直接引用 Markdown：![{preview_asset.file_name}]({self.asset_manager.public_url(preview_asset.asset_id)})",
                chunk_index=0,
            )
        ]
        emit_progress(progress_callback, 100, "done", f"BIN 文件 {asset.file_name} 处理完成", file_name=asset.file_name)
        return UploadProcessingResult(
            asset=asset,
            upload_kind="signal_bin",
            summary=summary,
            preview_asset=preview_asset,
            chunks=chunks,
        )


def emit_progress(
    progress_callback: ProgressCallback | None,
    percent: int | float,
    stage: str,
    message: str,
    **extra: Any,
) -> None:
    if progress_callback is None:
        return
    try:
        progress_callback(
            {
                "percent": max(0, min(100, int(round(float(percent))))),
                "stage": stage,
                "message": message,
                **extra,
            }
        )
    except Exception:
        # Progress reporting must never break parsing/indexing.
        pass


def decode_text_content(content: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "gb18030", "latin-1"):
        try:
            return content.decode(encoding)
        except Exception:
            continue
    return content.decode("utf-8", errors="ignore")


def normalize_markdown_text(text: str) -> str:
    normalized = str(text or "").replace("\r\n", "\n").replace("\r", "\n")
    lines = [line.rstrip() for line in normalized.split("\n")]
    return collapse_blank_lines("\n".join(lines)).strip()


def normalize_vlm_markdown(text: str) -> str:
    markdown = normalize_markdown_text(text)
    if markdown.startswith("```markdown"):
        markdown = markdown[len("```markdown") :].strip()
    elif markdown.startswith("```"):
        markdown = markdown[3:].strip()
    if markdown.endswith("```"):
        markdown = markdown[:-3].strip()
    return normalize_markdown_text(markdown)


def collapse_blank_lines(text: str, *, max_blank_lines: int = 2) -> str:
    result: list[str] = []
    blank_count = 0
    for line in text.split("\n"):
        if line.strip():
            blank_count = 0
            result.append(line)
            continue
        blank_count += 1
        if blank_count <= max_blank_lines:
            result.append("")
    return "\n".join(result)


def count_visible_chars(text: str) -> int:
    return sum(1 for char in str(text or "") if not char.isspace())


def extract_pdf_native_pages(
    *,
    content: bytes,
    progress_callback: ProgressCallback | None = None,
    file_name: str = "",
) -> list[tuple[str, str | None, str | None]]:
    reader = PdfReader(io.BytesIO(content))
    total_pages = len(reader.pages)
    extracted: list[tuple[str, str | None, str | None]] = []
    if total_pages <= 0:
        return extracted
    for page_index, page in enumerate(reader.pages, start=1):
        text = normalize_markdown_text(page.extract_text() or "")
        markdown = f"## 第 {page_index} 页\n\n{text}" if total_pages > 1 and text else text
        extracted.append((markdown, str(page_index), None))
        emit_progress(
            progress_callback,
            18 + int(20 * page_index / total_pages),
            "pdf_native_page",
            f"原生抽取 PDF 第 {page_index}/{total_pages} 页",
            file_name=file_name,
            page=page_index,
        )
    return extracted


def render_pdf_pages(
    *,
    content: bytes,
    dpi: int,
    max_pages: int = 0,
    progress_callback: ProgressCallback | None = None,
    file_name: str = "",
    start_percent: int = 50,
    end_percent: int = 55,
) -> list[RenderedPDFPage]:
    if fitz is None:
        raise UploadProcessingError("扫描 PDF 渲染需要安装 PyMuPDF：pip install PyMuPDF")
    doc = fitz.open(stream=content, filetype="pdf")
    try:
        page_count = int(getattr(doc, "page_count", 0) or len(doc))
        if page_count <= 0:
            raise UploadProcessingError("PDF 未包含可渲染页面。")
        render_count = min(page_count, max_pages) if max_pages and max_pages > 0 else page_count
        zoom = max(36, int(dpi)) / 72.0
        matrix = fitz.Matrix(zoom, zoom)
        pages: list[RenderedPDFPage] = []
        span = max(1, end_percent - start_percent)
        for zero_index in range(render_count):
            page_number = zero_index + 1
            page = doc.load_page(zero_index)
            pix = page.get_pixmap(matrix=matrix, alpha=False)
            pages.append(RenderedPDFPage(page_number=page_number, mime_type="image/png", content=pix.tobytes("png")))
            emit_progress(
                progress_callback,
                start_percent + int(span * page_number / render_count),
                "pdf_render_page",
                f"已将 PDF 第 {page_number}/{render_count} 页渲染为图片",
                file_name=file_name,
                page=page_number,
            )
        return pages
    finally:
        doc.close()


def docx_to_markdown(document: Document) -> str:
    parts: list[str] = []
    for paragraph in document.paragraphs:
        text = normalize_markdown_text(paragraph.text)
        if text:
            parts.append(text)
    for table in document.tables:
        rows = [[normalize_markdown_text(cell.text).replace("\n", "<br>") for cell in row.cells] for row in table.rows]
        table_md = markdown_table(rows)
        if table_md:
            parts.append(table_md)
    return normalize_markdown_text("\n\n".join(parts))


def markdown_table(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    width = max((len(row) for row in rows), default=0)
    if width <= 0:
        return ""
    normalized = [(row + [""] * width)[:width] for row in rows]
    header = normalized[0]
    body = normalized[1:]
    lines = ["| " + " | ".join(escape_table_cell(cell) for cell in header) + " |"]
    lines.append("| " + " | ".join("---" for _ in header) + " |")
    for row in body:
        lines.append("| " + " | ".join(escape_table_cell(cell) for cell in row) + " |")
    return "\n".join(lines)


def escape_table_cell(value: str) -> str:
    return str(value or "").replace("|", "\\|")


def split_text_chunks(text: str, *, chunk_size: int = 1200, overlap: int = 120) -> list[str]:
    normalized = normalize_markdown_text(text)
    if not normalized:
        return []
    chunks: list[str] = []
    start = 0
    while start < len(normalized):
        end = min(len(normalized), start + chunk_size)
        if end < len(normalized):
            paragraph_break = normalized.rfind("\n\n", start + max(200, chunk_size // 3), end)
            if paragraph_break > start:
                end = paragraph_break
        chunks.append(normalized[start:end].strip())
        if end >= len(normalized):
            break
        start = max(end - overlap, start + 1)
    return [chunk for chunk in chunks if chunk]


def build_summary(file_name: str, chunks: list[str], *, fallback: str) -> str:
    if not chunks:
        return fallback
    preview = " ".join(chunks[0][:240].split())
    return f"{file_name} 已解析为 Markdown，内容摘要：{preview}"


def build_signal_preview(content: bytes) -> tuple[bytes, int]:
    if len(content) >= 2 and len(content) % 2 == 0:
        samples = np.frombuffer(content, dtype=np.int16).astype(np.float32)
    else:
        samples = np.frombuffer(content, dtype=np.uint8).astype(np.float32)
    if samples.size == 0:
        samples = np.zeros(256, dtype=np.float32)
    fig = Figure(figsize=(8, 3))
    ax = fig.subplots()
    ax.specgram(samples, NFFT=256, Fs=1.0, noverlap=128, cmap="viridis")
    ax.set_title("Binary Signal Spectrogram")
    ax.set_xlabel("Time")
    ax.set_ylabel("Frequency")
    buffer = io.BytesIO()
    fig.savefig(buffer, format="png", bbox_inches="tight")
    return buffer.getvalue(), int(samples.size)


def encode_image_as_data_url(mime_type: str, content: bytes) -> str:
    encoded = base64.b64encode(content).decode("ascii")
    return f"data:{mime_type};base64,{encoded}"
