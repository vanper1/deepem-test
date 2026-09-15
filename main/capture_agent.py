from __future__ import annotations

import copy
import hashlib
import json
import math
import os
import re
import shutil
import threading
import time
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable
from uuid import uuid4

from docx import Document

from deepem.agent.profiles import CAPTURE_AGENT
from deepem.nl2sql_config import NL2SQLSessionConfig
from deepem.protocol import (
    ChatMessage,
    ChatRole,
    Run,
    RunStatus,
    RunTriggerKind,
    new_id,
    utc_now,
)
from deepem.tools import ToolRegistry, register_builtin_tools
from deepem.tools.autonomous_usrp import AUTONOMOUS_OUTPUT_ROOT
from deepem.tools.base import ToolContext
from deepem.upload_processing import docx_to_markdown
from .probe_evidence_compaction import build_probe_evidence_pack
from .probe_evidence_store import get_probe_evidence_store
from .device_evidence_chain import build_device_evidence_chain
from .capture_planner import (
    CaptureLLMPlanner,
    CapturePlanningUnavailable,
    CapturePlanValidationError,
)


CAPTURE_AGENT_ROOT = Path(__file__).resolve().parent / "data" / "capture_agent"
TASKS_DIR = CAPTURE_AGENT_ROOT / "tasks"
TEMPLATES_DIR = CAPTURE_AGENT_ROOT / "templates"
REPORTS_DIR = CAPTURE_AGENT_ROOT / "reports"

TERMINAL_TASK_STATES = {"completed", "failed", "cancelled"}
TERMINAL_NODE_STATES = {"completed", "failed", "cancelled", "skipped"}

_MISSING = object()


def utc_now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def _safe_name(value: str, fallback: str = "capture-template.docx") -> str:
    clean = Path(value or fallback).name
    clean = re.sub(r"[^0-9A-Za-z_.\-\u4e00-\u9fff]+", "_", clean).strip("._")
    return clean or fallback


def _deepcopy_json(value: Any) -> Any:
    return json.loads(json.dumps(value, ensure_ascii=False, default=str))


def _normalize_reasoning_mode(value: Any) -> str:
    normalized = str(value or "").strip().lower()
    if normalized == "deep":
        return "deep"
    return "fast"


def _initial_probe_assistance(selected_probes: list[dict[str, Any]]) -> dict[str, Any]:
    enabled = bool(selected_probes)
    return {
        "enabled": enabled,
        "status": "waiting_for_usrp" if enabled else "not_selected",
        "status_label": "等待 USRP 采集完成" if enabled else "未选择探针",
        "selected_probes": _deepcopy_json(selected_probes),
        "confirmed_at": None,
        "started_at": None,
        "ended_at": None,
        "summary": "",
        "summary_stream": "",
        "reasoning": "",
        "details": {},
        "errors": [],
        "steps": [
            {"id": "confirm", "title": "确认调用 WiFi/蓝牙探针", "status": "pending", "message": "等待 USRP 采集完成后由用户确认"},
            {"id": "devices", "title": "查询探针列表与在线状态", "status": "pending", "message": "等待执行"},
            {"id": "targets", "title": "查询 WiFi/蓝牙目标汇总", "status": "pending", "message": "等待执行"},
            {"id": "observations", "title": "查询目标观测时段明细", "status": "pending", "message": "等待执行"},
            {"id": "summary", "title": "智能总结辅助证据", "status": "pending", "message": "等待执行"},
        ],
    }


def _freq_value(number: str, unit: str | None) -> float:
    """Return a frequency value normalized to MHz for all platform/planner fields."""
    value = float(number)
    unit_text = (unit or "mhz").lower()
    if unit_text in {"g", "ghz"}:
        return value * 1e3
    if unit_text in {"m", "mhz"}:
        return value
    if unit_text in {"k", "khz"}:
        return value / 1e3
    if unit_text in {"h", "hz"}:
        return value / 1e6
    return value


def _human_frequency(value: float | int | None) -> str:
    if value is None:
        return "-"
    number = float(value)
    return f"{number:g} MHz"


def _device_mhz_value(value: Any) -> float:
    number = float(value)
    return number / 1e6 if abs(number) > 100000 else number


class CapturePlanBuilder:
    """Explicit-parameter parsing and safety helpers around the real-LLM planner.

    This helper never supplies collection defaults. It extracts only values that
    are actually stated by the operator or a template, verifies the complete
    parameter set designed by the LLM, and enforces hard platform limits.
    """

    REQUIRED_CONSTRAINT_KEYS: tuple[str, ...] = (
        "room_name",
        "mode",
        "freq_start_mhz",
        "freq_stop_mhz",
        "freq_step_mhz",
        "dwell_time_sec",
        "repeat_count",
        "aggregation",
        "sample_rate",
        "bandwidth",
        "gain",
        "antenna",
        "expected_fft_frames_per_capture",
    )

    @classmethod
    def extract_constraints(cls, instruction: str, template_text: str = "") -> dict[str, Any]:
        source = f"{template_text}\n{instruction}".strip()
        result: dict[str, Any] = {}

        room_match = re.search(
            r"([\u4e00-\u9fffA-Za-z0-9_-]{1,32})\s*(会议室|房间|实验室|教室|机房|区域|room)",
            source,
            flags=re.I,
        )
        if room_match:
            prefix = re.sub(r"^(请|帮我|采集|检测|检查|扫描|对|在)", "", room_match.group(1)).strip()
            result["room_name"] = (prefix + room_match.group(2)).strip() or room_match.group(0).replace(" ", "")

        if any(keyword in source.lower() for keyword in ["背景", "background", "底噪", "空场"]):
            result["mode"] = "background"
        elif any(keyword in source.lower() for keyword in ["正常", "normal", "基线"]):
            result["mode"] = "normal"

        range_patterns = [
            r"([0-9]+(?:\.[0-9]+)?)\s*(GHz|G|MHz|M|kHz|K|Hz)\s*(?:-|~|到|至|—|–)\s*([0-9]+(?:\.[0-9]+)?)\s*(GHz|G|MHz|M|kHz|K|Hz)",
            r"(?:起始频率|start(?:_frequency)?)\s*[:：=]?\s*([0-9]+(?:\.[0-9]+)?)\s*(GHz|G|MHz|M|kHz|K|Hz).*?(?:终止频率|结束频率|stop(?:_frequency)?)\s*[:：=]?\s*([0-9]+(?:\.[0-9]+)?)\s*(GHz|G|MHz|M|kHz|K|Hz)",
        ]
        for pattern in range_patterns:
            match = re.search(pattern, source, flags=re.I | re.S)
            if match:
                result["freq_start_mhz"] = _freq_value(match.group(1), match.group(2))
                result["freq_stop_mhz"] = _freq_value(match.group(3), match.group(4))
                break

        single_freq = re.search(
            r"(?:中心频率|频点|frequency|freq)\s*[:：=]?\s*([0-9]+(?:\.[0-9]+)?)\s*(GHz|G|MHz|M|kHz|K|Hz)",
            source,
            flags=re.I,
        )
        if single_freq and not any(re.search(pattern, source, flags=re.I | re.S) for pattern in range_patterns):
            value = _freq_value(single_freq.group(1), single_freq.group(2))
            result["freq_start_mhz"] = value
            result["freq_stop_mhz"] = value

        mappings = [
            ("freq_step_mhz", r"(?:步长|步进|step)[^0-9]{0,16}([0-9]+(?:\.[0-9]+)?)\s*(GHz|G|MHz|M|kHz|K|Hz)", "freq"),
            ("sample_rate", r"(?:采样率|sample[_ ]?rate)[^0-9]{0,16}([0-9]+(?:\.[0-9]+)?)\s*(GHz|G|MHz|M|kHz|K|Hz)", "freq"),
            ("bandwidth", r"(?:带宽|bandwidth)[^0-9]{0,16}([0-9]+(?:\.[0-9]+)?)\s*(GHz|G|MHz|M|kHz|K|Hz)", "freq"),
            ("gain", r"(?:增益|gain)[^0-9]{0,16}([0-9]+(?:\.[0-9]+)?)\s*(?:dB)?", "number"),
        ]
        for key, pattern, kind in mappings:
            match = re.search(pattern, source, flags=re.I)
            if match:
                result[key] = _freq_value(match.group(1), match.group(2)) if kind == "freq" else float(match.group(1))

        dwell_match = re.search(
            r"(?:扫描时间|采集时间|驻留|dwell|duration)[^0-9]{0,16}([0-9]+(?:\.[0-9]+)?)\s*(ms|毫秒|s|秒)",
            source,
            flags=re.I,
        )
        if dwell_match:
            result["dwell_time_sec"] = float(dwell_match.group(1)) / 1000 if dwell_match.group(2).lower() in {"ms", "毫秒"} else float(dwell_match.group(1))

        repeat_match = re.search(r"(?:扫描次数|采集次数|重复|repeat)[^0-9]{0,16}([0-9]+)\s*(?:次)?", source, flags=re.I)
        if repeat_match:
            result["repeat_count"] = int(repeat_match.group(1))

        aggregation_match = re.search(
            r"(?:聚合方式|聚合|aggregation)\s*[:：=]?\s*(mean|median|均值|平均值|中位数)",
            source,
            flags=re.I,
        )
        if aggregation_match:
            aggregation_text = aggregation_match.group(1).lower()
            result["aggregation"] = "median" if aggregation_text in {"median", "中位数"} else "mean"

        fft_frames_match = re.search(
            r"(?:FFT\s*帧数|期望FFT帧数|expected[_ ]?fft[_ ]?frames(?:[_ ]?per[_ ]?capture)?)[^0-9]{0,16}([0-9]+)",
            source,
            flags=re.I,
        )
        if fft_frames_match:
            result["expected_fft_frames_per_capture"] = int(fft_frames_match.group(1))

        antenna_match = re.search(r"(?:天线|antenna)\s*[:：=]?\s*(TX/RX|RX2)", source, flags=re.I)
        if antenna_match:
            result["antenna"] = antenna_match.group(1).upper()

        for key in (
            "freq_start_mhz",
            "freq_stop_mhz",
            "freq_step_mhz",
            "dwell_time_sec",
            "sample_rate",
            "bandwidth",
            "gain",
        ):
            if key in result:
                result[key] = float(result[key])
        for key in ("repeat_count", "expected_fft_frames_per_capture"):
            if key in result:
                result[key] = int(result[key])
        return result

    @classmethod
    def normalize_constraints(cls, constraints: dict[str, Any]) -> dict[str, Any]:
        missing = [
            key
            for key in cls.REQUIRED_CONSTRAINT_KEYS
            if key not in constraints or constraints.get(key) is None or constraints.get(key) == ""
        ]
        if missing:
            raise ValueError(
                "采集参数不完整，禁止使用默认值补齐；请让意图分析器根据用户意图设计缺失字段："
                + "、".join(missing)
            )
        result = {key: constraints[key] for key in cls.REQUIRED_CONSTRAINT_KEYS}
        result["room_name"] = str(result["room_name"]).strip()
        result["mode"] = str(result["mode"]).strip()
        result["aggregation"] = str(result["aggregation"]).strip().lower()
        result["antenna"] = str(result["antenna"]).strip().upper()
        for key in (
            "freq_start_mhz",
            "freq_stop_mhz",
            "freq_step_mhz",
            "dwell_time_sec",
            "sample_rate",
            "bandwidth",
            "gain",
        ):
            result[key] = float(result[key])
        for key in ("repeat_count", "expected_fft_frames_per_capture"):
            result[key] = int(result[key])
        if not result["room_name"] or not result["mode"] or not result["antenna"]:
            raise ValueError("room_name、mode、antenna 必须由用户、模板或意图设计明确给出")
        if result["aggregation"] not in {"mean", "median"}:
            raise ValueError("aggregation 必须为 mean 或 median")
        return result

    @classmethod
    def validate_constraints(cls, constraints: dict[str, Any]) -> dict[str, Any]:
        normalized = cls.normalize_constraints(constraints)
        start = float(normalized["freq_start_mhz"])
        stop = float(normalized["freq_stop_mhz"])
        step = float(normalized["freq_step_mhz"])
        repeats = int(normalized["repeat_count"])
        dwell = float(normalized["dwell_time_sec"])
        if start <= 0 or stop <= 0 or stop < start:
            raise ValueError("频率范围无效：终止频率必须大于等于起始频率")
        if step <= 0:
            raise ValueError("频率步长必须大于 0")
        if repeats < 1 or repeats > 50:
            raise ValueError("重复次数必须在 1 到 50 之间")
        if dwell <= 0 or dwell > 60:
            raise ValueError("单频点驻留时间必须大于 0 且不超过 60 秒")
        count = int(math.floor((stop - start) / step)) + 1
        if count > 800000:
            raise ValueError(f"频点数量为 {count}，超过安全上限 800000；请增大步长或缩小频率范围")
        if float(normalized["bandwidth"]) > float(normalized["sample_rate"]):
            raise ValueError("带宽不能大于采样率")
        return {
            "frequency_count": count,
            "estimated_capture_seconds": round(count * repeats * dwell, 3),
            "range_text": f"{_human_frequency(start)} - {_human_frequency(stop)}",
            "step_text": _human_frequency(step),
        }

    @classmethod
    def build_bootstrap_nodes(cls) -> list[dict[str, Any]]:
        specs = [
            ("template_analysis", "解析采集模板", [], "解析 DOCX 的标题、段落和表格，为设备判断和后续意图解析构建任务背景。", None, "template_analysis"),
            ("device_discovery", "感知当前可用设备", ["template_analysis"], "自动探测 USRP 与 WiFi/蓝牙探针，把在线状态、能力与设备标识提供给规划器。", None, "device_discovery"),
            ("device_selection", "LLM 自主选择采集设备", ["template_analysis", "device_discovery"], "真实 LLM 根据任务背景和可用设备清单自主决定调用 USRP、WiFi/蓝牙探针或两者组合，不使用固定设备顺序。", None, "llm_device_selection"),
            ("instruction_analysis", "LLM 理解采集意图", ["device_selection"], "真实 LLM 融合模板、用户指令和设备选择；仅在选择 USRP 时解析完整频谱采集参数。", None, "llm_intent_analysis"),
            ("plan_generation", "生成自适应多设备计划", ["instruction_analysis"], "按 LLM 的设备选择编排现有 USRP 内部节点和 WiFi/蓝牙探针节点，设备组顺序由任务决定。", None, "adaptive_device_plan_generation"),
            ("approval", "等待用户批准", ["plan_generation"], "用户批准前禁止执行真实采集；批准后锁定设备选择、参数与计划指纹。", None, "approval_gate"),
        ]
        now = utc_now_iso()
        return [
            {
                "id": node_id,
                "title": title,
                "description": description,
                "executor": executor,
                "dependencies": dependencies,
                "allowed_tool": tool,
                "allowed_tools": [tool] if tool else [],
                "status": "pending",
                "progress": 0,
                "attempts": 0,
                "started_at": None,
                "ended_at": None,
                "updated_at": now,
                "summary": "等待执行",
                "success_criteria": cls.success_criteria(node_id),
                "expected_evidence": [],
                "tool_input_overrides": {},
                "risk_level": "low",
                "planner_origin": "system_bootstrap",
                "inputs": {},
                "outputs": {},
                "logs": [],
                "notes": "",
            }
            for node_id, title, dependencies, description, tool, executor in specs
        ]

    @classmethod
    def build_nodes(cls) -> list[dict[str, Any]]:
        """Backward-compatible alias for bootstrap planning nodes."""
        return cls.build_bootstrap_nodes()

    @staticmethod
    def success_criteria(node_id: str) -> list[str]:
        criteria = {
            "template_analysis": ["全部选定 DOCX 已解析，或明确记录未上传模板"],
            "device_discovery": ["已探测 USRP 与 WiFi/蓝牙探针", "可用设备清单包含在线状态和设备标识"],
            "device_selection": ["真实 LLM 已基于任务背景与可用设备选择设备类型", "设备选择理由可审计且不强制 USRP 优先"],
            "instruction_analysis": ["真实 LLM 返回符合任务所需的结构化意图", "选择 USRP 时参数来源、假设和未决问题可追踪"],
            "plan_generation": ["仅编排 LLM 选择的可用设备", "USRP 内部节点保持原流程", "探针节点使用落库+按需检索机制"],
            "approval": ["用户显式批准当前计划版本", "批准时锁定计划指纹"],
            "device_scan": ["已查询可用设备列表", "已选择空闲 USRP 设备或给出明确失败原因"],
            "collection_execution": ["execute_spectrum_collection 返回成功", "生成标准化 .npz 输出文件"],
            "output_file_check": ["输出文件存在", "输出文件扩展名为 .npz", "输出文件大小大于 0"],
            "execution_summary": ["形成采集任务执行摘要"],
            "report_generation": ["已生成并检查 Markdown 研判报告", "报告已融合本任务中各类设备的采集情况与证据摘要"],
        }
        return criteria.get(node_id, [])


class CaptureAgentService:
    def __init__(self, platform: Any, root: Path | None = None, llm_client: Any | None = None) -> None:
        self.platform = platform
        runtime_llm = llm_client
        if runtime_llm is None:
            try:
                runtime_llm = platform.app.runtime.llm_client
            except Exception:
                runtime_llm = None
        self.planner = CaptureLLMPlanner(runtime_llm)
        self._tool_definitions = self._resolve_capture_tool_definitions()
        self.tool_catalog = {
            name: {
                "name": definition.name,
                "description": definition.description,
                "input_schema": definition.input_schema,
            }
            for name, definition in self._tool_definitions.items()
        }
        self.root = Path(root or CAPTURE_AGENT_ROOT)
        self.tasks_dir = self.root / "tasks"
        self.templates_dir = self.root / "templates"
        self.reports_dir = self.root / "reports"
        for directory in (self.root, self.tasks_dir, self.templates_dir, self.reports_dir):
            directory.mkdir(parents=True, exist_ok=True)
        self._lock = threading.RLock()
        self._workers: dict[str, threading.Thread] = {}
        self._restart_pending: set[str] = set()
        self._pause_events: dict[str, threading.Event] = {}
        self._cancel_events: dict[str, threading.Event] = {}
        self._summary_lock = threading.RLock()
        self._summary_workers: dict[tuple[str, str], threading.Thread] = {}
        self._summary_pending: dict[tuple[str, str], dict[str, Any]] = {}
        self._summary_signatures: dict[tuple[str, str], str] = {}
        self._summary_last_requested: dict[tuple[str, str], float] = {}
        self._probe_workers: dict[str, threading.Thread] = {}
        self._recover_interrupted_tasks()

    def runtime_info(self) -> dict[str, Any]:
        info = self.planner.runtime_info()
        return {**info, "tool_count": len(self.tool_catalog), "planning_mode": "adaptive_device_selection_with_stable_device_nodes"}

    def _resolve_capture_tool_definitions(self) -> dict[str, Any]:
        registry = None
        try:
            registry = self.platform.app.runtime.tool_registry
        except Exception:
            registry = register_builtin_tools(ToolRegistry())
        definitions: dict[str, Any] = {}
        for name in CAPTURE_AGENT.allowed_tools:
            try:
                definitions[name] = registry.get(name)
            except Exception:
                continue
        if not definitions:
            fallback = register_builtin_tools(ToolRegistry())
            for name in CAPTURE_AGENT.allowed_tools:
                try:
                    definitions[name] = fallback.get(name)
                except Exception:
                    continue
        return definitions

    # ---------- template library ----------
    def list_templates(self) -> list[dict[str, Any]]:
        items = []
        with self._lock:
            for path in sorted(self.templates_dir.glob("*.json"), key=lambda item: item.stat().st_mtime, reverse=True):
                try:
                    payload = json.loads(path.read_text(encoding="utf-8"))
                except Exception:
                    continue
                items.append(self._public_template(payload))
        return items

    def upload_template(self, filename: str, content: bytes, metadata: dict[str, Any] | None = None) -> dict[str, Any]:
        safe_filename = _safe_name(filename)
        if Path(safe_filename).suffix.lower() != ".docx":
            raise ValueError("采集模板初版仅支持 .docx 文件")
        if not content:
            raise ValueError("上传文件为空")
        template_id = f"tpl_{uuid4().hex[:12]}"
        template_dir = self.templates_dir / template_id
        template_dir.mkdir(parents=True, exist_ok=False)
        original_path = template_dir / safe_filename
        original_path.write_bytes(content)
        try:
            document = Document(str(original_path))
            markdown = docx_to_markdown(document)
            headings = []
            tables = []
            for paragraph in document.paragraphs:
                text = paragraph.text.strip()
                if text and str(paragraph.style.name or "").lower().startswith("heading"):
                    headings.append(text)
            for table_index, table in enumerate(document.tables, start=1):
                rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
                tables.append({"index": table_index, "rows": rows})
        except Exception as exc:
            shutil.rmtree(template_dir, ignore_errors=True)
            raise ValueError(f"DOCX 解析失败：{exc}") from exc
        now = utc_now_iso()
        metadata = dict(metadata or {})
        display_name = str(metadata.get("name") or Path(safe_filename).stem).strip() or Path(safe_filename).stem
        payload = {
            "id": template_id,
            "name": display_name,
            "file_name": safe_filename,
            "file_path": str(original_path),
            "size_bytes": len(content),
            "markdown": markdown,
            "text_preview": markdown[:1200],
            "headings": headings[:100],
            "tables": tables[:50],
            "scene": str(metadata.get("scene") or "通用检测场景").strip() or "通用检测场景",
            "operator": str(metadata.get("operator") or "operator").strip() or "operator",
            "created_at": now,
            "updated_at": now,
        }
        self._write_json(self.templates_dir / f"{template_id}.json", payload)
        return self._public_template(payload)

    def get_template(self, template_id: str) -> dict[str, Any]:
        path = self.templates_dir / f"{template_id}.json"
        if not path.exists():
            raise KeyError(template_id)
        return json.loads(path.read_text(encoding="utf-8"))

    def get_template_detail(self, template_id: str) -> dict[str, Any]:
        with self._lock:
            payload = self.get_template(template_id)
            public = self._public_template(payload)
            return {
                **public,
                "markdown": payload.get("markdown", ""),
                "text": payload.get("markdown", ""),
                "tables": payload.get("tables", []),
            }

    def rename_template(self, template_id: str, name: str) -> dict[str, Any]:
        display_name = str(name or "").strip()
        if not display_name:
            raise ValueError("模板名称不能为空")
        safe_filename = _safe_name(display_name if display_name.lower().endswith(".docx") else f"{display_name}.docx")
        with self._lock:
            payload = self.get_template(template_id)
            old_path = Path(str(payload.get("file_path") or ""))
            new_path = old_path.with_name(safe_filename) if old_path.name else self.templates_dir / template_id / safe_filename
            if old_path.exists() and old_path.resolve() != new_path.resolve():
                if new_path.exists():
                    new_path = new_path.with_name(f"{new_path.stem}_{uuid4().hex[:6]}{new_path.suffix}")
                    safe_filename = new_path.name
                old_path.rename(new_path)
                payload["file_path"] = str(new_path)
            payload["name"] = Path(safe_filename).stem
            payload["file_name"] = safe_filename
            payload["updated_at"] = utc_now_iso()
            self._write_json(self.templates_dir / f"{template_id}.json", payload)
            return self._public_template(payload)

    def delete_template(self, template_id: str) -> None:
        with self._lock:
            payload = self.get_template(template_id)
            json_path = self.templates_dir / f"{template_id}.json"
            template_dir = Path(str(payload.get("file_path") or "")).parent
            if json_path.exists():
                json_path.unlink()
            if template_dir.exists() and template_dir.parent.resolve() == self.templates_dir.resolve():
                shutil.rmtree(template_dir, ignore_errors=True)

    # ---------- task lifecycle ----------
    def list_tasks(self, limit: int = 50) -> list[dict[str, Any]]:
        with self._lock:
            payloads = self._all_task_payloads_locked()
            items = []
            for index, payload in enumerate(payloads, start=1):
                summary = self._task_summary(payload)
                summary.update(self._task_display_fields(payload, index))
                items.append(summary)
                if len(items) >= limit:
                    break
        return items

    def get_device_dashboard(self, *, refresh: bool = False) -> dict[str, Any]:
        """Return live device status plus capture-agent invocation history for the middle workbench."""
        errors: list[dict[str, str]] = []
        try:
            usrp_payload = dict((self.platform.scan_devices() if refresh else self.platform.list_devices()) or {})
        except Exception as exc:
            usrp_payload = {"devices": []}
            errors.append({"device_type": "usrp", "message": str(exc)})
        try:
            adapter = self.platform.app.runtime.device_registry.find_by_capability("collect_wifi_bluetooth_probe_evidence")
            client = getattr(adapter, "client", None) if adapter is not None else None
            if client is None:
                raise RuntimeError("探针适配器不可用")
            probes = [dict(item or {}) for item in client.list_devices()]
        except Exception as exc:
            probes = []
            errors.append({"device_type": "wifi_bluetooth_probe", "message": str(exc)})

        with self._lock:
            task_payloads = self._all_task_payloads_locked()

        cards: list[dict[str, Any]] = []
        for raw in usrp_payload.get("devices") or []:
            item = dict(raw or {})
            dev_id = str(item.get("dev_id") or item.get("id") or "usrp")
            status = str(item.get("status") or "UNKNOWN").upper()
            try:
                stream_info = dict(self.platform.get_usrp_stream_info(dev_id) or {})
            except Exception:
                stream_info = {}
            cards.append({
                "device_type": "usrp",
                "device_id": dev_id,
                "display_name": str(item.get("name") or item.get("product") or f"USRP {dev_id}"),
                "status": status,
                "online": status not in {"OFFLINE", "ERROR", "UNAVAILABLE"},
                "available_for_new_task": status in {"IDLE", "ONLINE", "READY"},
                "details": item,
                "stream_info": stream_info,
                "history": self._device_invocation_history(task_payloads, "usrp", dev_id),
            })
        for raw in probes:
            item = dict(raw or {})
            probe_id = str(item.get("probe_id") or item.get("id") or "probe")
            status = str(item.get("status") or "ONLINE").upper()
            cards.append({
                "device_type": "wifi_bluetooth_probe",
                "device_id": probe_id,
                "display_name": str(item.get("name") or item.get("label") or f"WiFi/蓝牙探针 {probe_id}"),
                "status": status,
                "online": status not in {"OFFLINE", "ERROR", "UNAVAILABLE"},
                "available_for_new_task": status not in {"OFFLINE", "ERROR", "UNAVAILABLE"},
                "details": item,
                "history": self._device_invocation_history(task_payloads, "wifi_bluetooth_probe", probe_id),
            })
        cards.sort(key=lambda item: (0 if item.get("device_type") == "usrp" else 1, str(item.get("device_id") or "")))
        return {
            "queried_at": utc_now_iso(),
            "refresh": bool(refresh),
            "online_count": sum(1 for item in cards if item.get("online")),
            "devices": cards,
            "errors": errors + ([{"device_type": "usrp", "message": str(usrp_payload.get("scan_error") or usrp_payload.get("error"))}] if (usrp_payload.get("scan_error") or usrp_payload.get("error")) else []),
        }

    def _device_invocation_history(self, tasks: list[dict[str, Any]], device_type: str, device_id: str) -> list[dict[str, Any]]:
        history: list[dict[str, Any]] = []
        for task_number, task in enumerate(tasks, start=1):
            nodes = {str(node.get("id")): node for node in task.get("nodes") or []}
            target_node: dict[str, Any] | None = None
            if device_type == "usrp":
                scan = nodes.get("device_scan") or {}
                selected = dict((scan.get("outputs") or {}).get("selected_device") or {})
                selected_id = str(selected.get("dev_id") or "")
                if selected_id and selected_id != device_id:
                    continue
                target_node = nodes.get("collection_execution")
                if target_node is None or target_node.get("status") not in {"in_progress", "completed", "failed", "cancelled"}:
                    continue
                if not selected_id:
                    selected_candidates = [str(item.get("dev_id") or "") for item in task.get("selected_usrp_devices") or []]
                    if selected_candidates and device_id not in selected_candidates:
                        continue
            else:
                target_node = nodes.get("wifi_bluetooth_probe_collection")
                if target_node is None or target_node.get("status") not in {"in_progress", "completed", "failed", "cancelled"}:
                    continue
                args = dict(target_node.get("tool_arguments") or {})
                if device_id not in [str(item) for item in args.get("probe_ids") or []]:
                    continue
            public_artifacts = []
            for artifact in task.get("artifacts") or []:
                artifact_id = str(artifact.get("id") or "")
                item = {
                    "id": artifact_id,
                    "file_name": artifact.get("file_name"),
                    "kind": artifact.get("kind"),
                    "size_bytes": artifact.get("size_bytes"),
                    "download_url": f"/api/capture-agent/tasks/{task.get('id')}/artifacts/{artifact_id}" if artifact_id else "",
                }
                if str(artifact.get("file_name") or "").lower().endswith(".npz") and artifact_id:
                    item["spectrum_url"] = f"/api/capture-agent/tasks/{task.get('id')}/artifacts/{artifact_id}/spectrum"
                public_artifacts.append(item)
            outputs = dict((target_node or {}).get("outputs") or {})
            task_title = str(task.get("title") or task.get("instruction") or task.get("id") or "采集任务")
            history.append({
                "task_id": task.get("id"),
                "task_number": task_number,
                "task_title": task_title,
                "task_display_title": f"任务{task_number}：{task_title}",
                "instruction": task.get("instruction") or "",
                "task_status": task.get("status"),
                "node_id": (target_node or {}).get("id"),
                "node_title": (target_node or {}).get("title"),
                "node_status": (target_node or {}).get("status"),
                "started_at": (target_node or {}).get("started_at"),
                "ended_at": (target_node or {}).get("ended_at"),
                "summary": (target_node or {}).get("summary") or "",
                "intelligent_summary": outputs.get("intelligent_summary") or task.get("latest_probe_summary") or "",
                "evidence_set_id": outputs.get("evidence_set_id") or "",
                "artifacts": public_artifacts,
            })
        history.sort(key=lambda item: str(item.get("ended_at") or item.get("started_at") or ""), reverse=True)
        return history[:30]

    def get_probe_live_snapshot(self, probe_id: str, *, page_size: int = 40) -> dict[str, Any]:
        probe_id = str(probe_id or "").strip()
        if not probe_id:
            raise ValueError("probe_id 不能为空")
        adapter = self.platform.app.runtime.device_registry.find_by_capability("collect_wifi_bluetooth_probe_evidence")
        client = getattr(adapter, "client", None) if adapter is not None else None
        if client is None:
            raise RuntimeError("WiFi/蓝牙探针适配器不可用")
        result: dict[str, Any] = {"probe_id": probe_id, "queried_at": utc_now_iso(), "items": {}, "errors": []}
        for kind in ("wifi_ap", "wifi_client", "bluetooth"):
            try:
                targets = dict(client.list_targets(probe_id=probe_id, kind=kind, status="all", page=1, page_size=max(1, min(100, page_size))) or {})
                observations = dict(client.list_observations(probe_id=probe_id, kind=kind, view="current", active_minutes=30, page=1, page_size=max(1, min(100, page_size))) or {})
                result["items"][kind] = {
                    "targets": list(targets.get("items") or []),
                    "target_total": targets.get("total", len(targets.get("items") or [])),
                    "observations": list(observations.get("items") or []),
                    "observation_total": observations.get("total", len(observations.get("items") or [])),
                }
            except Exception as exc:
                result["items"][kind] = {"targets": [], "observations": [], "target_total": 0, "observation_total": 0}
                result["errors"].append({"kind": kind, "message": str(exc)})
        return result

    def get_device_evidence_chain(self, task_id: str) -> dict[str, Any]:
        with self._lock:
            task = _deepcopy_json(self._load_task_locked(task_id))
        return build_device_evidence_chain(task, store=get_probe_evidence_store())

    def create_task(self, instruction: str, template_ids: list[str], metadata: dict[str, Any] | None = None) -> dict[str, Any]:
        self.planner.ensure_available()
        instruction = str(instruction or "").strip()
        if not instruction:
            raise ValueError("请输入采集任务指令")
        selected_templates = []
        for template_id in template_ids:
            selected_templates.append(self.get_template(template_id))
        metadata = dict(metadata or {})
        task_id = f"capture_{uuid4().hex[:12]}"
        output_dir = AUTONOMOUS_OUTPUT_ROOT / task_id
        output_dir.mkdir(parents=True, exist_ok=True)
        now = utc_now_iso()
        task = {
            "id": task_id,
            "title": str(metadata.get("title") or instruction[:48]).strip()[:120] or instruction[:48],
            "instruction": instruction,
            "place": str(metadata.get("place") or "").strip(),
            "operator": str(metadata.get("operator") or "operator").strip() or "operator",
            "source": str(metadata.get("source") or "capture_agent").strip() or "capture_agent",
            "parent_task_id": str(metadata.get("parent_task_id") or "").strip(),
            "reasoning_mode": _normalize_reasoning_mode(metadata.get("reasoning_mode")),
            "selected_usrp_devices": _deepcopy_json(list(metadata.get("selected_usrp_devices") or [])),
            "selected_probe_devices": _deepcopy_json(list(metadata.get("selected_probe_devices") or [])),
            "available_devices": {},
            "device_selection": {},
            "probe_evidence_sets": [],
            "probe_assistance": _initial_probe_assistance([]),
            "template_ids": [item["id"] for item in selected_templates],
            "templates": [self._public_template(item) for item in selected_templates],
            "output_dir": str(output_dir.resolve()),
            "status": "planning",
            "pause_reason": None,
            "plan_version": 1,
            "plan_locked": False,
            "approved_at": None,
            "approved_by": None,
            "current_node_id": "template_analysis",
            "selected_node_id": "template_analysis",
            "progress": 0,
            "constraints": {},
            "constraint_sources": {},
            "validation": {},
            "intent_analysis": {},
            "plan_candidate": {},
            "plan_fingerprint": None,
            "approved_plan_fingerprint": None,
            "plan_history": [],
            "planner": self.planner.runtime_info(),
            "nodes": CapturePlanBuilder.build_bootstrap_nodes(),
            "messages": [
                {
                    "id": f"msg_{uuid4().hex[:10]}",
                    "role": "operator",
                    "content": instruction,
                    "created_at": now,
                    "template_ids": [item["id"] for item in selected_templates],
                }
            ],
            "events": [],
            "artifacts": [],
            "generated_code": "",
            "execution_result": {},
            "error": None,
            "created_at": now,
            "updated_at": now,
        }
        self._append_event_locked(task, "task_created", "已创建采集智能体任务，开始生成任务规划。", node_id="template_analysis")
        self._save_task_locked(task)
        self._start_worker(task_id, mode="planning")
        return self._public_task(task)

    def get_task(self, task_id: str) -> dict[str, Any]:
        with self._lock:
            payload = self._load_task_locked(task_id)
            item = self._public_task(payload)
            index = self._task_index_locked(task_id)
            if index is not None:
                item.update(self._task_display_fields(payload, index))
            return item

    def confirm_probe_assistance(self, task_id: str, approved: bool) -> dict[str, Any]:
        with self._lock:
            task = self._load_task_locked(task_id)
            auxiliary = task.setdefault("probe_assistance", _initial_probe_assistance(list(task.get("selected_probe_devices") or [])))
            if not auxiliary.get("enabled"):
                raise ValueError("当前任务未选择 WiFi/蓝牙探针")
            if task.get("status") != "completed":
                raise ValueError("必须等待 USRP 采集流程完成后再调用探针")
            status = str(auxiliary.get("status") or "")
            if status in {"running", "queued"}:
                return self._public_task(task)
            if approved:
                auxiliary["status"] = "queued"
                auxiliary["status_label"] = "准备采集辅助证据"
                auxiliary["confirmed_at"] = utc_now_iso()
                self._set_probe_step(auxiliary, "confirm", "completed", "用户已确认调用 WiFi/蓝牙探针")
                self._append_event_locked(task, "probe_assistance_confirmed", "用户已确认调用 WiFi/蓝牙探针获取辅助证据。", data={"probe_ids": [item.get("probe_id") for item in auxiliary.get("selected_probes", [])]})
            else:
                auxiliary["status"] = "declined"
                auxiliary["status_label"] = "用户已取消辅助采集"
                auxiliary["ended_at"] = utc_now_iso()
                self._set_probe_step(auxiliary, "confirm", "skipped", "用户选择不调用探针")
                self._append_event_locked(task, "probe_assistance_declined", "用户选择不调用 WiFi/蓝牙探针；USRP 结果保持完成。", level="warning")
            self._save_task_locked(task)
        if approved:
            self._start_probe_worker(task_id)
        return self.get_task(task_id)

    def _start_probe_worker(self, task_id: str) -> None:
        with self._lock:
            active = self._probe_workers.get(task_id)
            if active is not None and active.is_alive():
                return
            thread = threading.Thread(target=self._run_probe_assistance, args=(task_id,), daemon=True, name=f"capture-probe-{task_id[-6:]}")
            self._probe_workers[task_id] = thread
            thread.start()

    @staticmethod
    def _set_probe_step(auxiliary: dict[str, Any], step_id: str, status: str, message: str) -> None:
        step = next((item for item in auxiliary.setdefault("steps", []) if item.get("id") == step_id), None)
        if step is None:
            step = {"id": step_id, "title": step_id}
            auxiliary["steps"].append(step)
        step.update({"status": status, "message": message, "updated_at": utc_now_iso()})

    @staticmethod
    def _probe_fallback_summary(result: dict[str, Any]) -> str:
        counts = dict(result.get("counts") or {})
        by_kind = dict(counts.get("by_kind") or {})
        parts = [
            f"已从 {counts.get('online_devices', 0)} 台在线探针获取辅助证据，共发现 {counts.get('targets', 0)} 个目标、{counts.get('observations', 0)} 条当前观测。",
            f"其中 WiFi 热点 {dict(by_kind.get('wifi_ap') or {}).get('targets', 0)} 个，WiFi 客户端 {dict(by_kind.get('wifi_client') or {}).get('targets', 0)} 个，蓝牙设备 {dict(by_kind.get('bluetooth') or {}).get('targets', 0)} 个。",
            "以上信息仅作为环境侧辅助证据，不改变已经成功完成的 USRP 采集结果。",
        ]
        if result.get("errors"):
            parts.append(f"部分接口或设备返回失败，共 {len(result.get('errors') or [])} 项，请结合详细信息复核。")
        return "".join(parts)

    def _run_probe_assistance(self, task_id: str) -> None:
        try:
            with self._lock:
                task = self._load_task_locked(task_id)
                auxiliary = task.setdefault("probe_assistance", _initial_probe_assistance(list(task.get("selected_probe_devices") or [])))
                auxiliary.update({"status": "running", "status_label": "正在采集 WiFi/蓝牙辅助证据", "started_at": utc_now_iso(), "ended_at": None})
                self._append_event_locked(task, "probe_assistance_started", "WiFi/蓝牙探针辅助证据流程已启动。")
                self._save_task_locked(task)
                selected_probes = list(auxiliary.get("selected_probes") or [])
                reasoning_mode = _normalize_reasoning_mode(task.get("reasoning_mode"))

            runtime = getattr(getattr(self.platform, "app", None), "runtime", None)
            if runtime is None:
                raise RuntimeError("平台运行时不可用")
            run = Run(
                id=new_id("probe_run"), task_id=self.platform.task.id, trigger_kind=RunTriggerKind.CHAT,
                trigger_event_id=None, trigger_message_id=None, agent_profile="capture_agent.probe_assistance",
                status=RunStatus.RUNNING, step_budget=1, step_count=0, started_at=utc_now(), conversation_id=None,
            )
            message = ChatMessage(
                id=new_id("probe_message"), conversation_id="capture-agent-probe", task_id=self.platform.task.id,
                role=ChatRole.OPERATOR, content="采集 WiFi/蓝牙探针辅助证据", run_id=run.id, created_at=utc_now(),
            )

            def progress_handler(event_type: str, payload: dict[str, Any]) -> None:
                if event_type != "probe_progress":
                    return
                stage = str(payload.get("stage") or "")
                step_id = "devices" if stage.startswith("devices") else ("targets" if stage.startswith("targets") else "observations")
                status = "failed" if stage.endswith("failed") else ("completed" if stage.endswith("completed") else "in_progress")
                message_text = str(payload.get("message") or stage)
                with self._lock:
                    current = self._load_task_locked(task_id)
                    aux = current["probe_assistance"]
                    self._set_probe_step(aux, step_id, status, message_text)
                    aux.setdefault("logs", []).append({"timestamp": utc_now_iso(), "stage": stage, "message": message_text, "data": {k: v for k, v in payload.items() if k not in {"devices"}}})
                    self._append_event_locked(current, "probe_assistance_progress", message_text, data={"stage": stage, "step_id": step_id, **{k: v for k, v in payload.items() if k != "devices"}})
                    self._save_task_locked(current)

            outcome = runtime.tool_registry.execute(
                "collect_wifi_bluetooth_probe_evidence",
                {
                    "probe_ids": [str(item.get("probe_id") or "") for item in selected_probes if item.get("probe_id")],
                    "kinds": ["wifi_ap", "wifi_client", "bluetooth"],
                    "active_minutes": 30,
                    "page_size": 200,
                    "capture_task_id": task_id,
                },
                ToolContext(
                    task=self.platform.task, run=run, trigger_event=None, trigger_message=message,
                    state_repo=runtime.state_repo, case_repo=runtime.case_repo, knowledge_base=runtime.knowledge_base,
                    device_registry=runtime.device_registry, llm_client=runtime.llm_client, stream_handler=progress_handler,
                    nl2sql_options=NL2SQLSessionConfig(), asset_manager=runtime.asset_manager,
                    document_index=runtime.document_index, database_catalog=runtime.database_catalog, cancel_checker=None,
                ),
            )
            if outcome.result.status != "success":
                raise RuntimeError(outcome.result.error or "探针辅助证据工具执行失败")
            result = dict(outcome.result.data or {})
            evidence_pack = build_probe_evidence_pack(result)

            output_dir = Path(str(task.get("output_dir") or AUTONOMOUS_OUTPUT_ROOT / task_id)).resolve()
            output_dir.mkdir(parents=True, exist_ok=True)
            evidence_path = output_dir / "wifi_bluetooth_probe_evidence.json"
            evidence_path.write_text(
                json.dumps(
                    {
                        "summary": "",
                        "result": result,
                        "evidence_pack": {
                            "overview": evidence_pack["overview"],
                            "manifest": evidence_pack["manifest"],
                            "llm_evidence": evidence_pack["llm_evidence"],
                        },
                    },
                    ensure_ascii=False,
                    indent=2,
                    default=str,
                ),
                encoding="utf-8",
            )

            with self._lock:
                current = self._load_task_locked(task_id)
                aux = current["probe_assistance"]
                aux["details"] = evidence_pack["overview"]
                aux["errors"] = list(result.get("errors") or [])
                self._set_probe_step(aux, "summary", "in_progress", "正在由当前模型总结辅助证据")
                self._append_event_locked(current, "probe_assistance_summary_started", "正在智能总结 WiFi/蓝牙辅助证据。")
                self._save_task_locked(current)

            summary_error = ""
            summary_parts: list[str] = []
            reasoning_parts: list[str] = []

            def summary_stream(event_type: str, payload: dict[str, Any]) -> None:
                delta = str(payload.get("delta") or payload.get("text") or "")
                if event_type == "content_delta" and delta:
                    summary_parts.append(delta)
                elif event_type == "reasoning_delta" and delta:
                    reasoning_parts.append(delta)
                else:
                    return
                with self._lock:
                    current = self._load_task_locked(task_id)
                    aux = current["probe_assistance"]
                    if summary_parts:
                        aux["summary_stream"] = "".join(summary_parts)[-20000:]
                    if reasoning_parts and reasoning_mode == "deep":
                        aux["reasoning"] = "".join(reasoning_parts)[-30000:]
                    self._append_event_locked(current, "probe_assistance_model_delta", "辅助证据摘要正在生成。", data={"kind": "content" if event_type == "content_delta" else "reasoning", "delta": delta})
                    self._save_task_locked(current)

            try:
                summary = self.planner.summarize_probe_evidence(
                    task_context={"id": task_id, "instruction": task.get("instruction"), "place": task.get("place")},
                    evidence=evidence_pack["llm_evidence"], reasoning_mode=reasoning_mode, stream_handler=summary_stream, cancel_checker=None,
                )
            except Exception as exc:
                summary_error = str(exc)
                summary = self._probe_fallback_summary(result)

            evidence_path.write_text(
                json.dumps(
                    {
                        "summary": summary,
                        "result": result,
                        "evidence_pack": {
                            "overview": evidence_pack["overview"],
                            "manifest": evidence_pack["manifest"],
                            "llm_evidence": evidence_pack["llm_evidence"],
                        },
                    },
                    ensure_ascii=False,
                    indent=2,
                    default=str,
                ),
                encoding="utf-8",
            )

            with self._lock:
                current = self._load_task_locked(task_id)
                aux = current["probe_assistance"]
                aux["summary"] = summary
                aux["summary_stream"] = summary
                if summary_error:
                    aux.setdefault("errors", []).append({"stage": "summary", "message": summary_error})
                partial = bool(aux.get("errors") or result.get("partial_failure"))
                aux["status"] = "partial_failed" if partial else "completed"
                aux["status_label"] = "辅助证据部分失败" if partial else "辅助证据采集完成"
                aux["ended_at"] = utc_now_iso()
                self._set_probe_step(aux, "summary", "completed" if not summary_error else "failed", "辅助证据智能总结完成" if not summary_error else "模型总结失败，已生成规则化摘要")
                artifact = self._add_artifact_locked(current, evidence_path, "probe_auxiliary_evidence")
                aux["detail_artifact_id"] = artifact["id"]
                aux["detail_url"] = f"/api/capture-agent/tasks/{task_id}/probe-assistance/details"
                self._append_event_locked(
                    current, "probe_assistance_completed", aux["status_label"],
                    level="warning" if partial else "info", data={"status": aux["status"], "counts": result.get("counts", {}), "errors": aux.get("errors", [])},
                )
                self._save_task_locked(current)
        except Exception as exc:
            with self._lock:
                try:
                    task = self._load_task_locked(task_id)
                except KeyError:
                    return
                auxiliary = task.setdefault("probe_assistance", _initial_probe_assistance(list(task.get("selected_probe_devices") or [])))
                auxiliary["status"] = "partial_failed"
                auxiliary["status_label"] = "辅助证据部分失败"
                auxiliary["ended_at"] = utc_now_iso()
                auxiliary.setdefault("errors", []).append({"stage": "workflow", "message": str(exc)})
                self._append_event_locked(task, "probe_assistance_failed", f"辅助证据部分失败：{exc}", level="warning")
                self._save_task_locked(task)

    def rename_task(self, task_id: str, title: str) -> dict[str, Any]:
        clean_title = str(title or "").strip()
        if not clean_title:
            raise ValueError("任务名称不能为空")
        with self._lock:
            task = self._load_task_locked(task_id)
            task["title"] = clean_title[:120]
            task["updated_at"] = utc_now_iso()
            self._append_event_locked(task, "task_renamed", f"任务已重命名为：{task['title']}", node_id=task.get("current_node_id"))
            self._save_task_locked(task)
        return self.get_task(task_id)

    def delete_task(self, task_id: str) -> None:
        with self._lock:
            task = self._load_task_locked(task_id)
            self._cancel_events.setdefault(task_id, threading.Event()).set()
            output_dir = Path(str(task.get("output_dir") or ""))
            task_path = self.tasks_dir / f"{task_id}.json"
            if task_path.exists():
                task_path.unlink()
            if output_dir.exists() and output_dir.is_dir():
                try:
                    shutil.rmtree(output_dir, ignore_errors=True)
                except Exception:
                    pass
            self._workers.pop(task_id, None)
            self._pause_events.pop(task_id, None)

    def update_approval_parameters(self, task_id: str, parameters: dict[str, Any]) -> dict[str, Any]:
        with self._lock:
            task = self._load_task_locked(task_id)
            if task.get("status") != "awaiting_approval" or task.get("plan_locked"):
                raise ValueError("只有待批准且未锁定的任务可以修改采集参数")
            selected_types = set((task.get("device_selection") or {}).get("selected_device_types") or [])
            if "usrp" not in selected_types:
                raise ValueError("当前计划未选择 USRP，没有需要修改的频谱采集参数")
            constraints = self._coerce_approval_constraints(dict(task.get("constraints") or {}), dict(parameters or {}))
            validation = CapturePlanBuilder.validate_constraints(constraints)
            task["constraints"] = constraints
            task["validation"] = validation
            sources = dict(task.get("constraint_sources") or {})
            for key in parameters:
                sources[key] = "operator_approval_edit"
            task["constraint_sources"] = sources
            task["updated_at"] = utc_now_iso()
            self._save_task_locked(task)

        # 复用同一自适应编排器：USRP 组按新 plan_id 重建，探针组与 LLM 设备选择保持不变。
        self._generate_plan_summary(task_id)
        with self._lock:
            current = self._load_task_locked(task_id)
            approval = self._node(current, "approval")
            approval.update({"status": "blocked", "progress": 0, "updated_at": utc_now_iso(), "summary": "参数已由用户修改，等待重新批准"})
            approval["logs"].append(self._log_entry("用户在审批界面修改了 USRP 采集参数；已在不改变设备选择的前提下重新生成对应执行组。", "warning", parameters))
            current["current_node_id"] = "approval"
            current["updated_at"] = utc_now_iso()
            self._append_event_locked(current, "approval_parameters_updated", "USRP 审批参数已更新，自适应设备计划已重新生成。", node_id="approval", data={"validation": current.get("validation") or {}})
            self._save_task_locked(current)
        return self.get_task(task_id)

    def get_task_internal(self, task_id: str) -> dict[str, Any]:
        with self._lock:
            return _deepcopy_json(self._load_task_locked(task_id))

    def approve_task(self, task_id: str, approved_by: str = "operator") -> dict[str, Any]:
        with self._lock:
            task = self._load_task_locked(task_id)
            if task["status"] != "awaiting_approval":
                raise ValueError("当前任务不处于待批准状态")
            if self._plan_requires_device_access(task):
                conflict = self._find_conflicting_real_task_locked(task_id)
                if conflict is not None:
                    raise ValueError(f"真实设备当前已被任务 {conflict.get('id')} 占用或保留，请先完成/取消该任务")
            current_fingerprint = self._task_plan_fingerprint(task)
            expected_fingerprint = str(task.get("plan_fingerprint") or "")
            if not expected_fingerprint or current_fingerprint != expected_fingerprint:
                raise ValueError("计划结构在审批前发生变化，必须重新生成计划版本")
            node = self._node(task, "approval")
            existing_llm_outputs = {
                key: value
                for key, value in dict(node.get("outputs") or {}).items()
                if str(key).startswith("llm_")
            }
            node.update(
                {
                    "status": "completed",
                    "progress": 100,
                    "started_at": node.get("started_at") or utc_now_iso(),
                    "ended_at": utc_now_iso(),
                    "updated_at": utc_now_iso(),
                    "summary": f"自适应设备计划 v{task['plan_version']} 已由用户批准并锁定指纹",
                    "outputs": {
                        **existing_llm_outputs,
                        "approved_by": approved_by,
                        "plan_version": task["plan_version"],
                        "plan_fingerprint": expected_fingerprint,
                    },
                }
            )
            node["logs"].append(self._log_entry("用户批准了自适应设备计划；执行器已锁定设备选择、节点依赖、工具名称与工具参数。", "success"))
            task["status"] = "running"
            task["pause_reason"] = None
            task["plan_locked"] = True
            task["approved_at"] = utc_now_iso()
            task["approved_by"] = approved_by
            task["approved_plan_fingerprint"] = expected_fingerprint
            execution_nodes = self._execution_nodes_in_order(task)
            task["current_node_id"] = execution_nodes[0]["id"] if execution_nodes else "approval"
            self._append_event_locked(task, "plan_approved", "自适应设备计划已批准并锁定，开始严格按依赖执行。", node_id="approval")
            self._save_task_locked(task)
        self._schedule_status_summary(task_id, "approval", force=True, phase="node_terminal")
        self._start_worker(task_id, mode="execution")
        return self.get_task(task_id)

    def pause_task(self, task_id: str) -> dict[str, Any]:
        with self._lock:
            task = self._load_task_locked(task_id)
            if task["status"] not in {"running", "planning"}:
                raise ValueError("只有规划中或执行中的任务可以暂停")
            task["status"] = "paused"
            task["pause_reason"] = "manual"
            pause_event = self._pause_events.setdefault(task_id, threading.Event())
            pause_event.set()
            self._append_event_locked(task, "task_paused", "用户已暂停任务；当前工具调用完成后不会进入下一节点。", node_id=task.get("current_node_id"))
            self._save_task_locked(task)
        return self.get_task(task_id)

    def resume_task(self, task_id: str) -> dict[str, Any]:
        with self._lock:
            task = self._load_task_locked(task_id)
            if task["status"] != "paused":
                raise ValueError("当前任务未暂停")
            if task.get("plan_locked") and self._plan_requires_device_access(task):
                conflict = self._find_conflicting_real_task_locked(task_id)
                if conflict is not None:
                    raise ValueError(f"真实设备当前已被任务 {conflict.get('id')} 占用或保留，请先完成/取消该任务")
            if task.get("pause_reason") == "node_failure" and task.get("current_node_id"):
                failed_node = self._node(task, str(task["current_node_id"]))
                if failed_node.get("status") == "failed":
                    failed_node.update(
                        {
                            "status": "pending",
                            "progress": 0,
                            "ended_at": None,
                            "updated_at": utc_now_iso(),
                            "summary": "失败节点已重置，等待继续执行",
                        }
                    )
                    failed_node["logs"].append(self._log_entry("用户继续任务，失败节点已重置为待执行。", "warning"))
            task["status"] = "planning" if not task.get("plan_locked") else "running"
            task["pause_reason"] = None
            self._pause_events.setdefault(task_id, threading.Event()).clear()
            self._append_event_locked(task, "task_resumed", "任务已继续，将从未完成节点恢复。", node_id=task.get("current_node_id"))
            self._save_task_locked(task)
        self._start_worker(task_id, mode="planning" if not task.get("plan_locked") else "execution")
        return self.get_task(task_id)

    def cancel_task(self, task_id: str) -> dict[str, Any]:
        cancelled_node_id = ""
        with self._lock:
            task = self._load_task_locked(task_id)
            if task["status"] in TERMINAL_TASK_STATES:
                return self._public_task(task)
            self._cancel_events.setdefault(task_id, threading.Event()).set()
            task["status"] = "cancelled"
            current = task.get("current_node_id")
            if current:
                node = self._node(task, current)
                if node["status"] not in TERMINAL_NODE_STATES:
                    node.update({"status": "cancelled", "ended_at": utc_now_iso(), "updated_at": utc_now_iso(), "summary": "用户取消任务"})
                    cancelled_node_id = str(current)
            self._append_event_locked(task, "task_cancelled", "用户已取消采集任务。", node_id=current)
            self._save_task_locked(task)
        if cancelled_node_id:
            self._schedule_status_summary(task_id, cancelled_node_id, force=True, phase="node_terminal")
        return self.get_task(task_id)

    def add_message(self, task_id: str, content: str) -> dict[str, Any]:
        self.planner.ensure_available()
        content = str(content or "").strip()
        if not content:
            raise ValueError("消息不能为空")
        with self._lock:
            task = self._load_task_locked(task_id)
            allowed_revision_states = {"awaiting_approval", "failed", "completed", "cancelled"}
            if task.get("status") == "paused" and task.get("pause_reason") == "node_failure":
                allowed_revision_states.add("paused")
            if task["status"] not in allowed_revision_states:
                raise ValueError("任务正在执行；请先暂停或等待完成，再追加新指令")
            task["messages"].append(
                {"id": f"msg_{uuid4().hex[:10]}", "role": "operator", "content": content, "created_at": utc_now_iso(), "template_ids": []}
            )
            task["instruction"] = f"{task['instruction']}\n追加要求：{content}"
            task["plan_version"] = int(task.get("plan_version") or 1) + 1
            task["plan_locked"] = False
            task["approved_at"] = None
            task["approved_by"] = None
            task["status"] = "planning"
            task["pause_reason"] = None
            task["error"] = None
            task["generated_code"] = ""
            task["execution_result"] = {}
            task["approved_plan_fingerprint"] = None
            task["plan_fingerprint"] = None
            task["plan_candidate"] = {}
            task["intent_analysis"] = {}
            task["constraint_sources"] = {}
            task["available_devices"] = {}
            task["device_selection"] = {}
            task["probe_evidence_sets"] = []
            task["nodes"] = CapturePlanBuilder.build_bootstrap_nodes()
            task["probe_assistance"] = _initial_probe_assistance([])
            task["current_node_id"] = "template_analysis"
            self._cancel_events.setdefault(task_id, threading.Event()).clear()
            self._pause_events.setdefault(task_id, threading.Event()).clear()
            self._append_event_locked(task, "plan_revision_started", f"收到追加指令，开始生成计划版本 v{task['plan_version']}。", node_id="template_analysis")
            self._save_task_locked(task)
        self._start_worker(task_id, mode="planning")
        return self.get_task(task_id)

    def save_node_note(self, task_id: str, node_id: str, note: str) -> dict[str, Any]:
        with self._lock:
            task = self._load_task_locked(task_id)
            node = self._node(task, node_id)
            node["notes"] = str(note or "")
            node["updated_at"] = utc_now_iso()
            self._append_event_locked(task, "node_note_saved", f"已保存节点“{node['title']}”的人工备注。", node_id=node_id)
            self._save_task_locked(task)
        return self.get_task(task_id)

    def events_since(self, task_id: str, after_seq: int = 0) -> list[dict[str, Any]]:
        with self._lock:
            task = self._load_task_locked(task_id)
            return [item for item in task.get("events", []) if int(item.get("seq") or 0) > after_seq]


    def list_npz_artifacts(self, *, limit: int = 50000) -> list[dict[str, Any]]:
        """Return all known .npz capture products from capture-agent tasks and output folders."""
        items: list[dict[str, Any]] = []
        seen: set[str] = set()
        with self._lock:
            payloads = self._all_task_payloads_locked()
            task_index = {str(task.get("id") or ""): index for index, task in enumerate(payloads, start=1)}
            for task in reversed(payloads):
                task_id = str(task.get("id") or "")
                task_title = str(task.get("title") or task.get("instruction") or task_id)
                display = self._task_display_fields(task, task_index.get(task_id, 0) or 0).get("display_title") if task_id else task_title
                for artifact in task.get("artifacts", []) or []:
                    file_name = str(artifact.get("file_name") or "")
                    path = Path(str(artifact.get("path") or "")).expanduser().resolve()
                    if path.suffix.lower() != ".npz" or not path.exists() or not path.is_file():
                        continue
                    resolved = str(path)
                    if resolved in seen:
                        continue
                    seen.add(resolved)
                    items.append({
                        "id": f"artifact:{task_id}:{artifact.get('id')}",
                        "source": "capture_artifact",
                        "task_id": task_id,
                        "task_title": task_title,
                        "task_display_title": display or task_title,
                        "artifact_id": artifact.get("id"),
                        "file_name": path.name,
                        "path": resolved,
                        "size_bytes": path.stat().st_size,
                        "created_at": artifact.get("created_at") or task.get("created_at"),
                        "updated_at": artifact.get("updated_at") or task.get("updated_at"),
                    })
                    if len(items) >= limit:
                        return items
            scan_roots = [AUTONOMOUS_OUTPUT_ROOT, self.platform._uploads_root if hasattr(self.platform, "_uploads_root") else None]
            for root in scan_roots:
                if root is None:
                    continue
                root_path = Path(root).expanduser().resolve()
                if not root_path.exists():
                    continue
                for path in sorted(root_path.rglob("*.npz"), key=lambda item: item.stat().st_mtime if item.exists() else 0, reverse=True):
                    if not path.exists() or not path.is_file():
                        continue
                    resolved = str(path.resolve())
                    if resolved in seen:
                        continue
                    seen.add(resolved)
                    items.append({
                        "id": f"file:{hashlib.sha1(resolved.encode('utf-8')).hexdigest()[:16]}",
                        "source": "npz_file",
                        "task_id": "",
                        "task_title": "",
                        "task_display_title": "本地 NPZ 文件",
                        "artifact_id": "",
                        "file_name": path.name,
                        "path": resolved,
                        "size_bytes": path.stat().st_size,
                        "created_at": utc_now_iso(),
                        "updated_at": utc_now_iso(),
                    })
                    if len(items) >= limit:
                        return items
        return items

    def resolve_artifact(self, task_id: str, artifact_id: str) -> tuple[Path, str]:
        with self._lock:
            task = self._load_task_locked(task_id)
            artifact = next((item for item in task.get("artifacts", []) if item.get("id") == artifact_id), None)
            if artifact is None:
                raise KeyError(artifact_id)
            path = Path(str(artifact.get("path") or "")).resolve()
            if not path.exists() or not path.is_file():
                raise FileNotFoundError(path)
            return path, str(artifact.get("file_name") or path.name)

    def describe_artifact(self, task_id: str, artifact_id: str) -> dict[str, Any]:
        """Return safe, bounded metadata and a human-friendly preview for one artifact."""
        with self._lock:
            task = self._load_task_locked(task_id)
            artifact = next((item for item in task.get("artifacts", []) if item.get("id") == artifact_id), None)
            if artifact is None:
                raise KeyError(artifact_id)
            path = Path(str(artifact.get("path") or "")).resolve()
            if not path.exists() or not path.is_file():
                raise FileNotFoundError(path)
            public = {key: value for key, value in artifact.items() if key != "path"}
        stat = path.stat()
        suffix = path.suffix.lower()
        detail: dict[str, Any] = {
            "artifact": {
                **public,
                "extension": suffix or "无扩展名",
                "size_bytes": stat.st_size,
                "modified_at": datetime.fromtimestamp(stat.st_mtime, tz=timezone.utc).isoformat(),
                "download_url": f"/api/capture-agent/tasks/{task_id}/artifacts/{artifact_id}",
            },
            "preview": {"type": "metadata", "message": "该文件暂无可视化预览，可下载后查看。"},
        }
        if suffix == ".npz":
            import numpy as np

            arrays: list[dict[str, Any]] = []
            with np.load(path, allow_pickle=False) as payload:
                for key in payload.files:
                    array = payload[key]
                    item: dict[str, Any] = {
                        "name": key,
                        "shape": list(array.shape),
                        "dtype": str(array.dtype),
                        "elements": int(array.size),
                    }
                    if array.size and np.issubdtype(array.dtype, np.number):
                        flat = array.reshape(-1)
                        finite = flat[np.isfinite(flat)]
                        if finite.size:
                            item["min"] = float(np.min(finite))
                            item["mean"] = float(np.mean(finite))
                            item["max"] = float(np.max(finite))
                        item["sample"] = [float(value) for value in flat[:8]]
                    arrays.append(item)
            detail["preview"] = {"type": "npz", "arrays": arrays}
        elif suffix == ".json":
            payload = json.loads(path.read_text(encoding="utf-8"))
            if isinstance(payload, dict):
                top_level = [
                    {"key": str(key), "type": type(value).__name__, "count": len(value) if hasattr(value, "__len__") else None}
                    for key, value in payload.items()
                ]
                sample = {key: value for key, value in list(payload.items())[:12] if not isinstance(value, (list, dict))}
            else:
                top_level = [{"key": "root", "type": type(payload).__name__, "count": len(payload) if hasattr(payload, "__len__") else None}]
                sample = {}
            detail["preview"] = {"type": "json", "top_level": top_level, "sample": sample}
        elif suffix in {".txt", ".md", ".log", ".py", ".yaml", ".yml"}:
            detail["preview"] = {"type": "text", "text": path.read_text(encoding="utf-8", errors="replace")[:16000]}
        elif suffix == ".csv":
            import csv

            with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
                reader = csv.reader(handle)
                rows = []
                for index, row in enumerate(reader):
                    rows.append(row[:30])
                    if index >= 100:
                        break
            detail["preview"] = {"type": "table", "rows": rows}
        elif suffix == ".docx":
            document = Document(path)
            paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()][:300]
            detail["preview"] = {"type": "text", "text": "\n".join(paragraphs)[:16000]}
        return detail

    def get_probe_evidence_details(self, task_id: str) -> dict[str, Any]:
        """Load the complete probe artifact on demand instead of bloating every task response."""
        with self._lock:
            task = self._load_task_locked(task_id)
            artifact = next(
                (item for item in reversed(task.get("artifacts", [])) if item.get("kind") == "probe_auxiliary_evidence"),
                None,
            )
            auxiliary = _deepcopy_json(task.get("probe_assistance") or {})
            if artifact is None:
                return {
                    "summary": auxiliary.get("summary") or auxiliary.get("summary_stream") or "",
                    "result": {},
                    "overview": auxiliary.get("details") or {},
                    "status": auxiliary.get("status"),
                }
            path = Path(str(artifact.get("path") or "")).resolve()
            if not path.exists() or not path.is_file():
                raise FileNotFoundError(path)
        payload = json.loads(path.read_text(encoding="utf-8"))
        result = dict(payload.get("result") or {}) if isinstance(payload, dict) else {}
        evidence_pack = dict(payload.get("evidence_pack") or {}) if isinstance(payload, dict) else {}
        return {
            "summary": str(payload.get("summary") or auxiliary.get("summary") or "") if isinstance(payload, dict) else "",
            "status": auxiliary.get("status"),
            "status_label": auxiliary.get("status_label"),
            "result": result,
            "overview": evidence_pack.get("overview") or auxiliary.get("details") or {},
            "artifact": {
                "id": artifact.get("id"),
                "file_name": artifact.get("file_name"),
                "size_bytes": path.stat().st_size,
                "download_url": f"/api/capture-agent/tasks/{task_id}/artifacts/{artifact.get('id')}",
            },
        }

    # ---------- workers ----------
    def _start_worker(self, task_id: str, mode: str) -> None:
        with self._lock:
            active = self._workers.get(task_id)
            if active is not None and active.is_alive():
                if task_id not in self._restart_pending:
                    self._restart_pending.add(task_id)
                    waiter = threading.Thread(
                        target=self._restart_worker_after,
                        args=(task_id, active, mode),
                        daemon=True,
                        name=f"capture-agent-restart-{task_id[-6:]}",
                    )
                    waiter.start()
                return
            self._pause_events.setdefault(task_id, threading.Event())
            self._cancel_events.setdefault(task_id, threading.Event())
            thread = threading.Thread(target=self._worker, args=(task_id, mode), daemon=True, name=f"capture-agent-{task_id[-6:]}")
            self._workers[task_id] = thread
            thread.start()

    def _restart_worker_after(self, task_id: str, active: threading.Thread, requested_mode: str) -> None:
        active.join()
        should_restart = False
        mode = requested_mode
        with self._lock:
            self._restart_pending.discard(task_id)
            try:
                task = self._load_task_locked(task_id)
            except KeyError:
                return
            auxiliary = task.get("probe_assistance") if isinstance(task.get("probe_assistance"), dict) else None
            if task.get("status") == "completed" and auxiliary and auxiliary.get("status") in {"running", "queued"}:
                auxiliary["status"] = "partial_failed"
                auxiliary["status_label"] = "辅助证据部分失败"
                auxiliary["ended_at"] = utc_now_iso()
                auxiliary.setdefault("errors", []).append({"stage": "workflow", "message": "服务重启中断了探针辅助流程"})
                self._append_event_locked(task, "probe_assistance_recovered", "服务重启中断探针辅助流程，已标记为辅助证据部分失败。", level="warning")
                self._save_task_locked(task)
            if task.get("status") in {"planning", "running"}:
                should_restart = True
                mode = "execution" if task.get("plan_locked") else "planning"
        if should_restart:
            self._start_worker(task_id, mode)

    def _worker(self, task_id: str, mode: str) -> None:
        try:
            if mode == "planning":
                self._run_planning(task_id)
            else:
                self._run_execution(task_id)
        except Exception as exc:
            failed_node_id = ""
            with self._lock:
                try:
                    task = self._load_task_locked(task_id)
                except KeyError:
                    return
                if task.get("status") == "cancelled":
                    return
                current = task.get("current_node_id")
                if current:
                    node = self._node(task, current)
                    if node["status"] not in TERMINAL_NODE_STATES:
                        node.update(
                            {
                                "status": "failed",
                                "progress": node.get("progress", 0),
                                "ended_at": utc_now_iso(),
                                "updated_at": utc_now_iso(),
                                "summary": f"节点失败：{exc}",
                            }
                        )
                        node["logs"].append(self._log_entry(str(exc), "error"))
                        failed_node_id = str(current)
                task["status"] = "failed"
                task["error"] = str(exc)
                self._append_event_locked(task, "task_failed", f"任务失败：{exc}", node_id=current, level="error")
                self._save_task_locked(task)
            if failed_node_id:
                self._schedule_status_summary(task_id, failed_node_id, force=True, phase="node_terminal")

    def _run_planning(self, task_id: str) -> None:
        planning_nodes = ["template_analysis", "device_discovery", "device_selection", "instruction_analysis", "plan_generation"]
        for node_id in planning_nodes:
            if self._should_stop(task_id):
                return
            self._wait_if_paused(task_id)
            with self._lock:
                task = self._load_task_locked(task_id)
                if task["status"] == "cancelled":
                    return
                node = self._node(task, node_id)
                if node["status"] == "completed":
                    continue
                self._assert_dependencies(task, node)
                self._begin_node_locked(task, node, f"开始{node['title']}。")
                self._save_task_locked(task)
            self._schedule_status_summary(task_id, node_id, force=True, phase="node_started")
            if node_id == "template_analysis":
                output = self._analyze_templates(task_id)
                self._complete_node(task_id, node_id, "模板解析完成。", output)
            elif node_id == "device_discovery":
                output = self._discover_collection_devices(task_id, refresh_usrp=True)
                self._complete_node(task_id, node_id, "可用设备探测完成。", output)
            elif node_id == "device_selection":
                output = self._select_collection_devices(task_id)
                self._complete_node(task_id, node_id, "设备自主选择完成。", output)
            elif node_id == "instruction_analysis":
                output = self._analyze_instruction(task_id)
                self._complete_node(task_id, node_id, "采集意图与所需参数解析完成。", output)
            else:
                output = self._generate_plan_summary(task_id)
                self._complete_node(task_id, node_id, "自适应多设备任务规划已生成。", output)

        with self._lock:
            task = self._load_task_locked(task_id)
            if task["status"] == "cancelled":
                return
            self._preflight_plan_locked(task)
            approval = self._node(task, "approval")
            approval.update(
                {
                    "status": "blocked",
                    "progress": 0,
                    "updated_at": utc_now_iso(),
                    "summary": f"等待用户批准自适应设备计划 v{task['plan_version']}",
                    "inputs": {"plan_version": task["plan_version"], "requires_explicit_approval": True},
                }
            )
            approval["logs"].append(self._log_entry("批准前，所有真实 USRP / WiFi蓝牙探针采集节点都被状态机阻止。", "warning"))
            task["status"] = "awaiting_approval"
            task["current_node_id"] = "approval"
            task["progress"] = self._task_progress(task)
            self._append_event_locked(task, "plan_ready", "设备选择与自适应采集计划已生成，请检查后批准执行。", node_id="approval")
            self._save_task_locked(task)
        self._schedule_status_summary(task_id, "approval", force=True, phase="node_started")

    def _run_execution(self, task_id: str) -> None:
        while True:
            if self._should_stop(task_id):
                return
            self._wait_if_paused(task_id)
            with self._lock:
                task = self._load_task_locked(task_id)
                if task["status"] == "cancelled":
                    return
                if not task.get("plan_locked"):
                    raise RuntimeError("计划未锁定，执行器拒绝启动")
                self._assert_plan_integrity_locked(task)
                execution_nodes = self._execution_nodes_in_order(task)
                pending = [node for node in execution_nodes if node.get("status") not in TERMINAL_NODE_STATES]
                if not pending:
                    self._persist_task_bundle_locked(task)
                    task["status"] = "completed"
                    task["pause_reason"] = None
                    task["progress"] = 100
                    task["current_node_id"] = execution_nodes[-1]["id"] if execution_nodes else "approval"
                    selected_types = list((task.get("device_selection") or {}).get("selected_device_types") or [])
                    task["messages"].append(
                        {
                            "id": f"msg_{uuid4().hex[:10]}",
                            "role": "assistant",
                            "content": "任务已按智能体自主选择的设备计划完成，结果、探针证据索引与审计文件已统一保存。",
                            "created_at": utc_now_iso(),
                            "template_ids": [],
                        }
                    )
                    self._append_event_locked(
                        task,
                        "task_completed",
                        "自适应设备计划的所有节点均已完成，统一任务目录已写入审计产物。",
                        node_id=task["current_node_id"],
                        data={"selected_device_types": selected_types},
                    )
                    # 旧版“USRP 完成后再确认探针”流程保留兼容 API，但自适应计划不再自动触发。
                    task["probe_assistance"] = _initial_probe_assistance([])
                    self._save_task_locked(task)
                    return

                ready = []
                for candidate in pending:
                    try:
                        self._assert_dependencies(task, candidate)
                        ready.append(candidate)
                    except RuntimeError:
                        continue
                if not ready:
                    blocked = ", ".join(node["id"] for node in pending)
                    raise RuntimeError(f"剩余节点没有可执行项，依赖状态异常：{blocked}")
                node = ready[0]
                node_id = node["id"]
                self._begin_node_locked(task, node, f"开始{node['title']}。")
                self._save_task_locked(task)
            self._schedule_status_summary(task_id, node_id, force=True, phase="node_started")

            max_attempts = max(1, min(3, int(node.get("max_attempts") or 2)))
            completed = False
            for attempt_index in range(max_attempts):
                try:
                    output, summary = self._execute_planned_node(task_id, node_id)
                    self._complete_node(task_id, node_id, summary, output)
                    completed = True
                    break
                except Exception as exc:
                    if self._should_stop(task_id):
                        return
                    with self._lock:
                        current = self._load_task_locked(task_id)
                        failed_node = self._node(current, node_id)
                        failed_node["logs"].append(self._log_entry(f"第 {attempt_index + 1} 次执行失败：{exc}", "error"))
                        failed_node["updated_at"] = utc_now_iso()
                        if attempt_index + 1 < max_attempts:
                            failed_node["status"] = "pending"
                            failed_node["progress"] = 0
                            failed_node["summary"] = f"节点执行失败，准备第 {attempt_index + 2} 次尝试"
                            self._append_event_locked(current, "node_retry", f"节点“{failed_node['title']}”执行失败，将重试：{exc}", node_id=node_id, level="warning")
                            self._save_task_locked(current)
                        else:
                            failed_node["status"] = "failed"
                            failed_node["ended_at"] = utc_now_iso()
                            failed_node["summary"] = f"达到最大尝试次数后仍失败：{exc}"
                            current["status"] = "paused"
                            current["pause_reason"] = "node_failure"
                            current["error"] = str(exc)
                            self._pause_events.setdefault(task_id, threading.Event()).set()
                            self._append_event_locked(current, "task_paused_on_failure", f"节点“{failed_node['title']}”失败，任务已暂停，可追加指令让 LLM 生成新计划版本：{exc}", node_id=node_id, level="error")
                            self._save_task_locked(current)
                    if attempt_index + 1 < max_attempts:
                        time.sleep(0.35)
                        with self._lock:
                            current = self._load_task_locked(task_id)
                            self._assert_plan_integrity_locked(current)
                            retry_node = self._node(current, node_id)
                            self._begin_node_locked(current, retry_node, f"重试{retry_node['title']}。")
                            self._save_task_locked(current)
                        self._schedule_status_summary(task_id, node_id, force=True, phase="node_started")
                        continue
                    self._schedule_status_summary(task_id, node_id, force=True, phase="node_terminal")
                    return
            if not completed:
                return

    def _execute_planned_node(self, task_id: str, node_id: str) -> tuple[dict[str, Any], str]:
        task = self.get_task_internal(task_id)
        node = self._node(task, node_id)
        executor = str(node.get("executor") or "")
        if executor == "device_scan":
            return self._execute_device_scan(task_id, node_id)
        if executor == "output_file_check":
            return self._execute_output_file_check(task_id, node_id)
        if executor == "execution_summary":
            return self._execute_execution_summary(task_id, node_id)
        if executor == "report_generation":
            return self._execute_report_generation(task_id, node_id)
        kind = str(node.get("kind") or node.get("executor") or "")
        if kind == "tool_call":
            return self._execute_freeform_tool_node(task_id, node_id)
        if kind in {"llm_analysis", "llm_verification", "decision"}:
            return self._execute_freeform_llm_node(task_id, node_id)
        raise RuntimeError(f"节点 {node_id} 使用了不支持的 kind：{kind}")

    def _execute_freeform_llm_node(self, task_id: str, node_id: str) -> tuple[dict[str, Any], str]:
        task = self.get_task_internal(task_id)
        node = self._node(task, node_id)
        dependency_outputs = {
            dep: dict(self._node(task, dep).get("outputs") or {})
            for dep in node.get("dependencies", [])
            if dep != "approval"
        }
        task_context = {
            "id": task["id"],
            "instruction": task["instruction"],
            "constraints": task.get("constraints", {}),
            "intent_analysis": task.get("intent_analysis", {}),
            "plan_version": task.get("plan_version"),
            "completion_contract": (task.get("plan_candidate") or {}).get("completion_contract", []),
            "output_dir": task.get("output_dir"),
        }
        response = self.planner.execute_llm_node(
            task_context=task_context,
            node={
                "id": node.get("id"),
                "title": node.get("title"),
                "description": node.get("description"),
                "kind": node.get("kind"),
                "success_criteria": node.get("success_criteria", []),
                "expected_evidence": node.get("expected_evidence", []),
            },
            dependency_outputs=dependency_outputs,
            reasoning_mode=_normalize_reasoning_mode(task.get("reasoning_mode")),
            stream_handler=self._planner_stream_handler(task_id, node_id),
            cancel_checker=lambda: self._should_stop(task_id),
        )
        result = dict(response.get("result") or {})
        status = str(result.get("status") or "failed")
        if status == "blocked":
            raise RuntimeError(result.get("summary") or "LLM 节点被阻塞")
        if status != "completed":
            raise RuntimeError(result.get("summary") or "LLM 节点执行失败")
        if result.get("should_replan"):
            with self._lock:
                current = self._load_task_locked(task_id)
                self._append_event_locked(
                    current,
                    "llm_replan_suggested",
                    str(result.get("replan_reason") or "LLM 建议修订后续计划"),
                    node_id=node_id,
                    level="warning",
                )
                self._save_task_locked(current)
        return {**result, "raw_model_output": response.get("raw_model_output"), "model": response.get("model")}, str(result.get("summary") or "LLM 节点完成。")

    def _execute_freeform_tool_node(self, task_id: str, node_id: str) -> tuple[dict[str, Any], str]:
        task = self.get_task_internal(task_id)
        node = self._node(task, node_id)
        tool_name = str(node.get("tool_name") or node.get("allowed_tool") or "")
        if tool_name not in self._tool_definitions:
            raise RuntimeError(f"计划引用的工具不可用：{tool_name}")

        raw_args = self._inject_capture_tool_context(tool_name, dict(node.get("tool_arguments") or {}), task)
        args, resolution_notes = self._resolve_tool_arguments(tool_name, raw_args, task, node)
        output = self._execute_tool(task_id, node_id, tool_name, args)

        probe_summary = ""
        probe_set_meta: dict[str, Any] | None = None
        if tool_name == "collect_wifi_bluetooth_probe_evidence":
            evidence_set_id = str(output.get("evidence_set_id") or "")
            if evidence_set_id:
                probe_set_meta = get_probe_evidence_store().get_set(evidence_set_id)
            try:
                probe_summary = self.planner.summarize_probe_evidence(
                    task_context={
                        "id": task.get("id"),
                        "instruction": task.get("instruction"),
                        "device_selection": task.get("device_selection") or {},
                        "constraints": task.get("constraints") or {},
                    },
                    evidence=dict(output.get("llm_evidence") or {}),
                    reasoning_mode=_normalize_reasoning_mode(task.get("reasoning_mode")),
                    stream_handler=self._planner_stream_handler(task_id, node_id),
                    cancel_checker=lambda: self._should_stop(task_id),
                )
            except Exception as exc:
                probe_summary = f"探针证据已成功采集并落库，但智能摘要生成失败：{exc}"
            output["intelligent_summary"] = probe_summary

        with self._lock:
            current = self._load_task_locked(task_id)
            current_node = self._node(current, node_id)
            if resolution_notes:
                current_node["logs"].append(
                    self._log_entry("已按已批准任务参数解析计划参数：" + "；".join(resolution_notes), "warning")
                )
            if tool_name == "generate_usrp_task_code" and output.get("code"):
                current["generated_code"] = str(output.get("code") or "")
            if tool_name == "run_autonomous_usrp_task" and output.get("generated_code"):
                current["generated_code"] = str(output.get("generated_code") or "")
            execution_result = output.get("execution_result") if isinstance(output.get("execution_result"), dict) else None
            if execution_result:
                current["execution_result"] = dict(execution_result)
            if tool_name == "execute_spectrum_collection":
                current["execution_result"] = dict(output)
                output_file = output.get("output_file")
                if output_file:
                    self._add_artifact_locked(current, Path(str(output_file)), "collection_result")
            if tool_name == "collect_wifi_bluetooth_probe_evidence":
                evidence_set_id = str(output.get("evidence_set_id") or "")
                if probe_set_meta and probe_set_meta.get("raw_json_path"):
                    self._add_artifact_locked(current, Path(str(probe_set_meta["raw_json_path"])), "probe_evidence_raw")
                history = list(current.get("probe_evidence_sets") or [])
                history = [item for item in history if item.get("evidence_set_id") != evidence_set_id]
                history.append({
                    "evidence_set_id": evidence_set_id,
                    "created_at": (probe_set_meta or {}).get("created_at") or utc_now_iso(),
                    "counts": output.get("counts") or {},
                    "overview": output.get("overview") or {},
                    "manifest": output.get("manifest") or {},
                    "intelligent_summary": probe_summary,
                })
                current["probe_evidence_sets"] = history[-20:]
                current["latest_probe_summary"] = probe_summary
            for path in self._discover_output_files(output):
                self._add_artifact_locked(current, path, "tool_output")
            self._save_task_locked(current)

        summary = f"工具 {tool_name} 执行完成。"
        if tool_name == "execute_spectrum_collection":
            summary = (
                f"标准频谱采集完成：成功 {output.get('success_count', '-')} / "
                f"{output.get('freq_count', '-')} 个频点，输出 {Path(str(output.get('output_file') or '')).name or '-'}。"
            )
        if tool_name == "collect_wifi_bluetooth_probe_evidence":
            counts = dict(output.get("counts") or {})
            by_kind = dict(counts.get("by_kind") or {})
            summary = (
                f"WiFi/蓝牙探针采集完成并落库：热点 {((by_kind.get('wifi_ap') or {}).get('targets', 0))}、"
                f"WiFi客户端 {((by_kind.get('wifi_client') or {}).get('targets', 0))}、"
                f"蓝牙 {((by_kind.get('bluetooth') or {}).get('targets', 0))} 个目标，证据集 {output.get('evidence_set_id') or '-'}。"
            )
        if isinstance(output.get("execution_result"), dict):
            execution = output["execution_result"]
            summary = f"工具 {tool_name} 执行完成，状态：{execution.get('status', 'unknown')}。"
        return {
            "tool_name": tool_name,
            "resolved_arguments": args,
            "argument_resolution_notes": resolution_notes,
            "result": output,
            **output,
        }, summary

    def _inject_capture_tool_context(self, tool_name: str, args: dict[str, Any], task: dict[str, Any]) -> dict[str, Any]:
        result = _deepcopy_json(args)
        if tool_name in {"generate_usrp_task_code", "execute_usrp_task_code", "run_autonomous_usrp_task"}:
            result.setdefault("task_description", self._tool_task_description(task))
            if not isinstance(result.get("task_plan"), dict):
                result["task_plan"] = self._autonomous_tool_plan(task)
        if tool_name == "execute_usrp_task_code" and not result.get("code"):
            result["code"] = str(task.get("generated_code") or "")
        return result

    def _resolve_tool_arguments(
        self,
        tool_name: str,
        raw_args: dict[str, Any],
        task: dict[str, Any],
        node: dict[str, Any],
    ) -> tuple[dict[str, Any], list[str]]:
        definition = self._tool_definitions[tool_name]
        schema = dict(definition.input_schema or {})
        properties = dict(schema.get("properties") or {})
        required = set(schema.get("required") or [])
        resolved: dict[str, Any] = {}
        notes: list[str] = []

        for key, value in raw_args.items():
            try:
                resolved[key] = self._resolve_plan_value(value, task)
            except RuntimeError as exc:
                inferred = self._infer_tool_argument(tool_name, key, task, node)
                if inferred is _MISSING:
                    if key in required:
                        raise RuntimeError(
                            f"工具 {tool_name} 的必填参数 {key} 无法解析：{exc}。"
                            "请检查该参数引用的前置节点输出，或让计划直接使用 constraints/runtime 占位符。"
                        ) from exc
                    notes.append(f"忽略无法解析的可选参数 {key}")
                    continue
                resolved[key] = inferred
                notes.append(f"参数 {key} 的占位符不可用，已从已批准任务上下文自动推断")

        for key in sorted(required - set(resolved)):
            inferred = self._infer_tool_argument(tool_name, key, task, node)
            if inferred is _MISSING:
                raise RuntimeError(f"工具 {tool_name} 缺少必填参数 {key}，且无法从任务上下文自动推断")
            resolved[key] = inferred
            notes.append(f"已从已批准任务参数解析必填字段 {key}")

        # Resolve optional fields only from approved task context or dependency
        # outputs; never from collection defaults.
        for key in properties:
            if key in resolved:
                continue
            inferred = self._infer_tool_argument(tool_name, key, task, node)
            if inferred is not _MISSING:
                resolved[key] = inferred

        # A model may create a phase-specific autonomous sub-plan.  Resolve it,
        # then clamp it to the operator-approved envelope before any real device
        # action.  This keeps flexible planning without allowing a later node to
        # silently widen frequency, duration, gain or repetition limits.
        if tool_name in {"generate_usrp_task_code", "execute_usrp_task_code", "run_autonomous_usrp_task"}:
            supplied_plan = resolved.get("task_plan")
            if not isinstance(supplied_plan, dict):
                supplied_plan = {}
            resolved["task_plan"] = self._merge_safe_autonomous_subplan(task, supplied_plan)

        self._assert_no_unresolved_placeholders(resolved, tool_name)
        return resolved, notes

    def _infer_tool_argument(
        self,
        tool_name: str,
        key: str,
        task: dict[str, Any],
        node: dict[str, Any],
    ) -> Any:
        constraints = dict(task.get("constraints") or {})
        if key in constraints:
            return _deepcopy_json(constraints[key])
        aliases = {
            "freq": "freq_start_mhz",
            "slice_duration": "dwell_time_sec",
        }
        if key in aliases and aliases[key] in constraints:
            return _deepcopy_json(constraints[aliases[key]])
        if key == "task_description":
            return self._tool_task_description(task)
        if key == "task_plan":
            return self._autonomous_tool_plan(task)
        if key == "code":
            code = str(task.get("generated_code") or "")
            return code if code else _MISSING
        if key == "query":
            return str(task.get("instruction") or "")
        if key == "top_k":
            return 6
        if key == "dev_id":
            value = self._find_dependency_value(task, node, {"dev_id", "device_id"})
            if value is not _MISSING:
                return value
            configured = os.getenv("DEEPEM_USRP_DEVICE_ID")
            return configured.strip() if configured and configured.strip() else _MISSING
        if key == "task_id":
            # Only accept a real device/task identifier emitted by a dependency.
            # The capture-agent task id is a different namespace and must never
            # be sent to query_usrp_task by accident.
            return self._find_dependency_value(task, node, {"task_id", "usrp_task_id"})
        if key == "output_dir":
            return task.get("output_dir") or _MISSING
        return _MISSING

    def _find_dependency_value(
        self,
        task: dict[str, Any],
        node: dict[str, Any],
        keys: set[str],
    ) -> Any:
        by_id = {item.get("id"): item for item in task.get("nodes", [])}
        queue_ids = list(reversed(node.get("dependencies") or []))
        seen_nodes: set[str] = set()
        while queue_ids:
            dep_id = queue_ids.pop(0)
            if dep_id == "approval" or dep_id in seen_nodes:
                continue
            seen_nodes.add(dep_id)
            dep = by_id.get(dep_id) or {}
            found = self._find_nested_key(dep.get("outputs") or {}, keys)
            if found is not _MISSING:
                return found
            queue_ids.extend(reversed(dep.get("dependencies") or []))
        return _MISSING

    @classmethod
    def _find_nested_key(cls, value: Any, keys: set[str]) -> Any:
        if isinstance(value, dict):
            for key in keys:
                candidate = value.get(key, _MISSING)
                if candidate is not _MISSING and candidate not in (None, ""):
                    return candidate
            preferred = ["selected_device", "device", "execution_result", "result", "data", "output", "outputs"]
            for key in preferred:
                if key in value:
                    found = cls._find_nested_key(value[key], keys)
                    if found is not _MISSING:
                        return found
            for child in value.values():
                found = cls._find_nested_key(child, keys)
                if found is not _MISSING:
                    return found
        elif isinstance(value, list):
            for child in value:
                found = cls._find_nested_key(child, keys)
                if found is not _MISSING:
                    return found
        return _MISSING

    def _merge_safe_autonomous_subplan(self, task: dict[str, Any], supplied: dict[str, Any]) -> dict[str, Any]:
        """Allow task-specific collection phases without expanding the approved envelope."""
        base = self._autonomous_tool_plan(task)
        merged = {**base, **_deepcopy_json(supplied)}
        numeric_bounds = {
            "freq_start_mhz": (float(base["freq_start_mhz"]), None, "min"),
            "freq_stop_mhz": (None, float(base["freq_stop_mhz"]), "max"),
            "freq_step_mhz": (float(base["freq_step_mhz"]), None, "min"),
            "dwell_time_sec": (0.0, float(base["dwell_time_sec"]), "range"),
            "repeat_count": (1.0, float(base["repeat_count"]), "range"),
            "sample_rate": (0.0, float(base["sample_rate"]), "range"),
            "bandwidth": (0.0, float(base["bandwidth"]), "range"),
            "gain": (0.0, float(base["gain"]), "range"),
            "expected_fft_frames_per_capture": (1.0, float(base["expected_fft_frames_per_capture"]), "range"),
        }
        for key, (lower, upper, _) in numeric_bounds.items():
            try:
                value = float(merged[key])
            except (KeyError, TypeError, ValueError) as exc:
                raise RuntimeError(f"子计划参数 {key} 无效") from exc
            if lower is not None and value < lower:
                raise RuntimeError(f"子计划参数 {key}={value} 超出已批准下限 {lower}")
            if upper is not None and value > upper:
                raise RuntimeError(f"子计划参数 {key}={value} 超出已批准上限 {upper}")
        if float(merged["freq_stop_mhz"]) < float(merged["freq_start_mhz"]):
            raise RuntimeError("子计划终止频率不能低于起始频率")
        if float(merged["bandwidth"]) > float(merged["sample_rate"]):
            raise RuntimeError("子计划带宽不能大于采样率")
        estimated_count = int(math.floor((float(merged["freq_stop_mhz"]) - float(merged["freq_start_mhz"])) / float(merged["freq_step_mhz"]))) + 1
        if estimated_count > int(task.get("validation", {}).get("frequency_count") or 50000):
            raise RuntimeError("子计划频点数量超过已批准计划")
        merged["repeat_count"] = int(float(merged["repeat_count"]))
        merged["expected_fft_frames_per_capture"] = int(float(merged["expected_fft_frames_per_capture"]))
        for key in ("capture_agent_task_id", "output_dir", "strict_llm", "plan_version", "plan_fingerprint", "llm_plan", "task_type", "data_source", "antenna"):
            merged[key] = base.get(key)
        return merged

    _PLAN_PLACEHOLDER_RE = re.compile(r"\$\{([^{}]+)\}")

    def _resolve_plan_value(self, value: Any, task: dict[str, Any]) -> Any:
        if isinstance(value, dict):
            return {key: self._resolve_plan_value(item, task) for key, item in value.items()}
        if isinstance(value, list):
            return [self._resolve_plan_value(item, task) for item in value]
        if not isinstance(value, str):
            return value
        full = self._PLAN_PLACEHOLDER_RE.fullmatch(value.strip())
        if full:
            return _deepcopy_json(self._lookup_plan_reference(full.group(1), task))

        def interpolate(match: re.Match[str]) -> str:
            resolved = self._lookup_plan_reference(match.group(1), task)
            if resolved is None:
                return ""
            if isinstance(resolved, (dict, list)):
                return json.dumps(resolved, ensure_ascii=False, separators=(",", ":"), default=str)
            return str(resolved)

        result = self._PLAN_PLACEHOLDER_RE.sub(interpolate, value)
        if "${" in result:
            raise RuntimeError(f"参数中仍包含未闭合占位符：{result}")
        return result

    def _lookup_plan_reference(self, expression: str, task: dict[str, Any]) -> Any:
        normalized = self.planner._normalize_placeholder_expression(expression)
        value_expr, separator, default_expr = normalized.partition(":-")
        try:
            tokens = self._reference_tokens(value_expr)
            if not tokens:
                raise KeyError("empty")
            if tokens[0] == "runtime" and len(tokens) >= 2 and tokens[1] == "autonomous_task_plan":
                current: Any = self._autonomous_tool_plan(task)
                tokens = tokens[2:]
            else:
                roots: dict[str, Any] = {
                    "task": task,
                    "constraints": task.get("constraints", {}),
                    "generated_code": task.get("generated_code", ""),
                    "intent": task.get("intent_analysis", {}),
                    "nodes": {node["id"]: node for node in task.get("nodes", [])},
                    # Keep this root lazy: building the autonomous plan requires
                    # a fully initialized capture task and must not break unrelated
                    # placeholders such as nodes.* or constraints.*.
                    "runtime": {"output_dir": task.get("output_dir")},
                }
                if tokens[0] not in roots:
                    raise KeyError(tokens[0])
                current = roots[tokens[0]]
                tokens = tokens[1:]
            for token in tokens:
                current = self._traverse_reference(current, token)
            return current
        except (KeyError, IndexError, TypeError, ValueError) as exc:
            if separator:
                return self._parse_placeholder_default(default_expr)
            available = self._reference_debug_keys(value_expr, task)
            suffix = f"；当前可用字段：{', '.join(available[:12])}" if available else ""
            raise RuntimeError(f"无法解析计划占位符：${{{expression}}}{suffix}") from exc

    @staticmethod
    def _reference_tokens(expression: str) -> list[str]:
        expr = str(expression or "").strip()
        tokens: list[str] = []
        buffer = ""
        index = 0
        while index < len(expr):
            char = expr[index]
            if char == ".":
                if buffer:
                    tokens.append(buffer)
                    buffer = ""
                index += 1
                continue
            if char == "[":
                if buffer:
                    tokens.append(buffer)
                    buffer = ""
                end = expr.find("]", index + 1)
                if end < 0:
                    raise ValueError("unclosed bracket")
                raw = expr[index + 1 : end].strip()
                if (raw.startswith("'") and raw.endswith("'")) or (raw.startswith('"') and raw.endswith('"')):
                    raw = raw[1:-1]
                if not raw:
                    raise ValueError("empty bracket")
                tokens.append(raw)
                index = end + 1
                continue
            buffer += char
            index += 1
        if buffer:
            tokens.append(buffer)
        return [token for token in tokens if token]

    @classmethod
    def _traverse_reference(cls, current: Any, token: str) -> Any:
        if isinstance(current, (list, tuple)):
            if re.fullmatch(r"-?\d+", token):
                return current[int(token)]
            # Allow selecting a field from the first list item that has it.
            for item in current:
                try:
                    return cls._traverse_reference(item, token)
                except (KeyError, IndexError, TypeError, ValueError):
                    continue
            raise KeyError(token)
        if not isinstance(current, dict):
            raise TypeError(token)
        aliases = {
            "output": "outputs",
            "data": "result",
        }
        candidates = [token]
        if token in aliases:
            candidates.append(aliases[token])
        for candidate in candidates:
            if candidate in current:
                return current[candidate]
        # Tool outputs are deliberately available both flat and under result. Be
        # tolerant of either shape so LLM plans do not fail on harmless wrappers.
        for wrapper in ("outputs", "output", "result", "data", "execution_result"):
            nested = current.get(wrapper, _MISSING)
            if nested is _MISSING or nested is current:
                continue
            try:
                return cls._traverse_reference(nested, token)
            except (KeyError, IndexError, TypeError, ValueError):
                continue
        raise KeyError(token)

    @staticmethod
    def _parse_placeholder_default(text: str) -> Any:
        value = str(text or "").strip()
        if value == "":
            return ""
        try:
            return json.loads(value)
        except json.JSONDecodeError:
            return value

    def _reference_debug_keys(self, expression: str, task: dict[str, Any]) -> list[str]:
        try:
            tokens = self._reference_tokens(expression)
            if not tokens:
                return []
            roots = {
                "task": task,
                "constraints": task.get("constraints", {}),
                "generated_code": task.get("generated_code", ""),
                "intent": task.get("intent_analysis", {}),
                "nodes": {node["id"]: node for node in task.get("nodes", [])},
                "runtime": {"output_dir": task.get("output_dir")},
            }
            current: Any = roots.get(tokens[0], {})
            for token in tokens[1:-1]:
                current = self._traverse_reference(current, token)
            if isinstance(current, dict):
                return [str(key) for key in current.keys()]
            if isinstance(current, list):
                return [str(index) for index in range(min(len(current), 12))]
        except Exception:
            return []
        return []

    @classmethod
    def _assert_no_unresolved_placeholders(cls, value: Any, tool_name: str) -> None:
        if isinstance(value, dict):
            for child in value.values():
                cls._assert_no_unresolved_placeholders(child, tool_name)
        elif isinstance(value, list):
            for child in value:
                cls._assert_no_unresolved_placeholders(child, tool_name)
        elif isinstance(value, str) and ("${" in value or cls._PLAN_PLACEHOLDER_RE.search(value)):
            raise RuntimeError(f"工具 {tool_name} 参数仍包含未解析占位符：{value}")

    @staticmethod
    def _discover_output_files(value: Any) -> list[Path]:
        found: list[Path] = []
        def walk(item: Any) -> None:
            if isinstance(item, dict):
                for child in item.values():
                    walk(child)
            elif isinstance(item, list):
                for child in item:
                    walk(child)
            elif isinstance(item, str):
                candidate = Path(item)
                if candidate.is_absolute() and candidate.exists() and candidate.is_file():
                    resolved = candidate.resolve()
                    if resolved not in found:
                        found.append(resolved)
        walk(value)
        return found

    def _persist_task_bundle_locked(self, task: dict[str, Any]) -> None:
        output_dir = Path(str(task.get("output_dir") or (AUTONOMOUS_OUTPUT_ROOT / task["id"]))).resolve()
        output_dir.mkdir(parents=True, exist_ok=True)
        json_path = output_dir / f"{task['id']}_report.json"
        md_path = output_dir / f"{task['id']}_report.md"
        md_path.write_text(self._render_markdown_report(task), encoding="utf-8")
        # 先落盘再登记，确保审计产物的 size_bytes 不是 0。
        json_path.write_text(json.dumps(task, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
        self._add_artifact_locked(task, json_path, "audit_json")
        self._add_artifact_locked(task, md_path, "audit_markdown")
        code = str(task.get("generated_code") or "")
        if code:
            code_path = output_dir / "generated_task.py"
            code_path.write_text(code, encoding="utf-8")
            self._add_artifact_locked(task, code_path, "generated_code")
        # 再写一次，把刚登记的产物清单纳入最终审计快照，并刷新实际文件大小。
        json_path.write_text(json.dumps(task, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
        self._add_artifact_locked(task, json_path, "audit_json")

    # ---------- node implementations ----------
    def _analyze_templates(self, task_id: str) -> dict[str, Any]:
        with self._lock:
            task = self._load_task_locked(task_id)
            template_ids = list(task.get("template_ids") or [])
        if not template_ids:
            return {"template_count": 0, "message": "本任务未上传模板，将仅依据自然语言指令生成计划。", "combined_markdown": ""}
        items = []
        combined = []
        for template_id in template_ids:
            template = self.get_template(template_id)
            items.append(
                {
                    "id": template_id,
                    "file_name": template["file_name"],
                    "headings": template.get("headings", []),
                    "table_count": len(template.get("tables", [])),
                    "character_count": len(template.get("markdown", "")),
                }
            )
            combined.append(f"# 模板：{template['file_name']}\n{template.get('markdown', '')}")
        return {"template_count": len(items), "items": items, "combined_markdown": "\n\n".join(combined)[:50000]}

    def _discover_collection_devices(self, task_id: str, *, refresh_usrp: bool = True) -> dict[str, Any]:
        """Discover real acquisition devices before the LLM chooses a workflow."""
        usrp_payload: dict[str, Any] = {}
        probe_devices: list[dict[str, Any]] = []
        errors: list[dict[str, str]] = []
        try:
            usrp_payload = dict((self.platform.scan_devices() if refresh_usrp else self.platform.list_devices()) or {})
        except Exception as exc:
            errors.append({"device_type": "usrp", "message": str(exc)})
            usrp_payload = {"devices": []}

        try:
            adapter = self.platform.app.runtime.device_registry.find_by_capability("collect_wifi_bluetooth_probe_evidence")
            client = getattr(adapter, "client", None) if adapter is not None else None
            if client is None or not callable(getattr(client, "list_devices", None)):
                raise RuntimeError("WiFi/蓝牙探针适配器未注册或不支持设备列表")
            probe_devices = [dict(item or {}) for item in client.list_devices()]
        except Exception as exc:
            errors.append({"device_type": "wifi_bluetooth_probe", "message": str(exc)})

        usrp_devices: list[dict[str, Any]] = []
        for raw in list(usrp_payload.get("devices") or []):
            item = dict(raw or {})
            status = str(item.get("status") or "UNKNOWN").upper()
            item["device_type"] = "usrp"
            item["online"] = status not in {"OFFLINE", "ERROR", "UNAVAILABLE"}
            item["available_for_new_task"] = status in {"IDLE", "ONLINE", "READY"}
            usrp_devices.append(item)

        normalized_probes: list[dict[str, Any]] = []
        for raw in probe_devices:
            item = dict(raw or {})
            status = str(item.get("status") or "ONLINE").upper()
            item["device_type"] = "wifi_bluetooth_probe"
            item["online"] = status not in {"OFFLINE", "ERROR", "UNAVAILABLE"}
            item["available_for_new_task"] = bool(item["online"])
            normalized_probes.append(item)

        inventory = {
            "scanned_at": utc_now_iso(),
            "usrp": {
                "devices": usrp_devices,
                "online_count": sum(1 for item in usrp_devices if item.get("online")),
                "available_count": sum(1 for item in usrp_devices if item.get("available_for_new_task")),
                "error": usrp_payload.get("scan_error") or usrp_payload.get("error"),
            },
            "wifi_bluetooth_probe": {
                "devices": normalized_probes,
                "online_count": sum(1 for item in normalized_probes if item.get("online")),
                "available_count": sum(1 for item in normalized_probes if item.get("available_for_new_task")),
            },
            "errors": errors,
        }
        inventory["available_device_types"] = [
            device_type
            for device_type in ("usrp", "wifi_bluetooth_probe")
            if int((inventory.get(device_type) or {}).get("available_count") or 0) > 0
        ]
        with self._lock:
            current = self._load_task_locked(task_id)
            current["available_devices"] = inventory
            self._save_task_locked(current)
        return inventory

    def _select_collection_devices(self, task_id: str) -> dict[str, Any]:
        with self._lock:
            task = self._load_task_locked(task_id)
            template_markdown = str((self._node(task, "template_analysis").get("outputs") or {}).get("combined_markdown") or "")
            inventory = _deepcopy_json(task.get("available_devices") or {})
            instruction = str(task.get("instruction") or "")
            reasoning_mode = _normalize_reasoning_mode(task.get("reasoning_mode"))
        available_types = set(inventory.get("available_device_types") or [])
        if not available_types:
            errors = inventory.get("errors") or []
            raise RuntimeError(f"当前没有可用于新采集任务的设备：{errors or 'USRP/探针均未在线或忙碌'}")
        result = self.planner.select_devices(
            instruction=instruction,
            template_markdown=template_markdown,
            available_devices=inventory,
            reasoning_mode=reasoning_mode,
            stream_handler=self._planner_stream_handler(task_id, "device_selection"),
            cancel_checker=lambda: self._should_stop(task_id),
        )
        proposal = dict(result.get("proposal") or {})
        selected = [str(item) for item in proposal.get("selected_device_types") or []]
        unavailable = [item for item in selected if item not in available_types]
        if unavailable:
            raise RuntimeError(f"设备规划器选择了当前不可用设备：{', '.join(unavailable)}")
        if not selected:
            raise RuntimeError("设备规划器未选择任何可用采集设备")
        if "wifi_bluetooth_probe" in selected and not proposal.get("probe_kinds"):
            proposal["probe_kinds"] = ["wifi_ap", "wifi_client", "bluetooth"]
        selected_usrp = [item for item in (inventory.get("usrp") or {}).get("devices", []) if item.get("available_for_new_task")] if "usrp" in selected else []
        selected_probes = [item for item in (inventory.get("wifi_bluetooth_probe") or {}).get("devices", []) if item.get("available_for_new_task")] if "wifi_bluetooth_probe" in selected else []
        with self._lock:
            current = self._load_task_locked(task_id)
            current["device_selection"] = proposal
            current["selected_usrp_devices"] = _deepcopy_json(selected_usrp)
            current["selected_probe_devices"] = _deepcopy_json(selected_probes)
            current["planner"] = result.get("model") or self.planner.runtime_info()
            self._append_event_locked(
                current,
                "capture_devices_selected",
                "智能体已根据任务背景与在线设备自主选择：" + "、".join(selected),
                node_id="device_selection",
                data={"selected_device_types": selected, "rationale": proposal.get("rationale") or {}},
            )
            self._save_task_locked(current)
        return {
            "device_selection": proposal,
            "available_device_types": sorted(available_types),
            "selected_usrp_devices": selected_usrp,
            "selected_probe_devices": selected_probes,
            "model": result.get("model"),
            "raw_model_output": result.get("raw_model_output"),
        }

    def _analyze_instruction(self, task_id: str) -> dict[str, Any]:
        with self._lock:
            task = self._load_task_locked(task_id)
            template_output = dict(self._node(task, "template_analysis").get("outputs") or {})
            template_markdown = str(template_output.get("combined_markdown") or "")
            previous_analysis = dict(task.get("intent_analysis") or {})

        selected_types = set((task.get("device_selection") or {}).get("selected_device_types") or [])
        if "usrp" not in selected_types:
            selection = dict(task.get("device_selection") or {})
            proposal = {
                "objective": selection.get("objective") or task["instruction"],
                "task_summary": selection.get("background_summary") or task["instruction"],
                "constraints": {},
                "assumptions": [],
                "unresolved_questions": list(selection.get("unavailable_requirements") or []),
                "source_trace": {},
                "risk_notes": [],
                "device_selection": selection,
            }
            with self._lock:
                current = self._load_task_locked(task_id)
                current["constraints"] = {}
                current["constraint_sources"] = {}
                current["validation"] = {"device_mode": "probe_only", "frequency_count": 0, "estimated_capture_seconds": 0}
                current["intent_analysis"] = proposal
                current["parameter_resolution_policy"] = "未选择 USRP，不生成无关频谱采集参数；探针参数来自设备选择节点。"
                self._save_task_locked(current)
            return {
                "intent_analysis": proposal,
                "constraints": {},
                "constraint_sources": {},
                "validation": current["validation"],
                "parameter_resolution_policy": current["parameter_resolution_policy"],
                "sources": ["用户指令", "DOCX 模板", "设备选择结果"],
            }

        user_explicit = CapturePlanBuilder.extract_constraints(task["instruction"], "")
        template_explicit = CapturePlanBuilder.extract_constraints("", template_markdown)
        parameter_context = {
            "policy": (
                "禁止任何默认采集参数。用户明确值优先于模板明确值；模板明确值优先于意图设计。"
                "仅对用户和模板均未指定的字段，根据本次任务意图设计合适参数。"
            ),
            "required_fields": list(CapturePlanBuilder.REQUIRED_CONSTRAINT_KEYS),
            "usrp_device_parameter_ranges": self.planner.usrp_parameter_context()["ranges"],
            "usrp_parameter_relationships": self.planner.usrp_parameter_context()["relationships"],
            "user_explicit_constraints": user_explicit,
            "template_explicit_constraints": template_explicit,
            "template_uploaded": bool(template_markdown.strip()),
        }
        result = self.planner.analyze_intent(
            instruction=task["instruction"],
            template_markdown=template_markdown,
            parameter_context=parameter_context,
            previous_analysis=previous_analysis,
            reasoning_mode=_normalize_reasoning_mode(task.get("reasoning_mode")),
            stream_handler=self._planner_stream_handler(task_id, "instruction_analysis"),
            cancel_checker=lambda: self._should_stop(task_id),
        )
        proposal = dict(result.get("proposal") or {})
        model_constraints = dict(proposal.get("constraints") or {})
        model_source_trace = dict(proposal.get("source_trace") or {})
        constraints: dict[str, Any] = {}
        sources: dict[str, str] = {}
        for key in CapturePlanBuilder.REQUIRED_CONSTRAINT_KEYS:
            if key in user_explicit:
                constraints[key] = user_explicit[key]
                sources[key] = "user_instruction"
            elif key in template_explicit:
                constraints[key] = template_explicit[key]
                sources[key] = "template"
            elif model_constraints.get(key) is not None:
                constraints[key] = model_constraints[key]
                trace = str(model_source_trace.get(key) or "").lower()
                if "user" in trace or "用户" in trace:
                    sources[key] = "user_instruction"
                elif "template" in trace or "模板" in trace:
                    sources[key] = "template"
                else:
                    sources[key] = "llm_intent_design"
            else:
                raise ValueError(
                    f"意图分析器未设计必需采集参数 {key}；系统禁止使用默认值或固定预设补齐"
                )

        constraints = CapturePlanBuilder.normalize_constraints(constraints)
        validation = CapturePlanBuilder.validate_constraints(constraints)
        proposal["constraints"] = dict(constraints)
        proposal["source_trace"] = dict(sources)
        with self._lock:
            current = self._load_task_locked(task_id)
            current["constraints"] = constraints
            current["constraint_sources"] = sources
            current["validation"] = validation
            current["intent_analysis"] = proposal
            current["parameter_resolution_policy"] = parameter_context["policy"]
            current["planner"] = result.get("model") or self.planner.runtime_info()
            self._save_task_locked(current)
        return {
            "intent_analysis": proposal,
            "constraints": constraints,
            "constraint_sources": sources,
            "validation": validation,
            "parameter_resolution_policy": parameter_context["policy"],
            "model": result.get("model"),
            "raw_model_output": result.get("raw_model_output"),
            "sources": ["用户明确参数", "DOCX 模板明确参数", "真实 LLM 按任务意图设计的其余参数"],
        }

    def _probe_collection_nodes(self, task: dict[str, Any], dependency: str = "approval") -> list[dict[str, Any]]:
        selection = dict(task.get("device_selection") or {})
        probes = list(task.get("selected_probe_devices") or [])
        probe_ids = [str(item.get("probe_id") or "").strip() for item in probes if str(item.get("probe_id") or "").strip()]
        if not probe_ids:
            raise RuntimeError("设备规划选择了 WiFi/蓝牙探针，但当前没有可用 probe_id")
        kinds = [item for item in selection.get("probe_kinds") or [] if item in {"wifi_ap", "wifi_client", "bluetooth"}]
        if not kinds:
            kinds = ["wifi_ap", "wifi_client", "bluetooth"]
        now = utc_now_iso()
        return [{
            "id": "wifi_bluetooth_probe_collection",
            "title": "采集 WiFi/蓝牙探针信息",
            "description": "调用在线 WiFi/蓝牙探针获取热点、WiFi 客户端和蓝牙设备线索；完整结果先落 JSON + SQLite 索引，节点上下文只保留有界摘要，后续可按需检索。",
            "executor": "tool_call",
            "kind": "tool_call",
            "dependencies": [dependency],
            "tool_name": "collect_wifi_bluetooth_probe_evidence",
            "tool_arguments": {
                "probe_ids": probe_ids,
                "kinds": kinds,
                "active_minutes": int(selection.get("active_minutes") or 30),
                "page_size": 200,
                "capture_task_id": str(task.get("id") or ""),
            },
            "allowed_tool": "collect_wifi_bluetooth_probe_evidence",
            "allowed_tools": ["collect_wifi_bluetooth_probe_evidence", "query_probe_evidence"],
            "status": "pending",
            "progress": 0,
            "attempts": 0,
            "started_at": None,
            "ended_at": None,
            "updated_at": now,
            "summary": "等待执行",
            "success_criteria": [
                "探针接口调用完成并返回有界证据摘要",
                "完整探针结果已落库并生成 evidence_set_id",
                "可通过 query_probe_evidence 按需检索详细设备记录",
            ],
            "expected_evidence": ["WiFi 热点", "WiFi 客户端", "蓝牙设备", "探针证据集索引"],
            "tool_input_overrides": {},
            "risk_level": "low",
            "planner_origin": "adaptive_device_selection",
            "inputs": {},
            "outputs": {},
            "logs": [],
            "notes": "",
        }]

    def _report_generation_node(self, dependency: str) -> dict[str, Any]:
        now = utc_now_iso()
        return {
            "id": "report_generation",
            "title": "研判报告生成",
            "description": "检查并整理任务 Markdown 报告，融合本任务实际调用的 USRP、WiFi/蓝牙探针及对应采集结果、证据集与执行状态。",
            "executor": "report_generation",
            "kind": "system_summary",
            "dependencies": [dependency],
            "tool_name": None,
            "tool_arguments": {},
            "allowed_tool": None,
            "allowed_tools": [],
            "status": "pending",
            "progress": 0,
            "attempts": 0,
            "started_at": None,
            "ended_at": None,
            "updated_at": now,
            "summary": "等待生成研判报告",
            "success_criteria": CapturePlanBuilder.success_criteria("report_generation"),
            "expected_evidence": ["Markdown 研判报告", "各设备采集情况汇总"],
            "tool_input_overrides": {},
            "risk_level": "low",
            "planner_origin": "system_report_finalization",
            "inputs": {},
            "outputs": {},
            "logs": [],
            "notes": "",
        }

    @staticmethod
    def _bootstrap_node_ids() -> tuple[str, ...]:
        return ("template_analysis", "device_discovery", "device_selection", "instruction_analysis", "plan_generation", "approval")

    def _generate_plan_summary(self, task_id: str) -> dict[str, Any]:
        with self._lock:
            task = self._load_task_locked(task_id)
            instruction = task["instruction"]
            constraints = dict(task.get("constraints") or {})
            plan_version = int(task.get("plan_version") or 1)
            selection = dict(task.get("device_selection") or {})
            selected_types = [str(item) for item in selection.get("selected_device_types") or []]

        if not selected_types:
            raise RuntimeError("缺少设备选择结果，无法生成采集计划")

        prepare_result: dict[str, Any] | None = None
        usrp_nodes: list[dict[str, Any]] = []
        if "usrp" in selected_types:
            prepare_result = self._prepare_spectrum_collection_plan(task_id, constraints)
            usrp_nodes = self._fixed_collection_nodes(prepare_result)

        with self._lock:
            task_snapshot = self._load_task_locked(task_id)
        probe_nodes = self._probe_collection_nodes(task_snapshot) if "wifi_bluetooth_probe" in selected_types else []

        groups: dict[str, list[dict[str, Any]]] = {
            "usrp": usrp_nodes,
            "wifi_bluetooth_probe": probe_nodes,
        }
        compiled_nodes: list[dict[str, Any]] = []
        previous_last = "approval"
        for device_type in selected_types:
            group = groups.get(device_type) or []
            if not group:
                continue
            first = group[0]
            # 只改变设备组之间的根依赖；USRP 组内部 device_scan -> collection -> check -> summary 完全保持不变。
            original_first_dependencies = list(first.get("dependencies") or [])
            if original_first_dependencies == ["approval"] or first.get("id") == "wifi_bluetooth_probe_collection":
                first["dependencies"] = [previous_last]
            compiled_nodes.extend(group)
            previous_last = str(group[-1].get("id"))

        if not compiled_nodes:
            raise RuntimeError("设备选择结果没有产生任何可执行采集节点")

        # 所有设备采集完成后统一进入显式的“研判报告生成”节点。
        # 该节点只整理已有采集结果与 report.md，不再触发新的真实设备调用。
        compiled_nodes.append(self._report_generation_node(previous_last))

        device_labels = {"usrp": "USRP", "wifi_bluetooth_probe": "WiFi/蓝牙探针"}
        ordered_labels = [device_labels.get(item, item) for item in selected_types]
        completion_contract = [
            "仅执行设备选择节点选中的可用设备，不机械调用未选择设备",
            "所有真实采集节点仅在用户批准后执行",
            "USRP 被选中时保持原标准频谱采集节点内部流程不变",
            "探针被选中时完整设备列表先落库索引，模型只接收有界摘要并支持按需检索",
            "全部设备采集结束后执行研判报告生成，报告融合各设备采集情况并检查 Markdown 文件有效性",
            "最终产物与证据均进入任务审计记录",
        ]
        proposal = {
            "objective": selection.get("objective") or f"按任务背景完成采集：{instruction}",
            "strategy_summary": "规划器先感知在线设备，再由真实 LLM 自主选择设备。当前设备组执行顺序：" + " → ".join(ordered_labels) + "。",
            "completion_contract": completion_contract,
            "adaptive_device_flow": True,
            "device_selection": selection,
            "prepare_result": prepare_result or {},
            "plan_steps_text": self._fixed_flow_text(compiled_nodes),
        }
        fingerprint = self.planner.plan_fingerprint(
            compiled_nodes,
            constraints=constraints,
            plan_version=plan_version,
            objective=proposal["objective"],
            completion_contract=completion_contract,
        )
        compiled_public = [
            {
                "id": node.get("id"),
                "title": node.get("title"),
                "kind": node.get("kind"),
                "executor": node.get("executor"),
                "dependencies": node.get("dependencies", []),
                "tool_name": node.get("tool_name"),
                "tool_arguments": node.get("tool_arguments", {}),
                "allowed_tools": node.get("allowed_tools", []),
                "success_criteria": node.get("success_criteria", []),
                "planner_origin": node.get("planner_origin"),
            }
            for node in compiled_nodes
        ]
        with self._lock:
            current = self._load_task_locked(task_id)
            bootstrap_ids = set(self._bootstrap_node_ids())
            bootstrap_by_id = {node["id"]: node for node in current.get("nodes", []) if node.get("id") in bootstrap_ids}
            ordered_bootstrap = [bootstrap_by_id[node_id] for node_id in self._bootstrap_node_ids() if node_id in bootstrap_by_id]
            current["nodes"] = ordered_bootstrap + compiled_nodes
            current["plan_candidate"] = proposal
            current["plan_fingerprint"] = fingerprint
            current["approved_plan_fingerprint"] = None
            current["planner"] = self.planner.runtime_info()
            history_entry = {
                "plan_version": plan_version,
                "created_at": utc_now_iso(),
                "model": current["planner"],
                "proposal": proposal,
                "compiled_nodes": compiled_public,
                "plan_fingerprint": fingerprint,
                "raw_model_output": "",
            }
            current.setdefault("plan_history", []).append(history_entry)
            self._append_event_locked(
                current,
                "adaptive_device_plan_compiled",
                f"已根据 LLM 设备选择生成 {len(compiled_nodes)} 个执行节点：{' → '.join(ordered_labels)}，等待用户批准。",
                node_id="plan_generation",
                data={"plan_version": plan_version, "plan_fingerprint": fingerprint, "selected_device_types": selected_types},
            )
            self._save_task_locked(current)
        return {
            "objective": proposal["objective"],
            "strategy_summary": proposal["strategy_summary"],
            "plan_version": plan_version,
            "locked_after_approval": True,
            "plan_fingerprint": fingerprint,
            "model": self.planner.runtime_info(),
            "device_selection": selection,
            "compiled_nodes": compiled_public,
            "validator_summary": {
                "candidate_node_count": len(compiled_nodes),
                "executable_node_count": len(compiled_nodes),
                "dag_validated": True,
                "tool_references_validated": True,
                "approval_gate_bound_to_roots": True,
                "workflow_template_used": False,
                "workflow_mode": "adaptive device groups / stable device-internal nodes",
            },
            "completion_contract": completion_contract,
            "prepare_result": prepare_result or {},
            "plan_steps_text": proposal["plan_steps_text"],
            "raw_model_output": "",
        }

    @staticmethod
    def _local_prepare_spectrum_collection(args: dict[str, Any]) -> dict[str, Any]:
        required = {
            "room_name",
            "mode",
            "freq_start_mhz",
            "freq_stop_mhz",
            "freq_step_mhz",
            "dwell_ms",
            "repeat_count",
            "aggregation",
            "sample_rate",
            "bandwidth",
            "gain",
            "antenna",
        }
        missing = sorted(
            key for key in required
            if key not in args or args.get(key) is None or args.get(key) == ""
        )
        if missing:
            raise ValueError("本地采集计划缺少显式参数，禁止默认补齐：" + "、".join(missing))
        freq_start = float(args["freq_start_mhz"])
        freq_stop = float(args["freq_stop_mhz"])
        freq_step = float(args["freq_step_mhz"])
        repeat_count = int(args["repeat_count"])
        dwell_ms = float(args["dwell_ms"])
        freq_count = max(1, int(math.floor((freq_stop - freq_start) / freq_step)) + 1)
        return {
            "plan_id": new_id("scplan"),
            "room_name": str(args["room_name"]),
            "mode": str(args["mode"]),
            "freq_start_mhz": freq_start,
            "freq_stop_mhz": freq_stop,
            "freq_step_mhz": freq_step,
            "dwell_ms": dwell_ms,
            "repeat_count": repeat_count,
            "aggregation": str(args["aggregation"]),
            "sample_rate": float(args["sample_rate"]),
            "bandwidth": float(args["bandwidth"]),
            "gain": float(args["gain"]),
            "antenna": str(args["antenna"]),
            "freq_count": freq_count,
            "estimated_total_sec": round(freq_count * repeat_count * dwell_ms / 1000.0, 3),
            "device_id": args.get("device_id"),
            "device_status": {
                "status": "not_queried",
                "message": "当前运行环境没有完整平台上下文，仅校验本任务已解析参数。",
            },
            "dry_run": bool(args.get("dry_run", False)),
            "confirmation_prompt": "请确认本任务按意图解析的采集参数无误后执行。",
        }

    def _prepare_spectrum_collection_plan(self, task_id: str, constraints: dict[str, Any]) -> dict[str, Any]:
        node_id = "plan_generation"
        constraints = CapturePlanBuilder.normalize_constraints(constraints)
        args = {
            "room_name": constraints["room_name"],
            "mode": constraints["mode"],
            "freq_start_mhz": constraints["freq_start_mhz"],
            "freq_stop_mhz": constraints["freq_stop_mhz"],
            "freq_step_mhz": constraints["freq_step_mhz"],
            "dwell_ms": constraints["dwell_time_sec"] * 1000.0,
            "repeat_count": constraints["repeat_count"],
            "aggregation": constraints["aggregation"],
            "sample_rate": constraints["sample_rate"],
            "bandwidth": constraints["bandwidth"],
            "gain": constraints["gain"],
            "antenna": constraints["antenna"],
            "dry_run": False,
        }
        with self._lock:
            task = self._load_task_locked(task_id)
            node = self._node(task, node_id)
            node["logs"].append(self._log_entry("调用 prepare_spectrum_collection 生成待确认采集参数。", "info", args))
            self._append_event_locked(
                task,
                "tool_progress",
                "调用 B 方案 prepare_spectrum_collection 生成采集参数确认单。",
                node_id=node_id,
                data={"tool": "prepare_spectrum_collection", "stage": "prepare", "progress": 35},
            )
            self._save_task_locked(task)

        runtime = getattr(getattr(self.platform, "app", None), "runtime", None)
        platform_task = getattr(self.platform, "task", None)
        if runtime is None or platform_task is None or not hasattr(runtime, "tool_registry"):
            data = self._local_prepare_spectrum_collection(args)
            with self._lock:
                task = self._load_task_locked(task_id)
                node = self._node(task, node_id)
                node["logs"].append(self._log_entry("当前环境缺少完整平台上下文，已使用本地参数确认结果。", "warning", data))
                self._append_event_locked(
                    task,
                    "llm_reasoning_delta",
                    "已完成本地参数预检，等待用户在弹窗中批准执行。\n",
                    node_id=node_id,
                    data={"delta": "已完成本地参数预检，等待用户在弹窗中批准执行。\n", "progress": 70},
                )
                self._save_task_locked(task)
            return data

        run = Run(
            id=new_id("capture_run"),
            task_id=platform_task.id,
            trigger_kind=RunTriggerKind.CHAT,
            trigger_event_id=None,
            trigger_message_id=None,
            agent_profile="capture_agent",
            status=RunStatus.RUNNING,
            step_budget=1,
            step_count=0,
            started_at=utc_now(),
            conversation_id=None,
        )
        message = ChatMessage(
            id=new_id("capture_message"),
            conversation_id="capture-agent",
            task_id=platform_task.id,
            role=ChatRole.OPERATOR,
            content=self.get_task_internal(task_id)["instruction"],
            run_id=run.id,
            created_at=utc_now(),
        )
        context = ToolContext(
            task=platform_task,
            run=run,
            trigger_event=None,
            trigger_message=message,
            state_repo=runtime.state_repo,
            case_repo=runtime.case_repo,
            knowledge_base=runtime.knowledge_base,
            device_registry=runtime.device_registry,
            llm_client=runtime.llm_client,
            stream_handler=None,
            nl2sql_options=NL2SQLSessionConfig(),
            asset_manager=runtime.asset_manager,
            document_index=runtime.document_index,
            database_catalog=runtime.database_catalog,
            cancel_checker=lambda: self._cancel_events.setdefault(task_id, threading.Event()).is_set(),
        )
        result = runtime.tool_registry.execute("prepare_spectrum_collection", args, context)
        if result.result.status != "success":
            raise RuntimeError(result.result.error or "prepare_spectrum_collection 执行失败")
        data = dict(result.result.data or {})
        with self._lock:
            task = self._load_task_locked(task_id)
            node = self._node(task, node_id)
            node["logs"].append(self._log_entry("prepare_spectrum_collection 已返回待确认参数。", "success", data))
            self._append_event_locked(
                task,
                "llm_reasoning_delta",
                "B 方案工具已完成参数预检，等待用户在弹窗中批准执行。\n",
                node_id=node_id,
                data={"delta": "B 方案工具已完成参数预检，等待用户在弹窗中批准执行。\n", "progress": 70},
            )
            self._save_task_locked(task)
        return data

    def _fixed_collection_nodes(self, prepare_result: dict[str, Any]) -> list[dict[str, Any]]:
        now = utc_now_iso()
        plan_id = str(prepare_result.get("plan_id") or "")
        specs = [
            {
                "id": "device_scan",
                "title": "开始：查询可用设备",
                "description": "执行开始前查询当前可用 USRP 设备，选择空闲设备并记录设备状态，作为后续采集的前置条件。",
                "executor": "device_scan",
                "kind": "system_check",
                "dependencies": ["approval"],
                "tool_name": "scan_usrp_devices",
                "tool_arguments": {},
                "allowed_tool": "scan_usrp_devices",
                "allowed_tools": ["scan_usrp_devices"],
            },
            {
                "id": "collection_execution",
                "title": "执行标准频谱采集",
                "description": "固定调用 B 方案 execute_spectrum_collection，按已确认 plan_id 执行全频段/多频点采集并生成 .npz。",
                "executor": "collection_execution",
                "kind": "tool_call",
                "dependencies": ["device_scan"],
                "tool_name": "execute_spectrum_collection",
                "tool_arguments": {"plan_id": plan_id, "operator_confirmed": True},
                "allowed_tool": "execute_spectrum_collection",
                "allowed_tools": ["execute_spectrum_collection"],
            },
            {
                "id": "output_file_check",
                "title": "检查输出文件",
                "description": "只检查输出文件是否存在、是否为 .npz、文件大小是否大于 0。",
                "executor": "output_file_check",
                "kind": "system_check",
                "dependencies": ["collection_execution"],
                "tool_name": None,
                "tool_arguments": {},
                "allowed_tool": None,
                "allowed_tools": [],
            },
            {
                "id": "execution_summary",
                "title": "总结采集任务摘要",
                "description": "汇总采集参数、执行结果、输出文件检查结果和任务产物。",
                "executor": "execution_summary",
                "kind": "system_summary",
                "dependencies": ["output_file_check"],
                "tool_name": None,
                "tool_arguments": {},
                "allowed_tool": None,
                "allowed_tools": [],
            },
        ]
        nodes = []
        for spec in specs:
            nodes.append(
                {
                    **spec,
                    "status": "pending",
                    "progress": 0,
                    "attempts": 0,
                    "started_at": None,
                    "ended_at": None,
                    "updated_at": now,
                    "summary": "等待执行",
                    "success_criteria": CapturePlanBuilder.success_criteria(spec["id"]),
                    "expected_evidence": [],
                    "tool_input_overrides": {},
                    "risk_level": "medium" if spec["id"] == "collection_execution" else "low",
                    "planner_origin": "fixed_b_tool_flow",
                    "inputs": {},
                    "outputs": {},
                    "logs": [],
                    "notes": "",
                }
            )
        return nodes

    @staticmethod
    def _fixed_flow_text(nodes: list[dict[str, Any]]) -> str:
        return " → ".join(str(node.get("title") or node.get("id")) for node in nodes)

    def _execute_knowledge_retrieval(self, task_id: str, node_id: str) -> tuple[dict[str, Any], str]:
        task = self.get_task_internal(task_id)
        node = self._node(task, node_id)
        overrides = dict(node.get("tool_input_overrides") or {})
        query = str(overrides.get("query") or f"{task['instruction']} WebSocket FFT configure start stop 参数限制")
        top_k = max(1, min(12, int(overrides.get("top_k") or 6)))
        data = self._execute_tool(task_id, node_id, "retrieve_usrp_api_knowledge", {"query": query, "top_k": top_k})
        return data, f"已检索 {len(data.get('items') or [])} 个 USRP API 知识片段。"

    def _execute_device_scan(self, task_id: str, node_id: str) -> tuple[dict[str, Any], str]:
        task = self.get_task_internal(task_id)
        try:
            payload = self._execute_tool(task_id, node_id, "scan_usrp_devices", {})
        except Exception as exc:
            payload = {"devices": [], "scan_error": str(exc)}
        devices = list(payload.get("devices") or [])
        selected = next((item for item in devices if str(item.get("status") or "").upper() == "IDLE"), None)
        if selected is None:
            reason = payload.get("scan_error") or "未发现 IDLE 的可用 USRP 设备"
            return {**payload, "selected_device": None, "query_status": "no_idle_device"}, f"可用设备查询完成：{reason}；后续采集工具将继续执行自身设备检查。"
        return {**payload, "selected_device": selected, "query_status": "idle_device_selected"}, f"已选择空闲设备 {selected.get('dev_id')}。"

    def _execute_device_status(self, task_id: str, node_id: str) -> tuple[dict[str, Any], str]:
        task = self.get_task_internal(task_id)
        selected = dict((self._node_for_executor(task, "device_scan").get("outputs") or {}).get("selected_device") or {})
        payload = self._execute_tool(task_id, node_id, "list_usrp_devices", {})
        devices = list(payload.get("devices") or [])
        matched = next((item for item in devices if item.get("dev_id") == selected.get("dev_id")), None)
        if matched is None:
            raise RuntimeError("设备状态复核时未找到已选择设备")
        if str(matched.get("status") or "").upper() != "IDLE":
            raise RuntimeError(f"设备 {matched.get('dev_id')} 当前状态不是 IDLE")
        return {**payload, "selected_device": matched}, f"设备 {matched.get('dev_id')} 状态复核通过。"

    def _execute_parameter_validation(self, task_id: str, node_id: str) -> tuple[dict[str, Any], str]:
        task = self.get_task_internal(task_id)
        constraints = dict(task.get("constraints") or {})
        validation = CapturePlanBuilder.validate_constraints(constraints)
        device_output = self._node_for_executor(task, "device_scan").get("outputs") or {}
        selected = dict(device_output.get("selected_device") or {})
        config = dict(selected.get("dev_config") or {})
        freq_range = dict(config.get("freq_range") or {})
        min_freq = freq_range.get("min")
        max_freq = freq_range.get("max")
        if min_freq is not None and constraints["freq_start_mhz"] < _device_mhz_value(min_freq):
            raise RuntimeError(f"起始频率低于设备下限 {_human_frequency(_device_mhz_value(min_freq))}")
        if max_freq is not None and constraints["freq_stop_mhz"] > _device_mhz_value(max_freq):
            raise RuntimeError(f"终止频率高于设备上限 {_human_frequency(_device_mhz_value(max_freq))}")
        return {"status": "passed", "constraints": constraints, "validation": validation, "device_id": selected.get("dev_id")}, f"参数安全校验通过，共 {validation['frequency_count']} 个频点。"

    def _execute_device_configuration(self, task_id: str, node_id: str) -> tuple[dict[str, Any], str]:
        task = self.get_task_internal(task_id)
        constraints = dict(task["constraints"])
        selected = dict((self._node_for_executor(task, "device_scan").get("outputs") or {}).get("selected_device") or {})
        params = {
            "dev_id": selected.get("dev_id"),
            "freq": constraints["freq_start_mhz"],
            "sample_rate": constraints["sample_rate"],
            "bandwidth": constraints["bandwidth"],
            "gain": constraints["gain"],
            "slice_duration": constraints["dwell_time_sec"],
            "duration": constraints["dwell_time_sec"],
            "antenna": constraints["antenna"],
        }
        result = self._execute_tool(task_id, node_id, "configure_usrp_capture", params)
        return {"status": "configured", "params": params, "result": result}, "设备采集参数配置成功。"

    def _execute_code_generation(self, task_id: str, node_id: str) -> tuple[dict[str, Any], str]:
        task = self.get_task_internal(task_id)
        tool_plan = self._autonomous_tool_plan(task)
        data = self._execute_tool(
            task_id,
            node_id,
            "generate_usrp_task_code",
            {"task_description": self._tool_task_description(task), "task_plan": tool_plan},
        )
        code = str(data.get("code") or "")
        if not code:
            raise RuntimeError("采集代码生成器未返回代码")
        with self._lock:
            current = self._load_task_locked(task_id)
            current["generated_code"] = code
            self._save_task_locked(current)
        validation = dict(data.get("validation") or {})
        return data, f"真实 LLM 已生成受控 run_task(ctx)，安全检查状态：{validation.get('status', 'unknown')}。"

    def _execute_collection_execution(self, task_id: str, node_id: str) -> tuple[dict[str, Any], str]:
        task = self.get_task_internal(task_id)
        code = str(task.get("generated_code") or "")
        if not code:
            raise RuntimeError("缺少已验证的采集代码")
        data = self._execute_tool(
            task_id,
            node_id,
            "execute_usrp_task_code",
            {
                "code": code,
                "task_description": self._tool_task_description(task),
                "task_plan": self._autonomous_tool_plan(task),
            },
        )
        execution = dict(data.get("execution_result") or {})
        if execution.get("status") != "completed":
            raise RuntimeError(f"采集执行未完成：{execution.get('status') or 'unknown'}")
        with self._lock:
            current = self._load_task_locked(task_id)
            current["execution_result"] = execution
            output_file = execution.get("output_file")
            if output_file:
                self._add_artifact_locked(current, Path(str(output_file)), "collection_result")
            self._save_task_locked(current)
        return data, f"采集执行完成：{execution.get('freq_count', '-')} 个频点，重复 {execution.get('repeat_count', '-')} 次。"

    def _execute_result_verification(self, task_id: str, node_id: str) -> tuple[dict[str, Any], str]:
        task = self.get_task_internal(task_id)
        result = dict(task.get("execution_result") or {})
        expected = dict(task.get("validation") or {})
        errors = []
        if result.get("status") != "completed":
            errors.append("执行状态不是 completed")
        if int(result.get("freq_count") or -1) != int(expected.get("frequency_count") or -2):
            errors.append(f"频点数不匹配：实际 {result.get('freq_count')} / 计划 {expected.get('frequency_count')}")
        if int(result.get("repeat_count") or -1) != int(task["constraints"].get("repeat_count") or -2):
            errors.append("重复次数与计划不一致")
        output_file = str(result.get("output_file") or "")
        output_exists = bool(output_file and Path(output_file).exists())
        if not output_exists:
            errors.append("输出文件不存在")
        if errors:
            raise RuntimeError("；".join(errors))
        verification = {
            "status": "passed",
            "output_file": output_file,
            "output_exists": output_exists,
            "frequency_count": result.get("freq_count"),
            "repeat_count": result.get("repeat_count"),
            "data_source": result.get("data_source"),
            "approved_plan_fingerprint": task.get("approved_plan_fingerprint"),
        }
        return verification, "采集结果与批准的 LLM 计划一致，验收通过。"

    def _execute_artifact_save(self, task_id: str, node_id: str) -> tuple[dict[str, Any], str]:
        with self._lock:
            task = self._load_task_locked(task_id)
            report_dir = self.reports_dir / task_id
            report_dir.mkdir(parents=True, exist_ok=True)
            json_path = report_dir / f"{task_id}_report.json"
            md_path = report_dir / f"{task_id}_report.md"
            json_path.write_text(json.dumps(task, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
            md_path.write_text(self._render_markdown_report(task), encoding="utf-8")
            json_artifact = self._add_artifact_locked(task, json_path, "audit_json")
            md_artifact = self._add_artifact_locked(task, md_path, "audit_markdown")
            self._save_task_locked(task)
        return {"report_json": json_artifact, "report_markdown": md_artifact}, "任务审计报告已保存。"

    def _execute_output_file_check(self, task_id: str, node_id: str) -> tuple[dict[str, Any], str]:
        task = self.get_task_internal(task_id)
        result = dict(task.get("execution_result") or {})
        output_file = str(result.get("output_file") or "")
        path = Path(output_file) if output_file else Path()
        checks = {
            "exists": bool(output_file and path.exists() and path.is_file()),
            "is_npz": bool(output_file and path.suffix.lower() == ".npz"),
            "size_gt_zero": bool(output_file and path.exists() and path.is_file() and path.stat().st_size > 0),
        }
        if not all(checks.values()):
            failed = "；".join(key for key, passed in checks.items() if not passed)
            raise RuntimeError(f"输出文件检查失败：{failed}")
        size_bytes = path.stat().st_size
        verification = {
            "status": "passed",
            "output_file": str(path.resolve()),
            "file_name": path.name,
            "size_bytes": size_bytes,
            "checks": checks,
        }
        with self._lock:
            current = self._load_task_locked(task_id)
            self._add_artifact_locked(current, path, "collection_result")
            self._save_task_locked(current)
        return verification, f"输出文件检查通过：存在、.npz、大小 {size_bytes} B。"

    def _execute_execution_summary(self, task_id: str, node_id: str) -> tuple[dict[str, Any], str]:
        with self._lock:
            task = self._load_task_locked(task_id)
            prepare_result = dict((task.get("plan_candidate") or {}).get("prepare_result") or {})
            execution_result = dict(task.get("execution_result") or {})
            output_check = dict(self._node(task, "output_file_check").get("outputs") or {})
            summary = {
                "task_id": task["id"],
                "plan_version": task.get("plan_version"),
                "room_name": prepare_result.get("room_name") or execution_result.get("room_name"),
                "plan_id": prepare_result.get("plan_id") or execution_result.get("plan_id"),
                "freq_count": execution_result.get("freq_count") or prepare_result.get("freq_count"),
                "success_count": execution_result.get("success_count"),
                "failed_count": execution_result.get("failed_count", 0),
                "repeat_count": execution_result.get("repeat_count") or prepare_result.get("repeat_count"),
                "aggregation": execution_result.get("aggregation") or prepare_result.get("aggregation"),
                "device_id": execution_result.get("device_id") or prepare_result.get("device_id"),
                "output_file": execution_result.get("output_file") or output_check.get("output_file"),
                "output_check": output_check.get("checks") or {},
                "status": execution_result.get("status") or "completed",
            }
            task["execution_summary"] = summary
            task["messages"].append(
                {
                    "id": f"msg_{uuid4().hex[:10]}",
                    "role": "assistant",
                    "content": (
                        f"采集任务执行摘要：会议室/区域 {summary.get('room_name') or '-'}，"
                        f"计划 {summary.get('plan_id') or '-'}，输出文件 {Path(str(summary.get('output_file') or '')).name or '-'}。"
                    ),
                    "created_at": utc_now_iso(),
                    "template_ids": [],
                }
            )
            self._save_task_locked(task)
        return summary, "采集任务执行摘要已生成。"

    def _execute_report_generation(self, task_id: str, node_id: str) -> tuple[dict[str, Any], str]:
        with self._lock:
            task = self._load_task_locked(task_id)
            output_dir = Path(str(task.get("output_dir") or (AUTONOMOUS_OUTPUT_ROOT / task["id"]))).resolve()
            output_dir.mkdir(parents=True, exist_ok=True)
            report_path = output_dir / f"{task['id']}_report.md"
            report_text = self._render_markdown_report(task)
            report_path.write_text(report_text, encoding="utf-8")
            if not report_path.exists() or report_path.stat().st_size <= 0:
                raise RuntimeError("研判报告生成后检查失败：Markdown 报告文件不存在或为空")
            device_snapshot = self._device_collection_snapshot(task)
            artifact = self._add_artifact_locked(task, report_path, "audit_markdown")
            result = {
                "status": "passed",
                "report_file": str(report_path),
                "file_name": report_path.name,
                "size_bytes": report_path.stat().st_size,
                "device_collection": device_snapshot,
                "artifact_id": artifact.get("id"),
            }
            task["report_generation"] = result
            self._save_task_locked(task)
        return result, f"研判报告已生成并检查通过：{report_path.name}，已融合各设备采集情况。"

    # ---------- tool bridge ----------
    def _execute_tool(self, task_id: str, node_id: str, tool_name: str, args: dict[str, Any]) -> dict[str, Any]:
        with self._lock:
            task = self._load_task_locked(task_id)
            self._assert_plan_integrity_locked(task)
            node = self._node(task, node_id)
            allowed_tools = list(node.get("allowed_tools") or ([] if not node.get("allowed_tool") else [node.get("allowed_tool")]))
            if tool_name not in allowed_tools:
                raise RuntimeError(f"计划约束拒绝工具调用：节点 {node_id} 仅允许 {allowed_tools or ['无']}，收到 {tool_name}")
            task = _deepcopy_json(task)
        runtime = getattr(getattr(self.platform, "app", None), "runtime", None)
        platform_task = getattr(self.platform, "task", None)
        if runtime is None or platform_task is None or not hasattr(runtime, "tool_registry"):
            raise RuntimeError(f"当前平台上下文不支持执行工具 {tool_name}")
        run = Run(
            id=new_id("capture_run"),
            task_id=platform_task.id,
            trigger_kind=RunTriggerKind.CHAT,
            trigger_event_id=None,
            trigger_message_id=None,
            agent_profile="capture_agent",
            status=RunStatus.RUNNING,
            step_budget=1,
            step_count=0,
            started_at=utc_now(),
            conversation_id=None,
        )
        message = ChatMessage(
            id=new_id("capture_message"),
            conversation_id="capture-agent",
            task_id=platform_task.id,
            role=ChatRole.OPERATOR,
            content=task["instruction"],
            run_id=run.id,
            created_at=utc_now(),
        )

        stream_state = {
            "reasoning_buffer": "",
            "content_buffer": "",
            "last_reasoning_flush": 0.0,
            "last_content_flush": 0.0,
        }

        def stream_handler(event_type: str, payload: dict[str, Any]) -> None:
            payload = dict(payload or {})
            nested = payload.get("data")
            data = dict(nested) if isinstance(nested, dict) else {}
            tool_event_name = str(payload.get("tool_name") or data.get("tool_name") or tool_name)
            stage = str(payload.get("stage") or data.get("stage") or event_type)
            delta = str(payload.get("delta") or data.get("delta") or "")
            message_text = str(payload.get("message") or data.get("message") or "")
            now = time.monotonic()

            # LLM 的 reasoning/content 流需要写入节点输出，并转发成前端能识别的
            # llm_* 事件。旧实现把它们当作普通工具日志，导致右侧面板一直为空。
            if event_type == "reasoning_delta":
                stream_state["reasoning_buffer"] += delta
                if len(stream_state["reasoning_buffer"]) < 80 and now - float(stream_state["last_reasoning_flush"]) < 0.12:
                    return
            elif event_type == "content_delta":
                stream_state["content_buffer"] += delta
                if len(stream_state["content_buffer"]) < 160 and now - float(stream_state["last_content_flush"]) < 0.18:
                    return

            summary_event_name = ""
            with self._lock:
                current = self._load_task_locked(task_id)
                current_node = self._node(current, node_id)
                outputs = current_node.setdefault("outputs", {})
                event_name = "tool_progress"
                event_payload: dict[str, Any] = {"tool": tool_name, "tool_event": tool_event_name, "stage": stage}
                log_message = ""

                if event_type == "reasoning_start":
                    outputs.setdefault("llm_reasoning", "")
                    current_node["summary"] = "工具内智能体正在分析采集任务……"
                    current_node["progress"] = max(int(current_node.get("progress") or 0), 5)
                    event_name = "llm_reasoning_start"
                    message_text = message_text or "工具内模型开始输出实时推理。"
                elif event_type == "reasoning_delta":
                    chunk = stream_state["reasoning_buffer"]
                    stream_state["reasoning_buffer"] = ""
                    stream_state["last_reasoning_flush"] = now
                    outputs["llm_reasoning"] = (str(outputs.get("llm_reasoning") or "") + chunk)[-60000:]
                    current_node["summary"] = "工具内智能体正在生成采集函数……"
                    current_node["progress"] = max(int(current_node.get("progress") or 0), min(35, 7 + len(str(outputs["llm_reasoning"])) // 180))
                    event_name = "llm_reasoning_delta"
                    event_payload["delta"] = chunk
                    message_text = "工具内模型推理内容已更新。"
                elif event_type == "reasoning_done":
                    if stream_state["reasoning_buffer"]:
                        chunk = stream_state["reasoning_buffer"]
                        outputs["llm_reasoning"] = (str(outputs.get("llm_reasoning") or "") + chunk)[-60000:]
                        stream_state["reasoning_buffer"] = ""
                        event_payload["delta"] = chunk
                    event_name = "llm_reasoning_done"
                    message_text = message_text or "工具内模型推理阶段完成。"
                elif event_type == "content_start":
                    outputs.setdefault("llm_live_output", "")
                    event_name = "llm_content_start"
                    message_text = message_text or "工具内模型开始生成采集函数。"
                elif event_type == "content_delta":
                    chunk = stream_state["content_buffer"]
                    stream_state["content_buffer"] = ""
                    stream_state["last_content_flush"] = now
                    outputs["llm_live_output"] = (str(outputs.get("llm_live_output") or "") + chunk)[-80000:]
                    current_node["summary"] = "工具内智能体正在流式生成采集函数……"
                    current_node["progress"] = max(int(current_node.get("progress") or 0), min(45, 18 + len(str(outputs["llm_live_output"])) // 300))
                    event_name = "llm_content_delta"
                    event_payload["delta"] = chunk
                    message_text = "采集函数输出流已更新。"
                elif event_type == "content_done":
                    if stream_state["content_buffer"]:
                        chunk = stream_state["content_buffer"]
                        outputs["llm_live_output"] = (str(outputs.get("llm_live_output") or "") + chunk)[-80000:]
                        stream_state["content_buffer"] = ""
                        event_payload["delta"] = chunk
                    event_name = "llm_content_done"
                    message_text = message_text or "工具内模型代码生成完成。"
                else:
                    # SafeExecutionLogger 以及 run_autonomous_usrp_task 的阶段事件
                    # 作为可审计执行轨迹同步到右侧面板，避免模型不返回 reasoning
                    # 时出现空白区域。
                    if event_type == "tool_trace":
                        message_text = message_text or f"正在执行 {stage} 阶段。"
                    elif tool_event_name == "autonomous_usrp_progress":
                        message_text = message_text or "自主采集进度已更新。"
                    elif event_type == "tool_result" and tool_event_name == "generate_usrp_task_code":
                        message_text = "采集函数与静态校验结果已生成。"
                    elif event_type == "tool_result" and tool_event_name == "run_autonomous_usrp_task":
                        message_text = "端到端自主采集工具已返回最终结果。"
                    else:
                        message_text = message_text or f"工具事件：{event_type}"

                    trace_line = f"[{stage}] {message_text}\n"
                    outputs["llm_reasoning"] = (str(outputs.get("llm_reasoning") or "") + trace_line)[-60000:]
                    event_name = "llm_reasoning_delta"
                    event_payload["delta"] = trace_line
                    log_message = message_text

                # 将实时 FFT 连接信息透传给前端；仅保留轻量元数据，避免把 FFT 数组写入任务事件。
                detail_data = data.get("data") if isinstance(data.get("data"), dict) else {}
                if stage == "fft":
                    for key in ("ws_url", "task_id", "freq", "freq_mhz", "frame_count", "fft_size"):
                        if key in detail_data and detail_data.get(key) is not None:
                            event_payload[key] = detail_data.get(key)
                    legacy_freq = detail_data.get("freq_" + "hz")
                    if legacy_freq is not None and "freq_mhz" not in event_payload:
                        try:
                            event_payload["freq_mhz"] = float(legacy_freq) / 1e6
                        except Exception:
                            event_payload["freq_mhz"] = legacy_freq

                progress_data = data if data else payload
                progress = self._progress_from_tool_event(stage, progress_data)
                if progress is not None:
                    current_node["progress"] = max(int(current_node.get("progress") or 0), progress)
                event_payload["progress"] = current_node.get("progress", 0)
                current_node["updated_at"] = utc_now_iso()

                if log_message:
                    current_node["logs"].append(
                        self._log_entry(log_message, "info", {"stage": stage, "tool_event": tool_event_name})
                    )
                self._append_event_locked(current, event_name, message_text, node_id=node_id, data=event_payload)
                self._save_task_locked(current)
                summary_event_name = event_name

            if summary_event_name in {
                "llm_reasoning_delta",
                "llm_reasoning_done",
                "llm_content_delta",
                "llm_content_done",
            }:
                self._schedule_status_summary(
                    task_id,
                    node_id,
                    force=summary_event_name in {"llm_reasoning_done", "llm_content_done"},
                    phase=summary_event_name,
                )

        context = ToolContext(
            task=self.platform.task,
            run=run,
            trigger_event=None,
            trigger_message=message,
            state_repo=runtime.state_repo,
            case_repo=runtime.case_repo,
            knowledge_base=runtime.knowledge_base,
            device_registry=runtime.device_registry,
            llm_client=runtime.llm_client,
            stream_handler=stream_handler,
            nl2sql_options=NL2SQLSessionConfig(),
            asset_manager=runtime.asset_manager,
            document_index=runtime.document_index,
            database_catalog=runtime.database_catalog,
            cancel_checker=lambda: self._cancel_events.setdefault(task_id, threading.Event()).is_set(),
        )
        result = runtime.tool_registry.execute(tool_name, args, context)
        if result.result.status != "success":
            raise RuntimeError(result.result.error or f"工具 {tool_name} 执行失败")
        return dict(result.result.data or {})

    @staticmethod
    def _progress_from_tool_event(stage: str, data: dict[str, Any]) -> int | None:
        explicit = data.get("progress")
        if explicit is None and isinstance(data.get("data"), dict):
            explicit = data["data"].get("progress")
        if explicit is not None:
            try:
                return max(1, min(95, int(float(explicit))))
            except Exception:
                pass
        stages = {
            "initializing": 3,
            "knowledge": 8,
            "code_generation": 16,
            "validation": 24,
            "validate": 26,
            "sandbox": 30,
            "execution": 38,
            "progress": 65,
            "scan": 42,
            "plan": 45,
            "configure": 50,
            "start": 55,
            "capture": 65,
            "fft": 78,
            "save": 90,
            "artifact_relocated": 92,
            "completed": 95,
            "error": 95,
        }
        return stages.get(stage)

    # ---------- helpers ----------
    def _tool_task_description(self, task: dict[str, Any]) -> str:
        constraints = task["constraints"]
        plan_steps = "；".join(
            f"{index + 1}.{node.get('title')}"
            for index, node in enumerate(self._execution_nodes_in_order(task))
        )
        return (
            f"{task['instruction']}\n"
            f"必须严格使用已批准的真实 LLM 计划 v{task['plan_version']}（指纹 {task.get('approved_plan_fingerprint') or task.get('plan_fingerprint')}）。"
            f"采集约束：{_human_frequency(constraints['freq_start_mhz'])} 至 {_human_frequency(constraints['freq_stop_mhz'])}，"
            f"步长 {_human_frequency(constraints['freq_step_mhz'])}，每个频点重复 {constraints['repeat_count']} 次，"
            f"驻留 {constraints['dwell_time_sec']} 秒。不得扩大频率范围、增加重复次数或修改数据源。\n"
            f"批准节点顺序：{plan_steps}"
        )

    def _autonomous_tool_plan(self, task: dict[str, Any]) -> dict[str, Any]:
        constraints = dict(task["constraints"])
        selected: dict[str, Any] = {}
        for candidate in task.get("nodes", []):
            outputs = candidate.get("outputs") or {}
            if not isinstance(outputs, dict):
                continue
            possible = outputs.get("selected_device")
            if not isinstance(possible, dict) and isinstance(outputs.get("result"), dict):
                possible = outputs["result"].get("selected_device")
            if isinstance(possible, dict) and possible.get("dev_id"):
                selected = dict(possible)
                break
        return {
            "task_type": "llm_planned_capture_agent_usrp_fft_stream",
            "data_source": "usrp_websocket_fft",
            "room_name": constraints["room_name"],
            "mode": constraints["mode"],
            "freq_start_mhz": constraints["freq_start_mhz"],
            "freq_stop_mhz": constraints["freq_stop_mhz"],
            "freq_step_mhz": constraints["freq_step_mhz"],
            "dwell_time_sec": constraints["dwell_time_sec"],
            "repeat_count": constraints["repeat_count"],
            "aggregation": constraints["aggregation"],
            "sample_rate": constraints["sample_rate"],
            "bandwidth": constraints["bandwidth"],
            "gain": constraints["gain"],
            "antenna": constraints["antenna"],
            "expected_fft_frames_per_capture": constraints["expected_fft_frames_per_capture"],
            "dev_id": selected.get("dev_id"),
            "capture_agent_task_id": task["id"],
            "output_dir": task.get("output_dir"),
            "strict_llm": True,
            "reasoning_mode": _normalize_reasoning_mode(task.get("reasoning_mode")),
            "plan_version": task["plan_version"],
            "plan_fingerprint": task.get("approved_plan_fingerprint") or task.get("plan_fingerprint"),
            "llm_plan": [
                {
                    "id": node.get("id"),
                    "title": node.get("title"),
                    "kind": node.get("kind") or node.get("executor"),
                    "tool_name": node.get("tool_name") or node.get("allowed_tool"),
                    "dependencies": node.get("dependencies", []),
                    "success_criteria": node.get("success_criteria", []),
                }
                for node in self._execution_nodes_in_order(task)
            ],
        }

    def _node_inputs(self, task: dict[str, Any], node_id: str) -> dict[str, Any]:
        node = self._node(task, node_id)
        executor = str(node.get("executor") or node_id)
        constraints = dict(task.get("constraints") or {})
        dependency_outputs = {
            dep: (self._node(task, dep).get("outputs") or {})
            for dep in node.get("dependencies", [])
            if dep != "approval"
        }
        base = {
            "plan_version": task.get("plan_version"),
            "plan_fingerprint": task.get("approved_plan_fingerprint") or task.get("plan_fingerprint"),
            "planner": task.get("planner"),
            "kind": node.get("kind") or executor,
            "executor": executor,
            "tool_name": node.get("tool_name") or node.get("allowed_tool"),
            "tool_arguments": node.get("tool_arguments", {}),
            "dependencies": node.get("dependencies", []),
            "dependency_outputs": dependency_outputs,
            "output_dir": task.get("output_dir"),
        }
        if node.get("kind") == "tool_call":
            try:
                resolved = self._resolve_plan_value(dict(node.get("tool_arguments") or {}), task)
            except Exception as exc:
                resolved = {"resolution_error": str(exc), "raw": node.get("tool_arguments", {})}
            return {**base, "resolved_tool_arguments": resolved}
        if node.get("kind") in {"llm_analysis", "llm_verification", "decision"}:
            return {**base, "instruction": task.get("instruction"), "constraints": constraints}
        if node_id == "template_analysis":
            return {**base, "template_ids": task.get("template_ids", []), "templates": task.get("templates", [])}
        if node_id == "instruction_analysis":
            return {**base, "instruction": task.get("instruction"), "template_count": len(task.get("template_ids", [])), "requires_real_llm": True}
        if node_id == "plan_generation":
            return {
                **base,
                "instruction": task.get("instruction"),
                "intent_analysis": task.get("intent_analysis", {}),
                "constraints": constraints,
                "fixed_tool": "prepare_spectrum_collection",
            }
        if node_id == "approval":
            return {
                **base,
                "requires_explicit_approval": True,
                "prepare_result": (task.get("plan_candidate") or {}).get("prepare_result", {}),
                "plan_steps_text": (task.get("plan_candidate") or {}).get("plan_steps_text", ""),
            }
        if executor == "output_file_check":
            return {**base, "execution_result": task.get("execution_result", {}), "checks": ["exists", "is_npz", "size_gt_zero"]}
        if executor == "execution_summary":
            return {
                **base,
                "prepare_result": (task.get("plan_candidate") or {}).get("prepare_result", {}),
                "execution_result": task.get("execution_result", {}),
                "output_file_check": self._node(task, "output_file_check").get("outputs", {}),
            }
        if executor == "knowledge_retrieval":
            return {**base, "instruction": task.get("instruction"), "data_source": "usrp_websocket_fft"}
        if executor == "device_scan":
            return {**base, "device_concurrency_policy": "one_active_task_per_device"}
        selected_device = {}
        try:
            selected_device = dict((self._node_for_executor(task, "device_scan").get("outputs") or {}).get("selected_device") or {})
        except KeyError:
            pass
        if executor == "device_status":
            return {**base, "selected_device": selected_device}
        if executor == "parameter_validation":
            return {**base, "constraints": constraints, "selected_device": selected_device}
        if executor == "device_configuration":
            return {**base, "constraints": constraints, "selected_device": selected_device}
        if executor == "code_generation":
            return {**base, "task_description": self._tool_task_description(task), "task_plan": self._autonomous_tool_plan(task)}
        if executor == "collection_execution":
            return {**base, "task_plan": self._autonomous_tool_plan(task), "generated_code_length": len(str(task.get("generated_code") or ""))}
        if executor == "result_verification":
            return {**base, "expected": task.get("validation", {}), "execution_result": task.get("execution_result", {})}
        if executor == "artifact_save":
            return {**base, "task_id": task.get("id"), "existing_artifact_count": len(task.get("artifacts", []))}
        return base

    def _planner_stream_handler(self, task_id: str, node_id: str) -> Callable[[str, dict[str, Any]], None]:
        stream_state = {
            "content_buffer": "",
            "reasoning_buffer": "",
            "last_content_flush": 0.0,
            "last_reasoning_flush": 0.0,
        }

        def handler(event_type: str, payload: dict[str, Any]) -> None:
            now = time.monotonic()
            delta = str(payload.get("delta") or payload.get("text") or "")
            if event_type == "content_delta":
                stream_state["content_buffer"] += delta
                if len(stream_state["content_buffer"]) < 120 and now - float(stream_state["last_content_flush"]) < 0.18:
                    return
            elif event_type == "reasoning_delta":
                stream_state["reasoning_buffer"] += delta
                if len(stream_state["reasoning_buffer"]) < 80 and now - float(stream_state["last_reasoning_flush"]) < 0.12:
                    return

            should_refresh_summary = False
            force_summary = False
            with self._lock:
                try:
                    task = self._load_task_locked(task_id)
                    node = self._node(task, node_id)
                except KeyError:
                    return
                outputs = node.setdefault("outputs", {})
                event_data = {k: v for k, v in payload.items() if k not in {"delta", "text"}}

                if event_type == "content_delta":
                    chunk = stream_state["content_buffer"]
                    live = str(outputs.get("llm_live_output") or "") + chunk
                    outputs["llm_live_output"] = live[-50000:]
                    stream_state["content_buffer"] = ""
                    stream_state["last_content_flush"] = now
                    node["summary"] = "模型正在生成结构化结果……"
                    node["progress"] = max(int(node.get("progress") or 0), min(90, 12 + len(live) // 100))
                    message = "模型正在流式生成结构化结果。"
                    event_name = "llm_content_delta"
                    event_data["delta"] = chunk
                    should_refresh_summary = True
                elif event_type == "reasoning_start":
                    outputs.setdefault("llm_reasoning", "")
                    node["summary"] = "模型正在分析任务与依赖关系……"
                    node["progress"] = max(int(node.get("progress") or 0), 6)
                    message = "模型开始输出可审计推理过程。"
                    event_name = "llm_reasoning_start"
                elif event_type == "reasoning_delta":
                    chunk = stream_state["reasoning_buffer"]
                    reasoning = str(outputs.get("llm_reasoning") or "") + chunk
                    outputs["llm_reasoning"] = reasoning[-50000:]
                    stream_state["reasoning_buffer"] = ""
                    stream_state["last_reasoning_flush"] = now
                    node["summary"] = "模型正在分析任务与依赖关系……"
                    node["progress"] = max(int(node.get("progress") or 0), min(88, 8 + len(reasoning) // 120))
                    message = "模型推理内容已更新。"
                    event_name = "llm_reasoning_delta"
                    event_data["delta"] = chunk
                    should_refresh_summary = True
                elif event_type == "reasoning_done":
                    if stream_state["reasoning_buffer"]:
                        chunk = stream_state["reasoning_buffer"]
                        outputs["llm_reasoning"] = (str(outputs.get("llm_reasoning") or "") + chunk)[-50000:]
                        stream_state["reasoning_buffer"] = ""
                        event_data["delta"] = chunk
                    message = "模型推理阶段完成。"
                    event_name = "llm_reasoning_done"
                    should_refresh_summary = True
                    force_summary = True
                elif event_type == "llm_json_repair":
                    message = "模型首次输出未通过结构校验，正在自动修复 JSON。"
                    event_name = "llm_json_repair"
                    node["logs"].append(self._log_entry(message, "warning", {"error": payload.get("error")}))
                elif event_type == "llm_request_started":
                    message = f"调用模型 {payload.get('model') or task.get('planner', {}).get('model')}（第 {payload.get('attempt', 1)} 次）。"
                    event_name = "llm_request_started"
                    node["logs"].append(self._log_entry(message, "info"))
                elif event_type == "llm_request_completed":
                    if stream_state["content_buffer"]:
                        outputs["llm_live_output"] = (str(outputs.get("llm_live_output") or "") + stream_state["content_buffer"])[-50000:]
                        stream_state["content_buffer"] = ""
                    if stream_state["reasoning_buffer"]:
                        outputs["llm_reasoning"] = (str(outputs.get("llm_reasoning") or "") + stream_state["reasoning_buffer"])[-50000:]
                        stream_state["reasoning_buffer"] = ""
                    node["progress"] = max(int(node.get("progress") or 0), 95)
                    message = "模型输出完成，正在校验结构、占位符、工具参数和依赖 DAG。"
                    event_name = "llm_request_completed"
                    node["logs"].append(self._log_entry(message, "success"))
                    should_refresh_summary = True
                    force_summary = True
                else:
                    return
                node["updated_at"] = utc_now_iso()
                self._append_event_locked(task, event_name, message, node_id=node_id, data=event_data)
                self._save_task_locked(task)

            if should_refresh_summary:
                self._schedule_status_summary(task_id, node_id, force=force_summary, phase=event_name)

        return handler

    def _schedule_status_summary(self, task_id: str, node_id: str, *, force: bool = False, phase: str = "") -> None:
        """Queue lifecycle summaries and coalesced live summaries for one node.

        Lifecycle summaries are requested exactly once for node start and node end.
        Stream-driven live summaries are throttled, SHA1-deduplicated, and coalesced
        so a busy node keeps only the newest pending snapshot while one worker runs.
        """
        phase_name = str(phase or "").strip().lower()
        if phase_name == "node_started":
            summary_kind = "start"
            lifecycle_phase = "start"
            model_phase = "node_started"
        elif phase_name in {"node_terminal", "node_completed", "node_failed", "node_cancelled"}:
            summary_kind = "end"
            lifecycle_phase = "end"
            model_phase = "node_terminal"
        else:
            summary_kind = "live"
            lifecycle_phase = ""
            model_phase = phase_name or "live"

        try:
            with self._lock:
                task = self._load_task_locked(task_id)
                node = self._node(task, node_id)
                outputs = node.setdefault("outputs", {})

                if lifecycle_phase:
                    requested = [str(item) for item in outputs.get("llm_status_summary_requested_phases") or []]
                    if lifecycle_phase in requested:
                        return
                    requested.append(lifecycle_phase)
                    outputs["llm_status_summary_requested_phases"] = requested

                reasoning = str(outputs.get("llm_reasoning") or "")
                structured = str(outputs.get("llm_live_output") or "")
                if summary_kind == "end" and not structured:
                    output_keys = [
                        str(key)
                        for key in outputs.keys()
                        if not str(key).startswith("llm_")
                    ][:24]
                    structured = "final_output_keys=" + ", ".join(output_keys)

                source = (reasoning[-2400:] + "\n" + structured[-800:]).strip()
                if summary_kind == "live" and not source:
                    return

                snapshot = {
                    "task_context": {
                        "id": task.get("id"),
                        "instruction": task.get("instruction"),
                        "status": task.get("status"),
                        "progress": task.get("progress"),
                    },
                    "node_context": {
                        "id": node.get("id"),
                        "title": node.get("title"),
                        "description": node.get("description"),
                        "summary": node.get("summary"),
                        "status": node.get("status"),
                        "progress": node.get("progress"),
                        "phase": model_phase,
                        "started_at": node.get("started_at"),
                        "ended_at": node.get("ended_at"),
                    },
                    "latest_reasoning": reasoning,
                    "latest_structured_output": structured,
                    "_summary_kind": summary_kind,
                    "_lifecycle_phase": lifecycle_phase,
                }
                if lifecycle_phase:
                    self._save_task_locked(task)
        except KeyError:
            return

        key = (task_id, node_id)
        now = time.monotonic()
        signature_source = f"{summary_kind}\n{source}"
        signature = hashlib.sha1(signature_source.encode("utf-8", errors="ignore")).hexdigest()

        with self._summary_lock:
            last_requested = float(self._summary_last_requested.get(key) or 0.0)
            if summary_kind == "live" and not force and now - last_requested < 1.4:
                return
            if signature == self._summary_signatures.get(key):
                return

            self._summary_signatures[key] = signature
            self._summary_last_requested[key] = now
            snapshot["_summary_signature"] = signature

            worker = self._summary_workers.get(key)
            if worker is not None and worker.is_alive():
                # Coalesce aggressively: while a summary is running, keep only the
                # newest state. The worker consumes that latest snapshot next.
                self._summary_pending[key] = snapshot
                return

            # The first snapshot is passed directly to the worker so the one-off
            # node-start summary cannot be overwritten by an immediate stream update.
            worker = threading.Thread(
                target=self._status_summary_worker,
                args=(key, snapshot),
                daemon=True,
                name=f"capture-summary-{task_id[-5:]}-{node_id[-8:]}",
            )
            self._summary_workers[key] = worker
            worker.start()

    def _status_summary_worker(
        self,
        key: tuple[str, str],
        initial_snapshot: dict[str, Any] | None = None,
    ) -> None:
        task_id, node_id = key
        snapshot = initial_snapshot
        try:
            while True:
                if snapshot is None:
                    with self._summary_lock:
                        snapshot = self._summary_pending.pop(key, None)
                if snapshot is None:
                    return

                summary_kind = str(snapshot.pop("_summary_kind", "live") or "live")
                lifecycle_phase = str(snapshot.pop("_lifecycle_phase", "") or "")
                snapshot_signature = str(snapshot.pop("_summary_signature", "") or "")

                try:
                    summary = self.planner.summarize_live_status(**snapshot)
                except Exception as exc:
                    with self._summary_lock:
                        if snapshot_signature and self._summary_signatures.get(key) == snapshot_signature:
                            self._summary_signatures.pop(key, None)
                    with self._lock:
                        try:
                            task = self._load_task_locked(task_id)
                            node = self._node(task, node_id)
                            label = (
                                f"节点{lifecycle_phase}摘要"
                                if lifecycle_phase
                                else "实时状态摘要"
                            )
                            node.setdefault("logs", []).append(
                                self._log_entry(f"{label}生成失败：{exc}", "warning")
                            )
                            self._save_task_locked(task)
                        except KeyError:
                            pass
                    snapshot = None
                    continue

                if summary:
                    with self._lock:
                        try:
                            task = self._load_task_locked(task_id)
                            node = self._node(task, node_id)
                        except KeyError:
                            return
                        outputs = node.setdefault("outputs", {})

                        if lifecycle_phase:
                            outputs[f"llm_status_summary_{lifecycle_phase}"] = summary
                            generated = [
                                str(item)
                                for item in outputs.get("llm_status_summary_generated_phases") or []
                            ]
                            if lifecycle_phase not in generated:
                                generated.append(lifecycle_phase)
                            outputs["llm_status_summary_generated_phases"] = generated

                        # A fast node may terminate while its start/live summary is
                        # still being generated. Never let a delayed non-terminal
                        # response overwrite the node's final summary.
                        is_terminal = str(node.get("status") or "") in TERMINAL_NODE_STATES
                        visible_summary = summary_kind == "end" or not is_terminal
                        if visible_summary and str(outputs.get("llm_status_summary") or "") != summary:
                            outputs["llm_status_summary"] = summary
                            outputs["llm_status_summary_updated_at"] = utc_now_iso()
                            outputs["llm_status_summary_phase"] = lifecycle_phase or "live"
                            node["updated_at"] = utc_now_iso()
                            self._append_event_locked(
                                task,
                                "llm_status_summary",
                                summary,
                                node_id=node_id,
                                data={
                                    "summary": summary,
                                    "phase": lifecycle_phase or "live",
                                    "source_phase": snapshot.get("node_context", {}).get("phase"),
                                },
                            )
                        self._save_task_locked(task)

                snapshot = None
        finally:
            with self._summary_lock:
                self._summary_workers.pop(key, None)
                if key in self._summary_pending:
                    worker = threading.Thread(
                        target=self._status_summary_worker,
                        args=(key,),
                        daemon=True,
                        name=f"capture-summary-{task_id[-5:]}-{node_id[-8:]}",
                    )
                    self._summary_workers[key] = worker
                    worker.start()

    @staticmethod
    def _execution_nodes_in_order(task: dict[str, Any]) -> list[dict[str, Any]]:
        bootstrap_ids = set(CaptureAgentService._bootstrap_node_ids())
        return [node for node in task.get("nodes", []) if node.get("id") not in bootstrap_ids]

    @staticmethod
    def _node_for_executor(task: dict[str, Any], executor: str) -> dict[str, Any]:
        node = next((item for item in task.get("nodes", []) if item.get("executor") == executor), None)
        if node is None:
            raise KeyError(executor)
        return node

    def _preflight_plan_locked(self, task: dict[str, Any]) -> None:
        """Validate every generated node before asking the operator to approve it."""
        issues: list[str] = []
        for node in self._execution_nodes_in_order(task):
            if node.get("kind") != "tool_call":
                continue
            tool_name = str(node.get("tool_name") or node.get("allowed_tool") or "")
            if tool_name not in self._tool_definitions:
                issues.append(f"节点 {node.get('id')} 引用了不可用工具 {tool_name}")
                continue
            raw_args = self._inject_capture_tool_context(tool_name, dict(node.get("tool_arguments") or {}), task)
            # Static placeholders must resolve now. Dynamic node-output placeholders are
            # validated structurally by CaptureLLMPlanner and resolved when ancestors finish.
            for expression, location in self.planner._iter_placeholders(raw_args):
                tokens = self.planner._reference_tokens(expression)
                if tokens and tokens[0] == "nodes":
                    continue
                try:
                    self._lookup_plan_reference(expression, task)
                except RuntimeError as exc:
                    issues.append(f"节点 {node.get('id')} / {location}: {exc}")
            schema = dict(self._tool_definitions[tool_name].input_schema or {})
            required = set(schema.get("required") or [])
            supplied = set(raw_args)
            for key in sorted(required - supplied):
                if self._infer_tool_argument(tool_name, key, task, node) is _MISSING:
                    issues.append(f"节点 {node.get('id')} 的工具 {tool_name} 缺少必填参数 {key}")
        if issues:
            raise RuntimeError("LLM 计划预检失败：" + "；".join(issues[:12]))

    def _task_plan_fingerprint(self, task: dict[str, Any]) -> str:
        proposal = dict(task.get("plan_candidate") or {})
        return self.planner.plan_fingerprint(
            task.get("nodes", []),
            constraints=dict(task.get("constraints") or {}),
            plan_version=int(task.get("plan_version") or 1),
            objective=str(proposal.get("objective") or task.get("instruction") or ""),
            completion_contract=list(proposal.get("completion_contract") or []),
        )

    def _assert_plan_integrity_locked(self, task: dict[str, Any]) -> None:
        expected = str(task.get("approved_plan_fingerprint") or "")
        if not expected:
            raise RuntimeError("批准计划缺少锁定指纹")
        actual = self._task_plan_fingerprint(task)
        if actual != expected:
            raise RuntimeError("批准后的计划结构发生变化，执行器已拒绝继续")

    def _complete_node(self, task_id: str, node_id: str, summary: str, outputs: dict[str, Any]) -> None:
        with self._lock:
            task = self._load_task_locked(task_id)
            node = self._node(task, node_id)
            existing_outputs = dict(node.get("outputs") or {})
            streamed_outputs = {
                key: value
                for key, value in existing_outputs.items()
                if key.startswith("llm_")
            }
            node.update(
                {
                    "status": "completed",
                    "progress": 100,
                    "ended_at": utc_now_iso(),
                    "updated_at": utc_now_iso(),
                    "summary": summary,
                    # Streamed reasoning/content is produced before the final
                    # structured result. Preserve it instead of replacing the
                    # node output when the node is marked completed.
                    "outputs": {**streamed_outputs, **outputs},
                }
            )
            node["logs"].append(self._log_entry(summary, "success"))
            task["progress"] = self._task_progress(task)
            self._append_event_locked(task, "node_completed", summary, node_id=node_id, data={"progress": task["progress"]})
            self._save_task_locked(task)
        self._schedule_status_summary(task_id, node_id, force=True, phase="node_terminal")

    def _begin_node_locked(self, task: dict[str, Any], node: dict[str, Any], message: str) -> None:
        node["status"] = "in_progress"
        node["progress"] = max(1, int(node.get("progress") or 0))
        node["attempts"] = int(node.get("attempts") or 0) + 1
        node["started_at"] = node.get("started_at") or utc_now_iso()
        node["updated_at"] = utc_now_iso()
        node["summary"] = message
        node["inputs"] = self._node_inputs(task, node["id"])
        node["logs"].append(self._log_entry(message, "info"))
        task["current_node_id"] = node["id"]
        task["status"] = "planning" if node["id"] in {"template_analysis", "device_discovery", "device_selection", "instruction_analysis", "plan_generation"} else "running"
        self._append_event_locked(task, "node_started", message, node_id=node["id"])

    @staticmethod
    def _assert_dependencies(task: dict[str, Any], node: dict[str, Any]) -> None:
        nodes = {item["id"]: item for item in task["nodes"]}
        incomplete = [dep for dep in node.get("dependencies", []) if nodes.get(dep, {}).get("status") != "completed"]
        if incomplete:
            raise RuntimeError(f"节点 {node['id']} 的依赖尚未完成：{', '.join(incomplete)}")

    def _wait_if_paused(self, task_id: str) -> None:
        while self._pause_events.setdefault(task_id, threading.Event()).is_set():
            if self._cancel_events.setdefault(task_id, threading.Event()).is_set():
                return
            time.sleep(0.2)

    def _should_stop(self, task_id: str) -> bool:
        if self._cancel_events.setdefault(task_id, threading.Event()).is_set():
            return True
        with self._lock:
            try:
                return self._load_task_locked(task_id).get("status") == "cancelled"
            except KeyError:
                return True

    @staticmethod
    def _plan_requires_device_access(task: dict[str, Any]) -> bool:
        device_tools = {
            "scan_usrp_devices",
            "list_usrp_devices",
            "configure_usrp_capture",
            "query_usrp_task",
            "execute_usrp_task_code",
            "run_autonomous_usrp_task",
            "execute_spectrum_collection",
            "collect_wifi_bluetooth_probe_evidence",
        }
        return any(
            (node.get("tool_name") or node.get("allowed_tool")) in device_tools
            for node in task.get("nodes", [])
        )

    def _find_conflicting_real_task_locked(self, exclude_task_id: str) -> dict[str, Any] | None:
        for path in self.tasks_dir.glob("*.json"):
            try:
                candidate = json.loads(path.read_text(encoding="utf-8"))
            except Exception:
                continue
            if candidate.get("id") == exclude_task_id:
                continue
            if (
                candidate.get("plan_locked")
                and candidate.get("status") in {"running", "paused"}
                and self._plan_requires_device_access(candidate)
            ):
                return candidate
        return None

    def _recover_interrupted_tasks(self) -> None:
        for path in self.tasks_dir.glob("*.json"):
            try:
                task = json.loads(path.read_text(encoding="utf-8"))
            except Exception:
                continue
            if task.get("status") in {"planning", "running"}:
                task["status"] = "paused"
                task["pause_reason"] = "service_restart"
                task["error"] = "服务重启后任务已安全暂停，请人工继续。"
                current = task.get("current_node_id")
                if current:
                    node = next((item for item in task.get("nodes", []) if item.get("id") == current), None)
                    if node and node.get("status") == "in_progress":
                        node["status"] = "pending"
                        node["summary"] = "服务重启后等待恢复"
                self._append_event_locked(task, "task_recovered", "检测到服务重启，任务已安全暂停。", node_id=current, level="warning")
                self._save_task_locked(task)

    def _load_task_locked(self, task_id: str) -> dict[str, Any]:
        path = self.tasks_dir / f"{task_id}.json"
        if not path.exists():
            raise KeyError(task_id)
        return json.loads(path.read_text(encoding="utf-8"))

    def _save_task_locked(self, task: dict[str, Any]) -> None:
        task["updated_at"] = utc_now_iso()
        self._write_json(self.tasks_dir / f"{task['id']}.json", task)

    @staticmethod
    def _write_json(path: Path, payload: dict[str, Any]) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        temp = path.with_suffix(path.suffix + ".tmp")
        temp.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
        temp.replace(path)

    @staticmethod
    def _node(task: dict[str, Any], node_id: str) -> dict[str, Any]:
        node = next((item for item in task.get("nodes", []) if item.get("id") == node_id), None)
        if node is None:
            raise KeyError(node_id)
        return node

    @staticmethod
    def _log_entry(message: str, level: str = "info", data: dict[str, Any] | None = None) -> dict[str, Any]:
        return {"timestamp": utc_now_iso(), "level": level, "message": message, "data": data or {}}

    def _append_event_locked(
        self,
        task: dict[str, Any],
        event_type: str,
        message: str,
        *,
        node_id: str | None = None,
        level: str = "info",
        data: dict[str, Any] | None = None,
    ) -> None:
        events = task.setdefault("events", [])
        seq = int(events[-1].get("seq") or 0) + 1 if events else 1
        events.append(
            {
                "seq": seq,
                "type": event_type,
                "task_id": task["id"],
                "node_id": node_id,
                "message": message,
                "level": level,
                "data": data or {},
                "created_at": utc_now_iso(),
            }
        )
        if len(events) > 3000:
            del events[:-3000]

    @staticmethod
    def _task_progress(task: dict[str, Any]) -> int:
        nodes = task.get("nodes", [])
        if not nodes:
            return 0
        total = sum(int(item.get("progress") or 0) for item in nodes)
        return max(0, min(100, int(total / len(nodes))))

    def _add_artifact_locked(self, task: dict[str, Any], path: Path, kind: str) -> dict[str, Any]:
        resolved = path.resolve()
        existing = next((item for item in task.get("artifacts", []) if item.get("path") == str(resolved)), None)
        if existing:
            existing["size_bytes"] = resolved.stat().st_size if resolved.exists() else 0
            existing["updated_at"] = utc_now_iso()
            return existing
        artifact = {
            "id": f"artifact_{uuid4().hex[:12]}",
            "kind": kind,
            "file_name": resolved.name,
            "path": str(resolved),
            "size_bytes": resolved.stat().st_size if resolved.exists() else 0,
            "created_at": utc_now_iso(),
        }
        task.setdefault("artifacts", []).append(artifact)
        return artifact

    @staticmethod
    def _public_template(payload: dict[str, Any]) -> dict[str, Any]:
        return {
            "id": payload["id"],
            "name": payload.get("name"),
            "file_name": payload.get("file_name"),
            "size_bytes": payload.get("size_bytes"),
            "text_preview": payload.get("text_preview"),
            "headings": payload.get("headings", []),
            "table_count": len(payload.get("tables", [])),
            "scene": payload.get("scene") or "通用检测场景",
            "operator": payload.get("operator") or "operator",
            "created_at": payload.get("created_at"),
            "updated_at": payload.get("updated_at"),
        }

    def _all_task_payloads_locked(self) -> list[dict[str, Any]]:
        payloads: list[dict[str, Any]] = []
        for path in self.tasks_dir.glob("*.json"):
            try:
                payloads.append(json.loads(path.read_text(encoding="utf-8")))
            except Exception:
                continue
        payloads.sort(key=lambda item: (str(item.get("created_at") or ""), str(item.get("id") or "")))
        return payloads

    def _task_index_locked(self, task_id: str) -> int | None:
        for index, payload in enumerate(self._all_task_payloads_locked(), start=1):
            if payload.get("id") == task_id:
                return index
        return None

    @staticmethod
    def _task_display_fields(payload: dict[str, Any], index: int) -> dict[str, Any]:
        title = str(payload.get("title") or payload.get("instruction") or payload.get("id") or "采集任务").strip()
        return {"task_number": index, "display_title": f"任务{index}：{title}"}

    @staticmethod
    def _coerce_approval_constraints(current: dict[str, Any], parameters: dict[str, Any]) -> dict[str, Any]:
        result = dict(current or {})

        def has_value(key: str) -> bool:
            return key in parameters and parameters.get(key) not in {None, ""}

        text_keys = {"room_name", "mode", "aggregation", "antenna"}
        numeric_keys = {"freq_start_mhz", "freq_stop_mhz", "freq_step_mhz", "sample_rate", "bandwidth", "gain"}
        integer_keys = {"repeat_count", "expected_fft_frames_per_capture"}
        for key in text_keys:
            if has_value(key):
                result[key] = str(parameters[key]).strip()
        for key in numeric_keys:
            if has_value(key):
                result[key] = float(parameters[key])
        for key in integer_keys:
            if has_value(key):
                result[key] = int(float(parameters[key]))
        if has_value("dwell_ms"):
            result["dwell_time_sec"] = float(parameters["dwell_ms"]) / 1000.0
        elif has_value("dwell_time_sec"):
            result["dwell_time_sec"] = float(parameters["dwell_time_sec"])
        return CapturePlanBuilder.normalize_constraints(result)

    @staticmethod
    def _task_summary(payload: dict[str, Any]) -> dict[str, Any]:
        return {
            "id": payload.get("id"),
            "title": payload.get("title"),
            "instruction": payload.get("instruction"),
            "place": payload.get("place") or "",
            "operator": payload.get("operator") or "operator",
            "source": payload.get("source") or "capture_agent",
            "parent_task_id": payload.get("parent_task_id") or "",
            "reasoning_mode": _normalize_reasoning_mode(payload.get("reasoning_mode")),
            "selected_usrp_devices": payload.get("selected_usrp_devices", []),
            "selected_probe_devices": payload.get("selected_probe_devices", []),
            "available_devices": payload.get("available_devices", {}),
            "device_selection": payload.get("device_selection", {}),
            "probe_evidence_sets": payload.get("probe_evidence_sets", []),
            "probe_assistance": payload.get("probe_assistance", {}),
            "template_ids": payload.get("template_ids", []),
            "templates": payload.get("templates", []),
            "status": payload.get("status"),
            "pause_reason": payload.get("pause_reason"),
            "progress": payload.get("progress", 0),
            "plan_version": payload.get("plan_version", 1),
            "plan_locked": payload.get("plan_locked", False),
            "current_node_id": payload.get("current_node_id"),
            "created_at": payload.get("created_at"),
            "updated_at": payload.get("updated_at"),
        }

    @staticmethod
    def _public_task(payload: dict[str, Any]) -> dict[str, Any]:
        task = _deepcopy_json(payload)
        for node in task.get("nodes", []):
            outputs = node.get("outputs")
            if isinstance(outputs, dict) and "combined_markdown" in outputs:
                outputs["combined_markdown"] = str(outputs["combined_markdown"])[:4000]
        task["generated_code"] = str(task.get("generated_code") or "")
        for artifact in task.get("artifacts", []):
            artifact.pop("path", None)
            base_url = f"/api/capture-agent/tasks/{task['id']}/artifacts/{artifact['id']}"
            artifact["download_url"] = base_url
            if str(artifact.get("file_name") or "").lower().endswith(".npz"):
                artifact["spectrum_url"] = f"{base_url}/spectrum"
        return task

    @staticmethod
    def _device_collection_snapshot(task: dict[str, Any]) -> dict[str, Any]:
        selected_types = list((task.get("device_selection") or {}).get("selected_device_types") or [])
        nodes_by_id = {str(node.get("id") or ""): node for node in task.get("nodes", [])}
        snapshot: dict[str, Any] = {"selected_device_types": selected_types}
        if "usrp" in selected_types:
            snapshot["usrp"] = {
                "selected_devices": [
                    {
                        "device_id": item.get("device_id") or item.get("dev_id") or item.get("id"),
                        "name": item.get("name") or item.get("label") or item.get("product"),
                    }
                    for item in list(task.get("selected_usrp_devices") or [])
                ],
                "nodes": [
                    {"id": node_id, "title": nodes_by_id.get(node_id, {}).get("title"), "status": nodes_by_id.get(node_id, {}).get("status"), "summary": nodes_by_id.get(node_id, {}).get("summary")}
                    for node_id in ("device_scan", "collection_execution", "output_file_check", "execution_summary")
                    if node_id in nodes_by_id
                ],
                "execution_summary": dict(task.get("execution_summary") or {}),
            }
        if "wifi_bluetooth_probe" in selected_types:
            probe_node = nodes_by_id.get("wifi_bluetooth_probe_collection", {})
            snapshot["wifi_bluetooth_probe"] = {
                "selected_devices": [
                    {
                        "probe_id": item.get("probe_id") or item.get("device_id") or item.get("id"),
                        "name": item.get("name") or item.get("label"),
                    }
                    for item in list(task.get("selected_probe_devices") or [])
                ],
                "node": {
                    "status": probe_node.get("status"),
                    "summary": probe_node.get("summary"),
                },
                "evidence_sets": list(task.get("probe_evidence_sets") or []),
                "latest_summary": task.get("latest_probe_summary") or "",
            }
        return snapshot

    @staticmethod
    def _render_markdown_report(task: dict[str, Any]) -> str:
        lines = [
            f"# DeepEM 采集智能体任务报告：{task['id']}",
            "",
            f"- 状态：{task['status']}",
            f"- 计划版本：v{task['plan_version']}",
            f"- 创建时间：{task['created_at']}",
            f"- 指令：{task['instruction']}",
            f"- 模型模式：{'快速模式' if _normalize_reasoning_mode(task.get('reasoning_mode')) == 'fast' else '深度思考模式'}",
            "",
            "## 结构化采集参数",
            "",
            "```json",
            json.dumps(task.get("constraints", {}), ensure_ascii=False, indent=2),
            "```",
            "",
            "## 节点执行记录",
            "",
        ]
        for node in task.get("nodes", []):
            lines.extend(
                [
                    f"### {node['title']} `{node['status']}`",
                    "",
                    node.get("summary") or "",
                    "",
                    f"- 依赖：{', '.join(node.get('dependencies') or []) or '无'}",
                    f"- 允许工具：{node.get('allowed_tool') or '无'}",
                    f"- 开始：{node.get('started_at') or '-'}",
                    f"- 结束：{node.get('ended_at') or '-'}",
                    f"- 人工备注：{node.get('notes') or '-'}",
                    "",
                    "验收条件：",
                ]
            )
            lines.extend([f"- {item}" for item in node.get("success_criteria", [])])
            lines.extend(["", "执行日志："])
            lines.extend([f"- {entry.get('timestamp')} [{entry.get('level')}] {entry.get('message')}" for entry in node.get("logs", [])])
            lines.append("")
        lines.extend(["## 设备选择与执行策略", "", "```json", json.dumps(task.get("device_selection", {}), ensure_ascii=False, indent=2), "```", ""])
        lines.extend(["## 各设备采集情况", "", "本节按本任务实际选中的设备汇总采集状态、执行摘要与证据索引。", "", "```json", json.dumps(CaptureAgentService._device_collection_snapshot(task), ensure_ascii=False, indent=2), "```", ""])
        lines.extend(["## 执行结果", "", "```json", json.dumps(task.get("execution_result", {}), ensure_ascii=False, indent=2), "```", ""])
        lines.extend(["## WiFi/蓝牙探针证据集", "", "```json", json.dumps(task.get("probe_evidence_sets", []), ensure_ascii=False, indent=2), "```", ""])
        legacy_aux = task.get("probe_assistance") or {}
        if legacy_aux.get("enabled"):
            lines.extend(["## 兼容旧任务：探针辅助证据", "", "```json", json.dumps(legacy_aux, ensure_ascii=False, indent=2), "```", ""])
        return "\n".join(lines)
